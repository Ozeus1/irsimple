#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
IRSimple - calculadora de IRPF para renda variavel.

Aplicativo independente para importar os arquivos de negociacao/movimentacao
da B3, complementar dados manualmente e gerar demonstrativos mensais e anuais.

Aviso: o programa automatiza os calculos usuais de bolsa, mas nao substitui a
conferencia do contribuinte/contador, principalmente quando houver notas de
corretagem com custos, emprestimos de ativos, exercicio de opcoes, migracao de
corretora, desdobramentos, grupamentos ou bonificacoes.
"""

from __future__ import annotations

import csv
import json
import math
import os
import re
import sqlite3
import sys
import warnings
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable
from urllib.request import Request, urlopen

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except Exception:  # pragma: no cover - servidores web headless podem nao ter tkinter
    class _MissingTkBase:
        def __init__(self, *_args: Any, **_kwargs: Any) -> None:
            raise RuntimeError("tkinter nao esta instalado. Use a interface web ou instale python3-tk para a interface desktop.")

    class _MissingTk:
        Misc = object
        Toplevel = object
        Tk = _MissingTkBase
        StringVar = _MissingTkBase
        IntVar = _MissingTkBase

    class _MissingTtk:
        Frame = object

        def __getattr__(self, _name: str) -> Any:
            return _MissingTkBase

    class _MissingDialog:
        def __getattr__(self, _name: str) -> Any:
            def _missing(*_args: Any, **_kwargs: Any) -> Any:
                raise RuntimeError("tkinter nao esta instalado. Instale python3-tk para usar a interface desktop.")

            return _missing

    tk = _MissingTk()  # type: ignore[assignment]
    ttk = _MissingTtk()  # type: ignore[assignment]
    filedialog = _MissingDialog()  # type: ignore[assignment]
    messagebox = _MissingDialog()  # type: ignore[assignment]

try:
    import pandas as pd
except Exception:  # pragma: no cover - dependencia local
    pd = None

try:
    import requests
except Exception:  # pragma: no cover - dependencia local
    requests = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except Exception:  # pragma: no cover - dependencia local
    Workbook = None

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
except Exception:  # pragma: no cover - dependencia local
    SimpleDocTemplate = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - dependencia local
    try:
        from PyPDF2 import PdfReader
    except Exception:  # pragma: no cover - dependencia local
        PdfReader = None


APP_TITLE = "IRSimple - IRPF Bolsa"
BASE_DIR = Path(__file__).resolve().parent
LOG_FILE = BASE_DIR / "irsimple_debug.log"
OPTION_EVENTS_FILE = BASE_DIR / "conferencias_opcoes.json"
CONFIG_FILE = BASE_DIR / "irsimple_config.json"
DB_FILE = BASE_DIR / "irsimple.db"
DEFAULT_USER = os.environ.get("IRSIMPLE_USER", "local").strip() or "local"
CNPJ_STOCKS_URL = "https://www.idinheiro.com.br/investimentos/cnpj-empresas-listadas-b3/"
CNPJ_FIIS_URL = "https://www.empiricus.com.br/explica/cnpj-fundos-imobiliarios-fiis-listados-b3-declaracao-imposto-de-renda/"
CNPJ_ETFS_URL = "https://maisretorno.com/lista-etf"
BCB_SGS_URL = "https://api.bcb.gov.br/dados/serie/bcdata.sgs.{code}/dados"
YAHOO_CHART_URL = "https://query1.finance.yahoo.com/v8/finance/chart/%5EBVSP"
CNPJ_RE = re.compile(r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}")
KNOWN_FII_TICKERS = {"ITIT11"}
KNOWN_ETF_TICKERS = {"BOVA11", "BOVV11", "HASH11", "GOLD11", "IVVB11", "LFTS11"}
MONTHS = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
MONTH_NAME_TO_NUMBER = {
    "janeiro": 1,
    "fevereiro": 2,
    "marco": 3,
    "março": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}

warnings.filterwarnings("ignore", message="Workbook contains no default style, apply openpyxl's default")


def debug_log(message: str) -> None:
    try:
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with LOG_FILE.open("a", encoding="utf-8") as fh:
            fh.write(f"[{timestamp}] {message}\n")
    except Exception:
        pass


def latest_matching_file(patterns: list[str], fallback: Path) -> Path:
    candidates: list[Path] = []
    for folder in [Path.home() / "Downloads", BASE_DIR]:
        for pattern in patterns:
            candidates.extend(folder.glob(pattern))
    if not candidates:
        return fallback
    return max(candidates, key=lambda path: path.stat().st_mtime)


def fetch_url_text(url: str) -> str:
    if requests is not None:
        response = requests.get(url, headers={"User-Agent": "Mozilla/5.0 IRSimple/1.0"}, timeout=40)
        response.raise_for_status()
        return response.text
    req = Request(url, headers={"User-Agent": "Mozilla/5.0 IRSimple/1.0"})
    with urlopen(req, timeout=40) as response:
        return response.read().decode("utf-8", errors="ignore")


def fetch_json_url(url: str, params: dict[str, Any] | None = None) -> Any:
    if requests is not None:
        response = requests.get(url, params=params, headers={"User-Agent": "Mozilla/5.0 IRSimple/1.0"}, timeout=40)
        response.raise_for_status()
        return response.json()
    if params:
        query = "&".join(f"{key}={value}" for key, value in params.items())
        url = f"{url}?{query}"
    req = Request(url, headers={"User-Agent": "Mozilla/5.0 IRSimple/1.0"})
    with urlopen(req, timeout=40) as response:
        return json.loads(response.read().decode("utf-8", errors="ignore"))


def month_key(dt: date) -> str:
    return f"{dt.year}-{dt.month:02d}"


def parse_month_key(period: str) -> tuple[int, int] | None:
    match = re.fullmatch(r"(\d{4})(?:-(\d{2}))?", str(period))
    if not match:
        return None
    return int(match.group(1)), int(match.group(2) or "12")


def add_months(year: int, month: int, offset: int) -> tuple[int, int]:
    value = year * 12 + month - 1 + offset
    return value // 12, value % 12 + 1


def iter_month_keys(start_year: int, start_month: int, end_year: int, end_month: int) -> list[str]:
    keys: list[str] = []
    year, month = start_year, start_month
    while (year, month) <= (end_year, end_month):
        keys.append(f"{year}-{month:02d}")
        year, month = add_months(year, month, 1)
    return keys


def fetch_bcb_monthly_percent(code: int, start: date, end: date) -> dict[str, Decimal]:
    data = fetch_json_url(
        BCB_SGS_URL.format(code=code),
        {
            "formato": "json",
            "dataInicial": start.strftime("%d/%m/%Y"),
            "dataFinal": end.strftime("%d/%m/%Y"),
        },
    )
    values: dict[str, Decimal] = {}
    for item in data or []:
        dt = parse_date(item.get("data"))
        if dt is None:
            continue
        key = month_key(dt)
        if key not in values or dt.day == 1:
            values[key] = money(item.get("valor"))
    return values


def fetch_ibovespa_monthly_close(start: date, end: date) -> dict[str, Decimal]:
    period1 = int(datetime(start.year, start.month, start.day).timestamp())
    period2 = int(datetime(end.year, end.month, end.day, 23, 59, 59).timestamp())
    data = fetch_json_url(YAHOO_CHART_URL, {"period1": period1, "period2": period2, "interval": "1mo"})
    result = ((data or {}).get("chart") or {}).get("result") or []
    if not result:
        return {}
    payload = result[0]
    timestamps = payload.get("timestamp") or []
    closes = (((payload.get("indicators") or {}).get("quote") or [{}])[0]).get("close") or []
    values: dict[str, Decimal] = {}
    for ts, close in zip(timestamps, closes):
        if close is None:
            continue
        dt = datetime.fromtimestamp(int(ts)).date()
        values[month_key(dt)] = q2(close)
    return values


def parse_asset_cnpj_sources() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    if pd is None:
        raise RuntimeError("A biblioteca pandas nao esta instalada.")

    stocks_html = fetch_url_text(CNPJ_STOCKS_URL)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", FutureWarning)
        stock_tables = pd.read_html(stocks_html)
    for table in stock_tables:
        columns = [normalize_header(col) for col in table.columns]
        table.columns = columns
        if "codigo_s" not in columns or "cnpj" not in columns:
            continue
        name_col = "empresa" if "empresa" in columns else columns[0]
        for _, row in table.iterrows():
            cnpj = str(row.get("cnpj") or "").strip()
            if not CNPJ_RE.fullmatch(cnpj):
                continue
            for ticker in re.findall(r"[A-Z]{4}\d{1,2}", str(row.get("codigo_s") or "").upper()):
                rows.append({"ticker": normalize_ticker(ticker), "cnpj": cnpj, "nome": str(row.get(name_col) or ""), "tipo": "acao", "fonte": CNPJ_STOCKS_URL})

    fii_html = fetch_url_text(CNPJ_FIIS_URL)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", FutureWarning)
        fii_tables = pd.read_html(fii_html)
    for table in fii_tables:
        if table.shape[1] < 3:
            continue
        for _, row in table.iloc[1:].iterrows():
            name = str(row.iloc[0] or "")
            ticker = normalize_ticker(row.iloc[1])
            cnpj = str(row.iloc[2] or "").strip()
            if ticker and CNPJ_RE.fullmatch(cnpj):
                rows.append({"ticker": ticker, "cnpj": cnpj, "nome": name, "tipo": "fii", "fonte": CNPJ_FIIS_URL})

    etf_html = fetch_url_text(CNPJ_ETFS_URL)
    etf_pattern = re.compile(
        r'>([A-Z0-9]{4,12})</p></a><p[^>]*>(.*?)</p>.*?<p[^>]*>CNPJ</p><p[^>]*>(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})</p>',
        re.S,
    )
    for ticker, name, cnpj in etf_pattern.findall(etf_html):
        rows.append({"ticker": normalize_ticker(ticker), "cnpj": cnpj, "nome": re.sub(r"<[^>]+>", "", name).strip(), "tipo": "etf", "fonte": CNPJ_ETFS_URL})

    dedup: dict[str, dict[str, str]] = {}
    for row in rows:
        if row["ticker"]:
            dedup[row["ticker"]] = row
    return sorted(dedup.values(), key=lambda item: item["ticker"])


def parse_consolidated_period(path: Path) -> tuple[int, int] | None:
    name = path.stem.lower()
    annual = re.search(r"relatorio-consolidado-anual-(\d{4})", name)
    if annual:
        return int(annual.group(1)), 12
    monthly = re.search(r"relatorio-consolidado-mensal-(\d{4})-([a-zç]+)", name)
    if monthly:
        month = MONTH_NAME_TO_NUMBER.get(monthly.group(2))
        if month:
            return int(monthly.group(1)), month
    return None


def read_consolidated_positions(path: Path) -> dict[str, dict[str, Decimal]]:
    if pd is None:
        raise RuntimeError("A biblioteca pandas nao esta instalada.")
    result = {"carteira": defaultdict(Decimal), "emprestimos": defaultdict(Decimal), "opcoes": defaultdict(Decimal)}
    xls = pd.ExcelFile(path)
    position_sheets = {
        "posicao_acoes",
        "posicao_bdr",
        "posicao_etf",
        "posicao_fundos",
    }
    for sheet in xls.sheet_names:
        norm_sheet = normalize_header(sheet)
        df = pd.read_excel(path, sheet_name=sheet, dtype=object)
        df.columns = [normalize_header(col) for col in df.columns]
        if "codigo_de_negociacao" not in df.columns or "quantidade" not in df.columns:
            continue
        if norm_sheet in position_sheets:
            bucket = "carteira"
        elif norm_sheet == "posicao_emprestimos":
            bucket = "emprestimos"
        elif norm_sheet == "posicao_opcoes":
            bucket = "opcoes"
        else:
            continue
        for _, row in df.iterrows():
            code = normalize_ticker(row.get("codigo_de_negociacao"))
            qty = money(row.get("quantidade"))
            if not code or qty <= 0:
                continue
            result[bucket][code] += qty
    return {bucket: dict(values) for bucket, values in result.items()}


def fmt_sheet_cell(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd is not None and pd.isna(value):
            return ""
    except Exception:
        pass
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, Decimal):
        return fmt_decimal(value)
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return ""
        if value == int(value):
            return str(int(value))
        return str(Decimal(str(value)).normalize()).replace(".", ",")
    return str(value)


MONEY_RE = r"\d{1,3}(?:\.\d{3})*,\d{2}"
NOTE_MONEY_RE = r"-?\s*(?:R\$\s*)?\d{1,3}(?:\.\d{3})*,\d{2}"
NOTE_PRICE_RE = r"\d{1,3}(?:\.\d{3})*,\d{2,4}"


def note_money(value: Any) -> Decimal:
    text = str(value or "").replace("R$", "").replace(" ", "")
    return q2(abs(money(text)))


def note_money_values(text: str) -> list[Decimal]:
    return [note_money(value) for value in re.findall(NOTE_MONEY_RE, text or "")]


def note_price(value: Any) -> Decimal:
    text = str(value or "").replace("R$", "").replace(" ", "")
    return money(text)


def extract_first_money_before(text: str, marker: str) -> Decimal:
    idx = text.find(marker)
    if idx < 0:
        return Decimal("0")
    values = note_money_values(text[max(0, idx - 140):idx])
    return values[-2] if len(values) >= 2 else values[-1] if values else Decimal("0")


def parse_note_header(text: str, path: Path, page_idx: int) -> tuple[str, date | None]:
    if "COMPROVANTE BOVESPA" in text.upper():
        date_match = re.search(r"(\d{2}/\d{2}/\d{4})\s*Data de refer", text, re.I)
        note_match = re.search(r"Comprovante\s+(\d{3,12})", text, re.I)
        return (note_match.group(1) if note_match else f"{path.stem}-{page_idx}", parse_date(date_match.group(1)) if date_match else None)
    dates = [parse_date(value) for value in re.findall(r"\d{2}/\d{2}/\d{4}", text)]
    dates = [value for value in dates if value is not None]
    trade_date = dates[0] if dates else None
    note_number = ""
    for pattern in [
        r"Nr\.?\s*[Nn]ota\s*(?:\d+\s+)?(\d{4,12})",
        r"NOTA DE CORRETAGEM\s+(\d{4,12})",
        r"\b(\d{4,12})\s+(?:GENIAL|CM CAPITAL|C\.N\.P\.J|Data preg)",
    ]:
        match = re.search(pattern, text, re.I)
        if match:
            note_number = match.group(1).replace(".", "")
            break
    if not note_number and trade_date:
        date_pos = text.find(trade_date.strftime("%d/%m/%Y"))
        tail = text[date_pos + 10: date_pos + 48] if date_pos >= 0 else ""
        match = re.search(r"(\d{4,12})", tail)
        if match:
            note_number = match.group(1)
    return note_number or f"{path.stem}-{page_idx}", trade_date


def parse_irrf_common(text: str) -> tuple[Decimal, Decimal]:
    value = Decimal("0")
    base = Decimal("0")

    match = re.search(rf"I\.?\s*R\.?\s*R\.?\s*F\.?\s*s/\s*opera\S*,?\s*base\s*R\$\s*({MONEY_RE})\s*({MONEY_RE})?", text, re.I)
    if match:
        base = note_money(match.group(1))
        value = note_money(match.group(2)) if match.group(2) else extract_first_money_before(text, match.group(0))

    toro_value = re.search(rf"IRRF\s+Opera\S*\s+Comum\s*(?:R\$)?\s*({NOTE_MONEY_RE})", text, re.I)
    toro_base = re.search(rf"IRRF\s+Opera\S*\s+Comum:\s*Base\s*({NOTE_MONEY_RE})", text, re.I)
    toro_projection = re.search(
        rf"IRRF\s+Opera\S*\s+Comum:\s*Base\s*{NOTE_MONEY_RE}\s*Proje\S*\s*({NOTE_MONEY_RE})",
        text,
        re.I,
    )
    if toro_base:
        base = note_money(toro_base.group(1))
    if toro_value:
        value = note_money(toro_value.group(1))
    if value == 0 and toro_projection:
        value = note_money(toro_projection.group(1))

    return q2(value), q2(base)


def parse_irrf_daytrade(text: str) -> Decimal:
    toro_projection = re.search(
        rf"IRRF\s+Day[- ]?Trade:\s*Base\s*{NOTE_MONEY_RE}\s*Proje\S*\s*({NOTE_MONEY_RE})",
        text,
        re.I,
    )
    toro_match = re.search(rf"IRRF\s+Day[- ]?Trade\s*(?:R\$)?\s*({NOTE_MONEY_RE})", text, re.I)
    if toro_match:
        value = note_money(toro_match.group(1))
        if value > 0:
            return q2(value)
        if toro_projection:
            return q2(note_money(toro_projection.group(1)))
    label = re.search(r"IRRF\s+IRRF\s+Day Trade", text, re.I)
    if label:
        values = note_money_values(text[label.end(): label.end() + 180])
        if len(values) >= 2:
            return q2(values[1])
    label = re.search(r"(?:I\.?\s*R\.?\s*R\.?\s*F\.?|IRRF)\s+Day[- ]?Trade\s*\([^)]*Proje", text, re.I)
    if label:
        values = note_money_values(text[label.end(): label.end() + 180])
        if values:
            return q2(values[0])
    if toro_projection:
        return q2(note_money(toro_projection.group(1)))
    return Decimal("0")


def add_brokerage_summary_from_labels(text: str, summary: defaultdict[str, Decimal]) -> None:
    label_map = [
        ("buy_normal", r"Compras à vista\s*(?:R\$)?\s*(" + MONEY_RE + ")"),
        ("sell_normal", r"Vendas à vista\s*(?:R\$)?\s*(" + MONEY_RE + ")"),
        ("buy_options", r"Opções\s*-\s*Compras\s*(" + MONEY_RE + ")"),
        ("sell_options", r"Opções\s*-\s*Vendas\s*(" + MONEY_RE + ")"),
    ]
    for key, pattern in label_map:
        match = re.search(pattern, text, re.I)
        if match and summary[key] == 0:
            summary[key] = money(match.group(1))


def parse_brokerage_note_pdfs(paths: list[Path]) -> list[dict[str, Any]]:
    if PdfReader is None:
        raise RuntimeError("Instale pypdf ou PyPDF2 para importar notas de corretagem em PDF.")
    notes: list[dict[str, Any]] = []
    for path in paths:
        reader = PdfReader(str(path))
        for page_idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            text_upper = text.upper()
            if "NOTA DE CORRETAGEM" not in text_upper and "COMPROVANTE BOVESPA" not in text_upper:
                continue
            note_number, trade_date = parse_note_header(text, path, page_idx)
            if trade_date is None:
                continue
            broker_match = re.search(r"\n([A-Z0-9 .&/-]+(?:CTVM|DTVM|CORRETORA|INVESTIMENTOS)[A-Z0-9 .&/-]*)\n", text, re.I)
            broker = broker_match.group(1).strip() if broker_match else ""
            irrf_common_total, irrf_base = parse_irrf_common(text)
            if "COMPROVANTE BOVESPA" in text_upper:
                common_value_match = re.search(rf"IRRF Operação Comum\s*R\$\s*({MONEY_RE})", text, re.I)
                common_base_match = re.search(rf"IRRF Operação Comum:\s*Base\s*R\$\s*({MONEY_RE})", text, re.I)
                irrf_common_total = money(common_value_match.group(1)) if common_value_match else irrf_common_total
                irrf_base = money(common_base_match.group(1)) if common_base_match else irrf_base
            if "COMPROVANTE BOVESPA" in text_upper:
                common_projection_match = re.search(
                    rf"IRRF\s+Opera\S*\s+Comum:\s*Base\s*{NOTE_MONEY_RE}\s*Proje\S*\s*({NOTE_MONEY_RE})",
                    text,
                    re.I,
                )
                if irrf_common_total == 0 and common_projection_match:
                    irrf_common_total = note_money(common_projection_match.group(1))
            irrf_daytrade = parse_irrf_daytrade(text)
            sold_fii = Decimal("0")
            sold_other = Decimal("0")
            summary = defaultdict(Decimal)
            for line in text.splitlines():
                if not line.startswith("1-BOVESPA ") and not line.strip() in {"COMPRA", "VENDA"}:
                    continue
                parts = line.split()
                if len(parts) < 8:
                    continue
                value = money(parts[-2])
                code = ""
                for token in parts[3:-3]:
                    if re.fullmatch(r"[A-Z]{4,5}\d{1,4}[A-Z0-9]*", token):
                        code = normalize_ticker(token)
                category = classify_asset(code, line)
                bucket = "fii" if category == "fii" or (code.endswith("11") and category not in {"etf", "opcoes", "futuro"}) else "options" if category == "opcoes" else "normal"
                side_key = "buy" if parts[1] == "C" else "sell" if parts[1] == "V" else ""
                if side_key:
                    summary[f"{side_key}_{bucket}"] += value
                    if side_key == "sell":
                        if bucket == "fii":
                            sold_fii += value
                        else:
                            sold_other += value
            add_brokerage_summary_from_labels(text, summary)
            total_sold = sold_fii + sold_other
            if total_sold > 0 and irrf_common_total > 0:
                irrf_fii = q2(irrf_common_total * sold_fii / total_sold)
                irrf_common = q2(irrf_common_total - irrf_fii)
            else:
                irrf_fii = Decimal("0")
                irrf_common = irrf_common_total
            notes.append(
                {
                    "note_number": note_number,
                    "trade_date": trade_date,
                    "broker": broker,
                    "source_file": str(path),
                    "page": page_idx,
                    "irrf_common": q2(irrf_common),
                    "irrf_daytrade": q2(irrf_daytrade),
                    "irrf_fii": q2(irrf_fii),
                    "irrf_base": q2(irrf_base),
                    "buy_normal": q2(summary["buy_normal"]),
                    "sell_normal": q2(summary["sell_normal"]),
                    "buy_fii": q2(summary["buy_fii"]),
                    "sell_fii": q2(summary["sell_fii"]),
                    "buy_options": q2(summary["buy_options"]),
                    "sell_options": q2(summary["sell_options"]),
                }
            )
    return notes


def read_consolidated_sheets(path: Path) -> list[dict[str, Any]]:
    if pd is None:
        raise RuntimeError("A biblioteca pandas nao esta instalada.")
    sheets: list[dict[str, Any]] = []
    xls = pd.ExcelFile(path)
    for sheet in xls.sheet_names:
        df = pd.read_excel(path, sheet_name=sheet, dtype=object)
        df = df.dropna(how="all").dropna(axis=1, how="all")
        columns = [str(col) for col in df.columns]
        if not columns:
            columns = ["Sem dados"]
            rows = [[""]]
        else:
            rows = [[fmt_sheet_cell(value) for value in row] for row in df.itertuples(index=False, name=None)]
            if not rows:
                rows = [["" for _ in columns]]
        sheets.append({"name": sheet, "columns": columns, "rows": rows})
    return sheets


def read_config() -> dict[str, Any]:
    if not CONFIG_FILE.exists():
        return {}
    try:
        return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def write_config(data: dict[str, Any]) -> None:
    CONFIG_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def db_connect() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    with db_connect() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL UNIQUE,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS app_config (
                user_id INTEGER NOT NULL,
                key TEXT NOT NULL,
                value TEXT,
                PRIMARY KEY (user_id, key),
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS source_files (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                file_type TEXT NOT NULL,
                path TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(user_id, file_type, path),
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS option_conferences (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                data TEXT,
                tipo TEXT,
                ativo TEXT,
                quantidade TEXT,
                valor TEXT,
                categoria TEXT,
                observacao TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS trades (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                dt TEXT NOT NULL,
                side TEXT,
                market TEXT,
                broker TEXT,
                code TEXT,
                qty TEXT,
                price TEXT,
                value TEXT,
                category TEXT,
                expiry TEXT,
                source_file TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS movements (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                dt TEXT NOT NULL,
                direction TEXT,
                kind TEXT,
                product TEXT,
                broker TEXT,
                code TEXT,
                qty TEXT,
                unit_price TEXT,
                value TEXT,
                category TEXT,
                source_file TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS manual_positions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                ativo TEXT,
                quantidade TEXT,
                custo TEXT,
                categoria TEXT,
                corretora TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS manual_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                data TEXT,
                tipo TEXT,
                ativo TEXT,
                quantidade TEXT,
                valor TEXT,
                categoria TEXT,
                observacao TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS asset_cnpjs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                ticker TEXT NOT NULL,
                cnpj TEXT NOT NULL,
                nome TEXT,
                tipo TEXT,
                fonte TEXT,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(user_id, ticker),
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS brokerage_note_taxes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                note_number TEXT,
                trade_date TEXT NOT NULL,
                broker TEXT,
                source_file TEXT,
                page INTEGER,
                irrf_common TEXT,
                irrf_daytrade TEXT,
                irrf_fii TEXT,
                irrf_base TEXT,
                buy_normal TEXT,
                sell_normal TEXT,
                buy_fii TEXT,
                sell_fii TEXT,
                buy_options TEXT,
                sell_options TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(user_id, note_number, trade_date, source_file, page),
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS password_reset_tokens (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                email TEXT NOT NULL,
                token TEXT NOT NULL UNIQUE,
                expires_at TEXT NOT NULL,
                used INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS smtp_config (
                id INTEGER PRIMARY KEY CHECK (id = 1),
                host TEXT NOT NULL DEFAULT '',
                port INTEGER NOT NULL DEFAULT 587,
                secure INTEGER NOT NULL DEFAULT 0,
                username TEXT NOT NULL DEFAULT '',
                password TEXT NOT NULL DEFAULT '',
                from_addr TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS custody_transfers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                transfer_date TEXT NOT NULL,
                protocol TEXT NOT NULL DEFAULT '',
                broker_from TEXT NOT NULL DEFAULT '',
                account_from TEXT NOT NULL DEFAULT '',
                broker_to TEXT NOT NULL DEFAULT '',
                account_to TEXT NOT NULL DEFAULT '',
                ticker TEXT NOT NULL,
                asset_type TEXT NOT NULL DEFAULT '',
                quantity TEXT NOT NULL DEFAULT '0',
                status TEXT NOT NULL DEFAULT 'finalizado',
                obs TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS corporate_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                event_date TEXT NOT NULL,
                event_type TEXT NOT NULL,
                ticker TEXT NOT NULL,
                ticker_new TEXT NOT NULL DEFAULT '',
                factor TEXT NOT NULL DEFAULT '1',
                bonus_qty TEXT NOT NULL DEFAULT '0',
                bonus_cost TEXT NOT NULL DEFAULT '0',
                broker_from TEXT NOT NULL DEFAULT '',
                broker_to TEXT NOT NULL DEFAULT '',
                obs TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            );
            """
        )
        existing_bn = {row["name"] for row in conn.execute("PRAGMA table_info(brokerage_note_taxes)").fetchall()}
        for column in ["buy_normal", "sell_normal", "buy_fii", "sell_fii", "buy_options", "sell_options"]:
            if column not in existing_bn:
                conn.execute(f"ALTER TABLE brokerage_note_taxes ADD COLUMN {column} TEXT")
        existing_u = {row["name"] for row in conn.execute("PRAGMA table_info(users)").fetchall()}
        for col_def in [
            ("role", "TEXT NOT NULL DEFAULT 'user'"),
            ("full_name", "TEXT NOT NULL DEFAULT ''"),
            ("whatsapp", "TEXT NOT NULL DEFAULT ''"),
            ("password_hash", "TEXT NOT NULL DEFAULT ''"),
            ("photo", "TEXT NOT NULL DEFAULT ''"),
            ("email", "TEXT NOT NULL DEFAULT ''"),
        ]:
            if col_def[0] not in existing_u:
                conn.execute(f"ALTER TABLE users ADD COLUMN {col_def[0]} {col_def[1]}")
        conn.execute("INSERT OR IGNORE INTO users (username) VALUES (?)", (DEFAULT_USER,))


def get_user_id(username: str = DEFAULT_USER) -> int:
    init_db()
    with db_connect() as conn:
        row = conn.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()
        if row:
            return int(row["id"])
        cur = conn.execute("INSERT INTO users (username) VALUES (?)", (username,))
        return int(cur.lastrowid)


def money(value: Any) -> Decimal:
    if value is None or value == "" or value == "-":
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return Decimal("0")
        return Decimal(str(value))
    text = str(value).strip().replace("\xa0", " ")
    text = re.sub(r"[R$\s]", "", text)
    if not text or text == "-":
        return Decimal("0")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except Exception:
        return Decimal("0")


def q2(value: Any) -> Decimal:
    return money(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def fmt_money(value: Any) -> str:
    val = q2(value)
    sign = "-" if val < 0 else ""
    val = abs(val)
    whole, cents = f"{val:.2f}".split(".")
    groups = []
    while whole:
        groups.append(whole[-3:])
        whole = whole[:-3]
    return f"{sign}{'.'.join(reversed(groups))},{cents}"


def fmt_decimal(value: Any) -> str:
    val = money(value)
    if val == val.to_integral():
        return str(int(val))
    return str(val.normalize()).replace(".", ",")


def parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    digits = re.sub(r"\D", "", text)
    if re.fullmatch(r"\d{8}", text):
        text = f"{digits[:2]}/{digits[2:4]}/{digits[4:]}"
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            pass
    return None


def normalize_header(text: Any) -> str:
    value = str(text or "").strip().lower()
    replacements = {
        "á": "a", "à": "a", "ã": "a", "â": "a",
        "é": "e", "ê": "e", "í": "i", "ó": "o",
        "ô": "o", "õ": "o", "ú": "u", "ç": "c",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    return re.sub(r"[^a-z0-9]+", "_", value).strip("_")


def asset_from_product(text: str) -> str:
    text = str(text or "").strip()
    match = re.search(r"\b([A-Z]{4}\d{1,2}[A-Z]?)\b", text)
    if match:
        return normalize_ticker(match.group(1))
    if " - " in text:
        return normalize_ticker(text.split(" - ", 1)[0].strip())
    return normalize_ticker(text[:20].strip().upper())


def normalize_ticker(code: Any) -> str:
    value = str(code or "").strip().upper()
    # A B3 usa o sufixo F para mercado fracionario. Para IR, ele compoe
    # a mesma posicao/custo medio do lote padrao.
    if re.fullmatch(r"[A-Z]{4}\d{1,2}F", value):
        return value[:-1]
    return value


def normalize_option_exercise_code(code: Any) -> str:
    value = str(code or "").strip().upper()
    if value.endswith("E") and re.match(r"^[A-Z]{5}", value):
        return value[:-1]
    return value


def option_root(code: Any) -> str:
    value = normalize_option_exercise_code(code)
    return value[:4] if re.match(r"^[A-Z]{4}", value) else value[:4]


def classify_asset(code: str, market: str = "", product: str = "") -> str:
    code = normalize_ticker(code)
    text = f"{market} {product}".upper()
    norm_text = normalize_header(text).upper()
    if "OPCAO" in norm_text or "OPÇÃO" in text:
        return "opcoes"
    if "FUTURO" in text or re.fullmatch(r"(WIN|IND|WDO|DOL)[FGHJKMNQUVXZ]\d{2}", code):
        return "futuro"
    if "ETF" in norm_text or code in KNOWN_ETF_TICKERS:
        return "etf"
    if code in KNOWN_FII_TICKERS or any(marker in norm_text for marker in ["FII", "FIAGRO", "FDO_INV_IMOB", "FUNDO_DE_INVESTIMENTO_IMOBILIARIO"]):
        return "fii"
    if re.fullmatch(r"[A-Z]{4}3[0-9]", code) or re.fullmatch(r"[A-Z]{4}34", code) or re.fullmatch(r"[A-Z]{4}39", code):
        return "bdr"
    return "normal"


def refine_asset_category(code: str, current: str = "", market: str = "", product: str = "") -> str:
    detected = classify_asset(code, market, product)
    current = (current or "").strip() or detected
    if detected in {"fii", "etf", "bdr", "opcoes", "futuro"}:
        return detected
    return current


def annual_asset_group(pos: Position) -> str:
    if pos.category == "fii":
        return "07 - Fundos imobiliarios/FIAGRO"
    if pos.category == "etf":
        return "74 - Fundos de indice (ETF)"
    if pos.category == "bdr":
        return "04 - BDR"
    return "03 - Acoes"


def infer_categories_from_movements(movements: list["Movement"]) -> dict[str, str]:
    categories: dict[str, str] = {}
    for mov in movements:
        if mov.category in {"fii", "opcoes", "futuro"}:
            categories[mov.code] = mov.category
    return categories


def clone_positions(positions: dict[str, Position]) -> dict[str, Position]:
    return {
            code: Position(
                code=pos.code,
                qty=q2(pos.qty),
                cost=q2(pos.cost),
                previous_qty=q2(pos.previous_qty),
                previous_cost=q2(pos.previous_cost),
                category=pos.category,
                broker=pos.broker,
        )
        for code, pos in positions.items()
    }


@dataclass
class Trade:
    dt: date
    side: str
    market: str
    broker: str
    code: str
    qty: Decimal
    price: Decimal
    value: Decimal
    category: str
    expiry: date | None = None


@dataclass
class Movement:
    dt: date
    direction: str
    kind: str
    product: str
    broker: str
    code: str
    qty: Decimal
    unit_price: Decimal
    value: Decimal
    category: str


@dataclass
class Position:
    code: str
    qty: Decimal = Decimal("0")
    cost: Decimal = Decimal("0")
    previous_qty: Decimal = Decimal("0")
    previous_cost: Decimal = Decimal("0")
    category: str = "normal"
    broker: str = ""

    @property
    def avg_price(self) -> Decimal:
        if self.qty == 0:
            return Decimal("0")
        return self.cost / self.qty


@dataclass
class MonthlyTax:
    month: int
    normal_result: Decimal = Decimal("0")
    daytrade_result: Decimal = Decimal("0")
    fii_result: Decimal = Decimal("0")
    options_result: Decimal = Decimal("0")
    future_result: Decimal = Decimal("0")
    future_dollar_common: Decimal = Decimal("0")
    future_index_common: Decimal = Decimal("0")
    future_dollar_daytrade: Decimal = Decimal("0")
    future_index_daytrade: Decimal = Decimal("0")
    normal_sales: Decimal = Decimal("0")
    daytrade_sales: Decimal = Decimal("0")
    fii_sales: Decimal = Decimal("0")
    normal_loss_before: Decimal = Decimal("0")
    daytrade_loss_before: Decimal = Decimal("0")
    fii_loss_before: Decimal = Decimal("0")
    options_loss_before: Decimal = Decimal("0")
    future_loss_before: Decimal = Decimal("0")
    normal_base: Decimal = Decimal("0")
    daytrade_base: Decimal = Decimal("0")
    fii_base: Decimal = Decimal("0")
    options_base: Decimal = Decimal("0")
    future_base: Decimal = Decimal("0")
    normal_loss_after: Decimal = Decimal("0")
    daytrade_loss_after: Decimal = Decimal("0")
    fii_loss_after: Decimal = Decimal("0")
    options_loss_after: Decimal = Decimal("0")
    future_loss_after: Decimal = Decimal("0")
    tax_due: Decimal = Decimal("0")
    tax_payable: Decimal = Decimal("0")
    irrf_common_month: Decimal = Decimal("0")
    irrf_common_before: Decimal = Decimal("0")
    irrf_common_after: Decimal = Decimal("0")
    irrf_daytrade_month: Decimal = Decimal("0")
    irrf_daytrade_before: Decimal = Decimal("0")
    irrf_daytrade_after: Decimal = Decimal("0")
    irrf_fii_month: Decimal = Decimal("0")
    irrf_fii_before: Decimal = Decimal("0")
    irrf_fii_after: Decimal = Decimal("0")
    exempt_stock_gain: Decimal = Decimal("0")
    warnings: list[str] = field(default_factory=list)


@dataclass
class BrokerageNoteTax:
    dt: date
    irrf_common: Decimal = Decimal("0")
    irrf_daytrade: Decimal = Decimal("0")
    irrf_fii: Decimal = Decimal("0")
    note_number: str = ""
    source_file: str = ""


@dataclass
class CalculationResult:
    year: int
    monthly: list[MonthlyTax]
    positions: dict[str, Position]
    exempt_income: list[dict[str, Any]]
    taxable_income: list[dict[str, Any]]
    debts: list[dict[str, Any]]
    warnings: list[str]
    monthly_positions: dict[int, dict[str, Position]] = field(default_factory=dict)
    monthly_loans: dict[int, dict[str, Decimal]] = field(default_factory=dict)
    pending_options: list[dict[str, Any]] = field(default_factory=list)


class IRSimpleEngine:
    def __init__(self, year: int, asset_cnpjs: dict[str, str] | None = None) -> None:
        self.year = year
        self.asset_cnpjs = asset_cnpjs or {}

    def calculate(
        self,
        trades: list[Trade],
        movements: list[Movement],
        initial_positions: list[Position],
        initial_losses: dict[str, Decimal],
        events: list[dict[str, Any]],
        dated_losses: dict[str, Decimal] | None = None,
        dated_loss_start: date | None = None,
        brokerage_taxes: list[BrokerageNoteTax] | None = None,
        custody_transfers: list[dict[str, Any]] | None = None,
        corporate_events: list[dict[str, Any]] | None = None,
    ) -> CalculationResult:
        positions: dict[str, Position] = {}
        warnings: list[str] = []
        for pos in initial_positions:
            pos.code = normalize_ticker(pos.code)
            if not pos.code:
                continue
            positions[pos.code] = Position(
                code=pos.code,
                qty=money(pos.qty),
                cost=q2(pos.cost),
                previous_qty=money(pos.qty),
                previous_cost=q2(pos.cost),
                category=pos.category or "normal",
                broker=pos.broker,
            )

        monthly = [MonthlyTax(month=i) for i in range(1, 13)]
        cutoff = self._cutoff_date(trades, movements)
        year_trades = [t for t in trades if t.dt.year == self.year]
        year_trades.sort(key=lambda t: (t.dt, t.code, t.side))

        if corporate_events:
            self._apply_corporate_events(positions, corporate_events, self.year, warnings)
        self._apply_events_before_trades(positions, events, warnings)
        trades_after_daytrade = self._split_daytrades(year_trades, monthly)
        normal_trades = [trade for trade in trades_after_daytrade if trade.category != "opcoes"]
        option_trades_for_year = [trade for trade in trades_after_daytrade if trade.category == "opcoes"]
        trades_for_options = [trade for trade in trades if trade.dt.year != self.year] + option_trades_for_year
        # atualiza broker pela ultima compra do ano (antes de aplicar transferencias de custodia)
        self._update_broker_from_trades(positions, year_trades)
        monthly_positions = self._process_normal_trades(normal_trades, positions, monthly, warnings, events)
        pending_options = self._process_option_trades(trades_for_options, monthly, warnings, events, cutoff)
        pending_options.extend(self._process_option_exercises(trades_for_options, positions, monthly_positions, monthly, warnings, cutoff, events))
        # transferencias de custodia aplicadas por ultimo — tem prioridade sobre trades
        if custody_transfers:
            self._apply_custody_transfers(positions, custody_transfers, self.year, warnings)
        self._apply_income_movements(movements, monthly)

        self._apply_monthly_tax(monthly, initial_losses, dated_losses or {}, dated_loss_start)
        self._apply_withholding(monthly, brokerage_taxes or [])
        exempt_income, taxable_income, debts = self._annual_tables(movements, positions)
        monthly_loans = self._monthly_loans(movements)
        return CalculationResult(self.year, monthly, positions, exempt_income, taxable_income, debts, warnings, monthly_positions, monthly_loans, pending_options)

    def _trade_detail(self, trade: Trade) -> str:
        expiry = f"; vencimento={trade.expiry:%d/%m/%Y}" if trade.expiry else ""
        return (
            f"Evento planilha negociacao: data={trade.dt:%d/%m/%Y}; tipo={trade.side}; mercado={trade.market}; "
            f"codigo={trade.code}; qtd={fmt_decimal(trade.qty)}; preco=R$ {fmt_money(trade.price)}; "
            f"valor=R$ {fmt_money(trade.value)}; corretora={trade.broker}{expiry}"
        )

    def _movement_detail(self, mov: Movement) -> str:
        return (
            f"Evento planilha movimentacao: data={mov.dt:%d/%m/%Y}; entrada/saida={mov.direction}; tipo={mov.kind}; "
            f"produto={mov.product}; codigo={mov.code}; qtd={fmt_decimal(mov.qty)}; "
            f"preco_unitario=R$ {fmt_money(mov.unit_price)}; valor=R$ {fmt_money(mov.value)}; corretora={mov.broker}"
        )

    def _cutoff_date(self, trades: list[Trade], movements: list[Movement]) -> date:
        dates = [trade.dt for trade in trades] + [mov.dt for mov in movements]
        max_year = max((item.year for item in dates), default=self.year)
        if self.year < max_year:
            return date(self.year, 12, 31)
        same_year_dates = [item for item in dates if item.year == self.year]
        return max(same_year_dates, default=date(self.year, 12, 31))

    def _split_daytrades(self, trades: list[Trade], monthly: list[MonthlyTax]) -> list[Trade]:
        remaining: list[Trade] = []
        by_key: dict[tuple[date, str, str, str], list[Trade]] = defaultdict(list)
        for trade in trades:
            by_key[(trade.dt, trade.code, trade.category, trade.broker)].append(trade)

        for (_dt, _code, category, _broker), group in by_key.items():
            buys = [t for t in group if t.side == "compra"]
            sells = [t for t in group if t.side == "venda"]
            buy_qty = sum((t.qty for t in buys), Decimal("0"))
            sell_qty = sum((t.qty for t in sells), Decimal("0"))
            dt_qty = min(buy_qty, sell_qty)
            if dt_qty <= 0:
                remaining.extend(group)
                continue

            buy_avg = sum((t.value for t in buys), Decimal("0")) / buy_qty if buy_qty else Decimal("0")
            sell_avg = sum((t.value for t in sells), Decimal("0")) / sell_qty if sell_qty else Decimal("0")
            result = q2((sell_avg - buy_avg) * dt_qty)
            month = monthly[group[0].dt.month - 1]
            if category == "futuro":
                month.future_result += result
                if self._future_kind(group[0].code) == "dolar":
                    month.future_dollar_daytrade += result
                else:
                    month.future_index_daytrade += result
            else:
                month.daytrade_result += result
            month.daytrade_sales += q2(sell_avg * dt_qty)

            for item in buys:
                rest = item.qty - (dt_qty * item.qty / buy_qty if buy_qty else Decimal("0"))
                if rest > 0:
                    remaining.append(Trade(item.dt, item.side, item.market, item.broker, item.code, q2(rest), item.price, q2(item.price * rest), item.category, item.expiry))
            for item in sells:
                rest = item.qty - (dt_qty * item.qty / sell_qty if sell_qty else Decimal("0"))
                if rest > 0:
                    remaining.append(Trade(item.dt, item.side, item.market, item.broker, item.code, q2(rest), item.price, q2(item.price * rest), item.category, item.expiry))
        return sorted(remaining, key=lambda t: (t.dt, t.code, t.side))

    def _process_normal_trades(
        self,
        trades: list[Trade],
        positions: dict[str, Position],
        monthly: list[MonthlyTax],
        warnings: list[str],
        events: list[dict[str, Any]],
    ) -> dict[int, dict[str, Position]]:
        monthly_positions: dict[int, dict[str, Position]] = {}
        short_positions: dict[str, dict[str, Any]] = {}
        trades_by_month: dict[int, list[Trade]] = defaultdict(list)
        for trade in trades:
            trades_by_month[trade.dt.month].append(trade)
        for month_number in range(1, 13):
            for trade in sorted(trades_by_month.get(month_number, []), key=lambda t: (t.dt, t.code, t.side)):
                pos = positions.setdefault(trade.code, Position(code=trade.code, category=trade.category, broker=trade.broker))
                if not pos.category or pos.category == "normal":
                    pos.category = trade.category
                if trade.side == "compra":
                    remaining_qty = trade.qty
                    remaining_value = trade.value
                    short = short_positions.get(trade.code)
                    if short and short["qty"] > 0:
                        cover_qty = min(remaining_qty, short["qty"])
                        cover_cost = q2(trade.value * cover_qty / trade.qty) if trade.qty else Decimal("0")
                        cover_credit = q2(short["credit"] * cover_qty / short["qty"]) if short["qty"] else Decimal("0")
                        result = q2(cover_credit - cover_cost)
                        month = monthly[trade.dt.month - 1]
                        self._add_trade_result(month, trade.category, result, code=trade.code)
                        short["qty"] = q2(short["qty"] - cover_qty)
                        short["credit"] = q2(short["credit"] - cover_credit)
                        remaining_qty = q2(remaining_qty - cover_qty)
                        remaining_value = q2(remaining_value - cover_cost)
                        if short["qty"] <= 0:
                            del short_positions[trade.code]
                    if remaining_qty > 0:
                        pos.qty += remaining_qty
                        pos.cost = q2(pos.cost + remaining_value)
                    continue

                month = monthly[trade.dt.month - 1]
                sell_remaining = trade.qty
                sell_remaining_value = trade.value
                if pos.qty > 0:
                    covered_qty = min(pos.qty, sell_remaining)
                    covered_value = q2(trade.value * covered_qty / trade.qty) if trade.qty else Decimal("0")
                    avg = pos.avg_price
                    cost = q2(avg * covered_qty)
                    result = q2(covered_value - cost)
                    self._add_trade_result(month, trade.category, result, covered_value, trade.code)
                    pos.qty = q2(pos.qty - covered_qty)
                    pos.cost = q2(max(Decimal("0"), pos.cost - cost))
                    sell_remaining = q2(sell_remaining - covered_qty)
                    sell_remaining_value = q2(sell_remaining_value - covered_value)
                    if pos.qty <= 0:
                        pos.qty = Decimal("0")
                        pos.cost = Decimal("0")
                if sell_remaining > 0:
                    short = short_positions.setdefault(
                        trade.code,
                        {"qty": Decimal("0"), "credit": Decimal("0"), "category": trade.category, "broker": trade.broker, "opened": trade.dt, "trade": trade},
                    )
                    short["qty"] = q2(short["qty"] + sell_remaining)
                    short["credit"] = q2(short["credit"] + sell_remaining_value)
                    short["opened"] = min(short["opened"], trade.dt)
                    short["trade"] = trade
                    self._add_trade_sale(month, trade.category, sell_remaining_value)
            monthly_positions[month_number] = clone_positions(positions)
        for code, short in sorted(short_positions.items()):
            if short["qty"] <= 0:
                continue
            if self._registered_active_short(code, short["qty"], events):
                continue
            trade = short.get("trade")
            detail = self._trade_detail(trade) if isinstance(trade, Trade) else f"ativo={code}"
            warnings.append(
                f"Venda descoberta em aberto: ativo={code}; qtd_vendida_sem_recompra={fmt_decimal(short['qty'])}; "
                f"credito_liquido=R$ {fmt_money(short['credit'])}. Possivel causa: operacao vendida ainda nao recomprada ou falta de evento de aluguel/recompra. {detail}"
            )
        return monthly_positions

    def _registered_active_short(self, code: str, qty: Decimal, events: list[dict[str, Any]]) -> bool:
        code = normalize_ticker(code)
        registered = Decimal("0")
        for event in events:
            if normalize_ticker(event.get("ativo")) != code:
                continue
            event_type = normalize_header(event.get("tipo"))
            if "venda_descoberta_ativa" in event_type or ("vendida" in event_type and "ativa" in event_type):
                registered += money(event.get("quantidade"))
        return registered >= qty

    def _add_trade_sale(self, month: MonthlyTax, category: str, value: Decimal) -> None:
        if category == "fii":
            month.fii_sales += q2(value)
        elif category not in {"opcoes", "futuro"}:
            month.normal_sales += q2(value)

    def _add_trade_result(self, month: MonthlyTax, category: str, result: Decimal, sale_value: Decimal = Decimal("0"), code: str = "") -> None:
        if category == "fii":
            month.fii_result += q2(result)
            month.fii_sales += q2(sale_value)
        elif category == "opcoes":
            month.options_result += q2(result)
        elif category == "futuro":
            month.future_result += q2(result)
            if self._future_kind(code) == "dolar":
                month.future_dollar_common += q2(result)
            else:
                month.future_index_common += q2(result)
        else:
            month.normal_result += q2(result)
            month.normal_sales += q2(sale_value)

    def _future_kind(self, code: str) -> str:
        code = normalize_ticker(code)
        if code.startswith(("WDO", "DOL")):
            return "dolar"
        return "indice"

    def _apply_events_before_trades(self, positions: dict[str, Position], events: list[dict[str, Any]], warnings: list[str]) -> None:
        for event in events:
            code = normalize_ticker(event.get("ativo", ""))
            if not code:
                continue
            pos = positions.setdefault(code, Position(code=code, category=str(event.get("categoria") or "normal")))
            kind = normalize_header(event.get("tipo", ""))
            qty = money(event.get("quantidade"))
            value = money(event.get("valor"))
            if "desdobramento" in kind or "grupamento" in kind:
                factor = qty if qty > 0 else Decimal("1")
                pos.qty = q2(pos.qty * factor)
                warnings.append(f"Evento aplicado em {code}: fator {fmt_decimal(factor)}. Confira se o fator informado esta correto.")
            elif "bonificacao" in kind or "subscricao" in kind:
                pos.qty += qty
                pos.cost = q2(pos.cost + value)
            elif "transferencia" in kind or "migracao" in kind:
                pos.qty += qty
                pos.cost = q2(pos.cost + value)
                broker_to = str(event.get("corretora") or event.get("broker_to") or "")
                if broker_to:
                    pos.broker = broker_to
            elif "exercicio" in kind:
                pos.qty += qty
                pos.cost = q2(pos.cost + value)

    def _update_broker_from_trades(
        self,
        positions: dict[str, Position],
        year_trades: list[Trade],
    ) -> None:
        """Atualiza o broker de cada posicao para a corretora da ultima compra do ano.

        Isso garante que compras feitas diretamente em uma nova corretora (sem
        transferencia de custodia registrada) reflitam a corretora correta na carteira.
        As transferencias de custodia sao aplicadas depois e tem prioridade sobre isso.
        """
        last_buy: dict[str, Trade] = {}
        for trade in sorted(year_trades, key=lambda t: (t.dt, t.code)):
            if trade.side == "compra" and trade.broker:
                last_buy[trade.code] = trade
        for code, trade in last_buy.items():
            if code in positions and positions[code].qty > 0:
                positions[code].broker = trade.broker

    def _apply_custody_transfers(
        self,
        positions: dict[str, Position],
        transfers: list[dict[str, Any]],
        year: int,
        warnings: list[str],
    ) -> None:
        """Aplica transferencias de custodia (portabilidade B3) nas posicoes.

        Quando a quantidade transferida e menor que a posicao total, considera
        que o saldo remanescente foi adquirido diretamente na corretora destino
        (compras posteriores) — logo a posicao inteira e marcada como broker_to,
        pois o PM ja e calculado de forma consolidada pelo engine.
        """
        # agrupa por codigo para aplicar multiplas transferencias em ordem cronologica
        from collections import defaultdict as _dd
        by_code: dict[str, list[dict]] = _dd(list)
        for t in transfers:
            dt = parse_date(t.get("transfer_date"))
            if dt is None or dt.year > year:
                continue
            code = normalize_ticker(t.get("ticker", ""))
            if code:
                by_code[code].append(t)

        for code, tlist in by_code.items():
            # ordena por data para aplicar na sequencia correta
            tlist.sort(key=lambda x: parse_date(x.get("transfer_date")) or date.min)
            for t in tlist:
                dt = parse_date(t.get("transfer_date"))
                qty_transferred = money(t.get("quantity", "0"))
                broker_to = str(t.get("broker_to") or "")
                if not broker_to:
                    continue
                if code in positions:
                    pos = positions[code]
                    pos_qty = pos.qty
                    # migra o broker independente da quantidade —
                    # transferencias parciais sao esperadas quando ha compras
                    # posteriores na corretora destino que nao geram nova transferencia
                    old_broker = pos.broker or "?"
                    pos.broker = broker_to
                    if qty_transferred > 0 and qty_transferred < pos_qty:
                        warnings.append(
                            f"Transferencia parcial de {code} em {dt:%d/%m/%Y}: "
                            f"{fmt_decimal(qty_transferred)} de {fmt_decimal(pos_qty)} acoes transferidas de "
                            f"'{old_broker}' para '{broker_to}'. "
                            f"O restante ({fmt_decimal(pos_qty - qty_transferred)}) foi adquirido diretamente na corretora destino. "
                            f"Posicao consolidada em '{broker_to}'."
                        )
                else:
                    warnings.append(
                        f"Transferencia de custodia: ativo {code} em {dt:%d/%m/%Y} nao encontrado nas posicoes. "
                        f"Corretora destino: {broker_to}. Verifique se o ativo foi importado corretamente."
                    )

    def _apply_corporate_events(
        self,
        positions: dict[str, Position],
        corp_events: list[dict[str, Any]],
        year: int,
        warnings: list[str],
    ) -> None:
        """Aplica eventos corporativos: split, grupamento, bonificacao, fusao de corretora."""
        for ev in corp_events:
            dt = parse_date(ev.get("event_date"))
            if dt is None or dt.year > year:
                continue
            etype = normalize_header(ev.get("event_type", ""))
            code = normalize_ticker(ev.get("ticker", ""))
            code_new = normalize_ticker(ev.get("ticker_new", "")) or code
            factor = money(ev.get("factor", "1"))
            bonus_qty = money(ev.get("bonus_qty", "0"))
            bonus_cost = money(ev.get("bonus_cost", "0"))
            broker_to = str(ev.get("broker_to") or "")

            if "desdobramento" in etype or "split" in etype:
                if code in positions:
                    f = factor if factor > 0 else Decimal("1")
                    positions[code].qty = q2(positions[code].qty * f)
                    warnings.append(
                        f"Split/desdobramento em {code} em {dt:%d/%m/%Y}: fator {fmt_decimal(f)}. "
                        f"Nova quantidade: {fmt_decimal(positions[code].qty)}."
                    )

            elif "grupamento" in etype or "inplit" in etype:
                if code in positions:
                    f = factor if factor > 0 else Decimal("1")
                    positions[code].qty = q2(positions[code].qty / f)
                    warnings.append(
                        f"Grupamento em {code} em {dt:%d/%m/%Y}: fator 1/{fmt_decimal(f)}. "
                        f"Nova quantidade: {fmt_decimal(positions[code].qty)}."
                    )

            elif "bonificacao" in etype or "subscricao" in etype:
                pos = positions.setdefault(code_new, Position(code=code_new, category="normal"))
                pos.qty += bonus_qty
                pos.cost = q2(pos.cost + bonus_cost)
                warnings.append(
                    f"Bonificacao em {code_new} em {dt:%d/%m/%Y}: +{fmt_decimal(bonus_qty)} cotas/acoes, "
                    f"custo adicional R$ {fmt_money(bonus_cost)}."
                )

            elif "fusao" in etype or "incorporacao" in etype or "migracao_corretora" in etype or "aquisicao" in etype:
                # Fusão de corretora: atualiza broker de todas as posicoes da corretora origem
                broker_from = str(ev.get("broker_from") or "")
                if broker_from and broker_to:
                    # palavras significativas do nome da corretora origem (ignora siglas genéricas)
                    _skip = {"s", "a", "sa", "ltda", "s_a", "de", "do", "da", "e", "em", "banco",
                             "corretora", "cctvm", "ctvm", "dtvm", "cv", "cc", "investimentos"}
                    _words_from = {w for w in normalize_header(broker_from).split("_") if w and w not in _skip and len(w) > 2}
                    def _broker_matches(pos_broker: str) -> bool:
                        if not pos_broker:
                            return False
                        norm_pos = normalize_header(pos_broker)
                        # correspondência exata
                        if norm_pos == normalize_header(broker_from):
                            return True
                        # correspondência por palavras-chave significativas
                        _words_pos = set(norm_pos.split("_"))
                        return bool(_words_from & _words_pos)
                    moved: list[str] = []
                    for pos in positions.values():
                        if _broker_matches(pos.broker or ""):
                            pos.broker = broker_to
                            moved.append(pos.code)
                    if moved:
                        warnings.append(
                            f"Fusao/aquisicao em {dt:%d/%m/%Y}: {len(moved)} ativo(s) migrados de "
                            f"'{broker_from}' para '{broker_to}': {', '.join(sorted(moved)[:10])}{'...' if len(moved)>10 else ''}."
                        )
                elif code and broker_to:
                    # Migração de ativo específico
                    if code in positions:
                        positions[code].broker = broker_to

            elif "renomeacao" in etype or "conversao" in etype:
                if code in positions and code_new and code_new != code:
                    pos_old = positions.pop(code)
                    pos_new = positions.setdefault(code_new, Position(code=code_new, category=pos_old.category, broker=pos_old.broker))
                    pos_new.qty += pos_old.qty
                    pos_new.cost = q2(pos_new.cost + pos_old.cost)
                    warnings.append(
                        f"Renomeacao/conversao: {code} -> {code_new} em {dt:%d/%m/%Y}. "
                        f"Quantidade transferida: {fmt_decimal(pos_old.qty)}."
                    )

    def _apply_income_movements(self, movements: list[Movement], monthly: list[MonthlyTax]) -> None:
        for mov in movements:
            if mov.dt.year != self.year:
                continue
            kind = normalize_header(mov.kind)
            if "rendimento" in kind and mov.category == "fii":
                pass

    def _apply_monthly_tax(
        self,
        monthly: list[MonthlyTax],
        initial_losses: dict[str, Decimal],
        dated_losses: dict[str, Decimal],
        dated_loss_start: date | None,
    ) -> None:
        normal_loss = abs(money(initial_losses.get("normal")))
        day_loss = abs(money(initial_losses.get("daytrade")))
        fii_loss = abs(money(initial_losses.get("fii")))
        for item in monthly:
            if dated_loss_start and dated_loss_start.year == self.year and item.month == dated_loss_start.month:
                normal_loss = abs(money(dated_losses.get("normal")))
                day_loss = abs(money(dated_losses.get("daytrade")))
                fii_loss = abs(money(dated_losses.get("fii")))
            item.normal_loss_before = q2(normal_loss)
            item.daytrade_loss_before = q2(day_loss)
            item.fii_loss_before = q2(fii_loss)
            item.options_loss_before = Decimal("0")
            item.future_loss_before = Decimal("0")

            taxable_normal = item.normal_result
            stock_exempt = item.normal_sales <= Decimal("20000") and item.normal_result > 0
            if stock_exempt:
                item.exempt_stock_gain = q2(item.normal_result)
                taxable_normal = Decimal("0")
            future_daytrade_result = q2(item.future_dollar_daytrade + item.future_index_daytrade)
            future_common_result = q2(item.future_result - future_daytrade_result)
            common_result = q2(taxable_normal + item.options_result + future_common_result)
            item.normal_base = q2(max(Decimal("0"), common_result - normal_loss))
            normal_loss = q2(max(Decimal("0"), normal_loss - common_result)) if common_result >= 0 else q2(normal_loss + abs(common_result))

            taxable_daytrade = q2(item.daytrade_result + future_daytrade_result)
            item.daytrade_base = q2(max(Decimal("0"), taxable_daytrade - day_loss))
            day_loss = q2(max(Decimal("0"), day_loss - taxable_daytrade)) if taxable_daytrade >= 0 else q2(day_loss + abs(taxable_daytrade))

            item.fii_base = q2(max(Decimal("0"), item.fii_result - fii_loss))
            fii_loss = q2(max(Decimal("0"), fii_loss - item.fii_result)) if item.fii_result >= 0 else q2(fii_loss + abs(item.fii_result))

            item.options_base = Decimal("0")
            item.future_base = Decimal("0")

            item.normal_loss_after = q2(normal_loss)
            item.daytrade_loss_after = q2(day_loss)
            item.fii_loss_after = q2(fii_loss)
            item.options_loss_after = Decimal("0")
            item.future_loss_after = Decimal("0")
            item.tax_due = q2(item.normal_base * Decimal("0.15") + item.daytrade_base * Decimal("0.20") + item.fii_base * Decimal("0.20"))
            item.tax_payable = item.tax_due

    def _apply_withholding(self, monthly: list[MonthlyTax], brokerage_taxes: list[BrokerageNoteTax]) -> None:
        by_month: dict[int, dict[str, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
        for note in brokerage_taxes:
            if note.dt.year != self.year:
                continue
            by_month[note.dt.month]["common"] += money(note.irrf_common)
            by_month[note.dt.month]["daytrade"] += money(note.irrf_daytrade)
            by_month[note.dt.month]["fii"] += money(note.irrf_fii)
        common_credit = Decimal("0")
        daytrade_credit = Decimal("0")
        fii_credit = Decimal("0")
        for item in monthly:
            item.irrf_common_before = q2(common_credit)
            item.irrf_daytrade_before = q2(daytrade_credit)
            item.irrf_fii_before = q2(fii_credit)
            item.irrf_common_month = q2(by_month[item.month]["common"])
            item.irrf_daytrade_month = q2(by_month[item.month]["daytrade"])
            item.irrf_fii_month = q2(by_month[item.month]["fii"])

            common_tax = q2(item.normal_base * Decimal("0.15"))
            daytrade_tax = q2(item.daytrade_base * Decimal("0.20"))
            fii_tax = q2(item.fii_base * Decimal("0.20"))

            common_available = q2(common_credit + item.irrf_common_month)
            daytrade_available = q2(daytrade_credit + item.irrf_daytrade_month)
            fii_available = q2(fii_credit + item.irrf_fii_month)

            common_payable = q2(max(Decimal("0"), common_tax - common_available))
            daytrade_payable = q2(max(Decimal("0"), daytrade_tax - daytrade_available))
            fii_payable = q2(max(Decimal("0"), fii_tax - fii_available))

            common_credit = q2(max(Decimal("0"), common_available - common_tax))
            daytrade_credit = q2(max(Decimal("0"), daytrade_available - daytrade_tax))
            fii_credit = q2(max(Decimal("0"), fii_available - fii_tax))
            item.irrf_common_after = common_credit
            item.irrf_daytrade_after = daytrade_credit
            item.irrf_fii_after = fii_credit
            item.tax_payable = q2(common_payable + daytrade_payable + fii_payable)

    def _process_option_trades(
        self,
        trades: list[Trade],
        monthly: list[MonthlyTax],
        warnings: list[str],
        events: list[dict[str, Any]],
        cutoff: date,
    ) -> list[dict[str, Any]]:
        resolved = self._resolved_option_events(events)
        option_positions: dict[str, dict[str, Any]] = {}
        option_trades = [
            trade
            for trade in trades
            if trade.category == "opcoes" and trade.dt <= cutoff and "exercicio" not in normalize_header(trade.market)
        ]
        for trade in sorted(option_trades, key=lambda t: (t.dt, t.code, t.side)):
            pos = option_positions.setdefault(
                trade.code,
                {
                    "codigo": trade.code,
                    "vencimento": trade.expiry,
                    "long_qty": Decimal("0"),
                    "long_cost": Decimal("0"),
                    "short_qty": Decimal("0"),
                    "short_credit": Decimal("0"),
                    "ultima_data": trade.dt,
                    "corretora": trade.broker,
                },
            )
            if trade.expiry and (pos["vencimento"] is None or trade.expiry > pos["vencimento"]):
                pos["vencimento"] = trade.expiry
            pos["ultima_data"] = max(pos["ultima_data"], trade.dt)
            if trade.side == "compra":
                qty_to_close = min(trade.qty, pos["short_qty"])
                if qty_to_close > 0:
                    avg_credit = pos["short_credit"] / pos["short_qty"] if pos["short_qty"] else Decimal("0")
                    buy_cost = q2(trade.price * qty_to_close)
                    result = q2(avg_credit * qty_to_close - buy_cost)
                    if trade.dt.year == self.year:
                        monthly[trade.dt.month - 1].options_result += result
                    pos["short_qty"] -= qty_to_close
                    pos["short_credit"] = q2(max(Decimal("0"), pos["short_credit"] - avg_credit * qty_to_close))
                remaining = trade.qty - qty_to_close
                if remaining > 0:
                    pos["long_qty"] += remaining
                    pos["long_cost"] = q2(pos["long_cost"] + trade.price * remaining)
            else:
                qty_to_close = min(trade.qty, pos["long_qty"])
                if qty_to_close > 0:
                    avg_cost = pos["long_cost"] / pos["long_qty"] if pos["long_qty"] else Decimal("0")
                    sell_value = q2(trade.price * qty_to_close)
                    result = q2(sell_value - avg_cost * qty_to_close)
                    if trade.dt.year == self.year:
                        monthly[trade.dt.month - 1].options_result += result
                    pos["long_qty"] -= qty_to_close
                    pos["long_cost"] = q2(max(Decimal("0"), pos["long_cost"] - avg_cost * qty_to_close))
                remaining = trade.qty - qty_to_close
                if remaining > 0:
                    pos["short_qty"] += remaining
                    pos["short_credit"] = q2(pos["short_credit"] + trade.price * remaining)

        pending: list[dict[str, Any]] = []
        for code, pos in sorted(option_positions.items()):
            expiry = pos["vencimento"]
            event_type = resolved.get(code)
            if event_type and any(marker in event_type for marker in ["liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao"]):
                continue
            if event_type and ("recompra" in event_type or "recomprada" in event_type):
                event = self._resolved_option_event_rows(events).get(code)
                if event:
                    event_date = parse_date(event.get("data")) or pos["vencimento"] or cutoff
                    if event_date.year == self.year:
                        repurchase_value = q2(event.get("valor"))
                        if pos["short_qty"] > 0:
                            monthly[event_date.month - 1].options_result += q2(pos["short_credit"] - repurchase_value)
                        elif pos["long_qty"] > 0:
                            monthly[event_date.month - 1].options_result += q2(repurchase_value - pos["long_cost"])
                continue
            if event_type and self._is_open_option_event(event_type):
                self._append_pending_option(pending, pos, "Registrada como vendida/comprada e ainda ativa")
                continue
            if event_type and ("exercicio" in event_type or "virou_po" in event_type or event_type.endswith("po")):
                continue
            if isinstance(expiry, date) and expiry <= cutoff:
                if pos["long_qty"] > 0 and expiry.year == self.year:
                    monthly[expiry.month - 1].options_result -= q2(pos["long_cost"])
                    self._append_pending_option(pending, pos, "SUPOSICAO: opcao comprada vencida foi tratada como virou po; confirme ou registre recompra/exercicio")
                    warnings.append(
                        f"Opcao comprada vencida sem venda/exercicio identificado: codigo={code}; vencimento={expiry:%d/%m/%Y}; "
                        f"qtd_aberta={fmt_decimal(pos['long_qty'])}; premio_pago=R$ {fmt_money(pos['long_cost'])}; "
                        f"efeito_calculo=perda do premio em Opcoes. Evento origem: ultima_operacao={pos['ultima_data']:%d/%m/%Y}; corretora={pos['corretora']}."
                    )
                if pos["short_qty"] > 0 and expiry.year == self.year:
                    monthly[expiry.month - 1].options_result += q2(pos["short_credit"])
                    self._append_pending_option(pending, pos, "SUPOSICAO: opcao vendida vencida foi tratada como sem recompra/exercicio; confirme ou registre recompra/exercicio")
                    warnings.append(
                        f"Opcao vendida vencida sem recompra/exercicio identificado: codigo={code}; vencimento={expiry:%d/%m/%Y}; "
                        f"qtd_aberta={fmt_decimal(pos['short_qty'])}; premio_recebido=R$ {fmt_money(pos['short_credit'])}; "
                        f"efeito_calculo=premio considerado como ganho em Opcoes. Evento origem: ultima_operacao={pos['ultima_data']:%d/%m/%Y}; corretora={pos['corretora']}."
                    )
                continue
            if pos["long_qty"] > 0 or pos["short_qty"] > 0:
                self._append_pending_option(pending, pos, "Ainda ativa no corte do calculo; confirme se foi recomprada, exercida ou virou po")
        return pending

    def _resolved_option_events(self, events: list[dict[str, Any]]) -> dict[str, str]:
        resolved: dict[str, str] = {}
        for event in events:
            code = normalize_ticker(event.get("ativo"))
            kind = normalize_header(event.get("tipo"))
            if not code:
                continue
            if any(marker in kind for marker in ["exercicio", "virou_po", "po", "ainda_ativa", "vendida_ativa", "ainda_em_aberto", "em_aberto", "recompra", "recomprada", "liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao"]):
                resolved[code] = kind
        return resolved

    def _is_open_option_event(self, event_type: str) -> bool:
        normalized = normalize_header(event_type)
        return "ativa" in normalized or "em_aberto" in normalized or ("aberto" in normalized and "opcao" in normalized)

    def _resolved_option_event_rows(self, events: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
        rows: dict[str, dict[str, Any]] = {}
        for event in events:
            code = normalize_ticker(event.get("ativo"))
            kind = normalize_header(event.get("tipo"))
            if code and any(marker in kind for marker in ["exercicio", "virou_po", "po", "ainda_ativa", "vendida_ativa", "ainda_em_aberto", "em_aberto", "recompra", "recomprada", "liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao"]):
                rows[code] = event
        return rows

    def _append_pending_option(self, pending: list[dict[str, Any]], pos: dict[str, Any], situation: str) -> None:
        net_short = q2(pos["short_qty"])
        net_long = q2(pos["long_qty"])
        qty = net_short if net_short > 0 else net_long
        if qty <= 0:
            return
        premium = q2(pos["short_credit"] if net_short > 0 else -pos["long_cost"])
        pending.append(
            {
                "codigo": pos["codigo"],
                "vencimento": pos["vencimento"],
                "quantidade_vendida": net_short,
                "quantidade_comprada": net_long,
                "quantidade_aberta": qty,
                "premio_liquido": premium,
                "ultima_data": pos["ultima_data"],
                "corretora": pos["corretora"],
                "situacao": situation,
            }
        )

    def _process_option_exercises(
        self,
        trades: list[Trade],
        positions: dict[str, Position],
        monthly_positions: dict[int, dict[str, Position]],
        monthly: list[MonthlyTax],
        warnings: list[str],
        cutoff: date,
        events: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        settlement_events = self._option_settlement_without_stock_events(events)
        confirmed_exercises = self._confirmed_option_exercise_events(events)
        pending_exercises: list[dict[str, Any]] = []
        exercises = [
            trade
            for trade in trades
            if trade.dt.year == self.year and trade.dt <= cutoff and "exercicio" in normalize_header(trade.market)
        ]
        for trade in sorted(exercises, key=lambda item: item.dt):
            exercise_code = normalize_option_exercise_code(trade.code)
            if exercise_code in settlement_events or trade.code in settlement_events:
                effect = q2(trade.value if trade.side == "venda" else -trade.value)
                monthly[trade.dt.month - 1].options_result += effect
                warnings.append(
                    f"Exercicio tratado como liquidacao financeira de trava sem movimentar acao: opcao={trade.code}; "
                    f"efeito_em_opcoes=R$ {fmt_money(effect)}. {self._trade_detail(trade)}"
                )
                continue
            if exercise_code not in confirmed_exercises and trade.code not in confirmed_exercises:
                pending_exercises.append(self._pending_exercise_row(trade))
                warnings.append(
                    f"Exercicio de opcao pendente de confirmacao: opcao={trade.code}; "
                    f"confirme se houve compra/venda da acao-base ou se foi liquidacao financeira de trava. {self._trade_detail(trade)}"
                )
                continue
            underlying = self._infer_underlying_from_option(trade.code, positions, trade.qty)
            if not underlying:
                warnings.append(
                    f"Exercicio de opcao sem ativo-objeto identificado: opcao={trade.code}. "
                    f"Revise o codigo e a posicao do ativo-base. {self._trade_detail(trade)}"
                )
                continue
            pos = positions.setdefault(underlying, Position(code=underlying, category=classify_asset(underlying), broker=trade.broker))
            avg = pos.avg_price
            cost = q2(avg * trade.qty)
            result = q2(trade.value - cost)
            month = monthly[trade.dt.month - 1]
            if trade.side == "venda":
                if pos.qty < trade.qty:
                    warnings.append(
                        f"Exercicio de call vendida excede posicao do ativo-base: opcao={trade.code}; ativo_base={underlying}; "
                        f"posicao_antes={fmt_decimal(pos.qty)}; qtd_exercida={fmt_decimal(trade.qty)}; "
                        f"falta={fmt_decimal(trade.qty - pos.qty)}. {self._trade_detail(trade)}"
                    )
                pos.qty -= trade.qty
                pos.cost = q2(max(Decimal("0"), pos.cost - cost))
                if pos.qty <= 0:
                    pos.qty = Decimal("0")
                    pos.cost = Decimal("0")
                self._apply_exercise_to_monthly_snapshots(monthly_positions, trade.dt.month, underlying, trade.qty, cost, "venda", pos)
                if pos.category == "fii":
                    month.fii_result += result
                    month.fii_sales += trade.value
                else:
                    month.normal_result += result
                    month.normal_sales += trade.value
                warnings.append(
                    f"Exercicio de call vendida contabilizado: opcao={trade.code}; ativo_base={underlying}; "
                    f"venda_automatica_qtd={fmt_decimal(trade.qty)}; strike/preco=R$ {fmt_money(trade.price)}; "
                    f"valor_venda=R$ {fmt_money(trade.value)}; custo_medio_base=R$ {fmt_money(avg)}; "
                    f"resultado_ativo=R$ {fmt_money(result)}. {self._trade_detail(trade)}"
                )
            elif trade.side == "compra":
                pos.qty += trade.qty
                pos.cost = q2(pos.cost + trade.value)
                self._apply_exercise_to_monthly_snapshots(monthly_positions, trade.dt.month, underlying, trade.qty, trade.value, "compra", pos)
                warnings.append(
                    f"Exercicio de call comprada contabilizado: opcao={trade.code}; ativo_base={underlying}; "
                    f"compra_automatica_qtd={fmt_decimal(trade.qty)}; strike/preco=R$ {fmt_money(trade.price)}; "
                    f"valor_compra=R$ {fmt_money(trade.value)}. {self._trade_detail(trade)}"
                )
        return pending_exercises

    def _pending_exercise_row(self, trade: Trade) -> dict[str, Any]:
        return {
            "codigo": normalize_option_exercise_code(trade.code),
            "vencimento": trade.expiry or trade.dt,
            "quantidade_vendida": trade.qty if trade.side == "venda" else Decimal("0"),
            "quantidade_comprada": trade.qty if trade.side == "compra" else Decimal("0"),
            "quantidade_aberta": trade.qty,
            "premio_liquido": q2(trade.value if trade.side == "venda" else -trade.value),
            "ultima_data": trade.dt,
            "corretora": trade.broker,
            "situacao": "Exercicio B3 pendente: confirme se movimentou a acao ou se foi liquidacao financeira de trava",
        }

    def _confirmed_option_exercise_events(self, events: list[dict[str, Any]]) -> set[str]:
        codes: set[str] = set()
        for event in events:
            kind = normalize_header(event.get("tipo"))
            if "exercicio" not in kind or any(marker in kind for marker in ["liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao"]):
                continue
            code = normalize_ticker(event.get("ativo"))
            if not code:
                continue
            codes.add(code)
            codes.add(normalize_option_exercise_code(code))
        return codes

    def _option_settlement_without_stock_events(self, events: list[dict[str, Any]]) -> set[str]:
        codes: set[str] = set()
        for event in events:
            kind = normalize_header(event.get("tipo"))
            if not any(marker in kind for marker in ["liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao"]):
                continue
            code = normalize_ticker(event.get("ativo"))
            if not code:
                continue
            codes.add(code)
            codes.add(normalize_option_exercise_code(code))
        return codes

    def _apply_exercise_to_monthly_snapshots(
        self,
        monthly_positions: dict[int, dict[str, Position]],
        start_month: int,
        code: str,
        qty: Decimal,
        cost: Decimal,
        side: str,
        reference: Position,
    ) -> None:
        for month in range(start_month, 13):
            snapshot = monthly_positions.setdefault(month, {})
            pos = snapshot.setdefault(code, Position(code=code, category=reference.category, broker=reference.broker))
            if side == "venda":
                pos.qty = max(Decimal("0"), pos.qty - qty)
                pos.cost = q2(max(Decimal("0"), pos.cost - cost))
                if pos.qty <= 0:
                    pos.qty = Decimal("0")
                    pos.cost = Decimal("0")
            else:
                pos.qty += qty
                pos.cost = q2(pos.cost + cost)

    def _infer_underlying_from_option(self, option_code: str, positions: dict[str, Position], qty: Decimal) -> str:
        root = option_root(option_code)
        candidates = [
            pos
            for code, pos in positions.items()
            if code.startswith(root) and pos.category not in {"opcoes", "futuro"} and re.match(r"^[A-Z]{4}\d", code)
        ]
        if candidates:
            candidates.sort(key=lambda pos: (pos.qty >= qty, pos.qty, pos.cost), reverse=True)
            return candidates[0].code
        common = {
            "ABEV": "ABEV3",
            "ASAI": "ASAI3",
            "BBAS": "BBAS3",
            "BBSE": "BBSE3",
            "BBDC": "BBDC4",
            "BRAP": "BRAP4",
            "CMIG": "CMIG4",
            "EGIE": "EGIE3",
            "GGBR": "GGBR4",
            "ITSA": "ITSA4",
            "LREN": "LREN3",
            "PETR": "PETR4",
            "RENT": "RENT3",
            "SUZB": "SUZB3",
            "USIM": "USIM5",
            "VALE": "VALE3",
            "WEGE": "WEGE3",
        }
        return common.get(root, f"{root}3")


    def _monthly_loans(self, movements: list[Movement]) -> dict[int, dict[str, Decimal]]:
        loans: dict[str, Decimal] = defaultdict(Decimal)
        snapshots: dict[int, dict[str, Decimal]] = {}
        year_movements = sorted((mov for mov in movements if mov.dt.year == self.year), key=lambda m: m.dt)
        by_month: dict[int, list[Movement]] = defaultdict(list)
        for mov in year_movements:
            by_month[mov.dt.month].append(mov)
        for month in range(1, 13):
            for mov in by_month.get(month, []):
                if "emprestimo" not in normalize_header(mov.kind) or mov.qty <= 0:
                    continue
                if normalize_header(mov.direction) == "debito":
                    loans[mov.code] = max(Decimal("0"), loans[mov.code] - mov.qty)
                else:
                    # A B3 costuma repetir/renovar linhas de emprestimo. Somar
                    # todas infla a quantidade; manter o maior saldo observado
                    # aproxima melhor as acoes efetivamente emprestadas no mes.
                    loans[mov.code] = max(loans[mov.code], mov.qty)
            snapshots[month] = {code: q2(qty) for code, qty in loans.items() if qty > 0}
        return snapshots

    def _pending_options(self, trades: list[Trade], events: list[dict[str, Any]]) -> list[dict[str, Any]]:
        option_trades = [t for t in trades if t.category == "opcoes"]
        positions: dict[str, dict[str, Any]] = {}
        exercised_codes = {t.code for t in trades if "exercicio" in normalize_header(t.market)}
        exercised_base_codes = {normalize_option_exercise_code(code) for code in exercised_codes}
        resolved_codes = {
            normalize_ticker(event.get("ativo"))
            for event in events
            if any(marker in normalize_header(event.get("tipo")) for marker in ["exercicio", "virou_po", "po", "ainda_ativa", "vendida_ativa", "ainda_em_aberto", "em_aberto", "liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao"])
        }
        resolved_base_codes = {normalize_option_exercise_code(code) for code in resolved_codes}
        for trade in option_trades:
            if "exercicio" in normalize_header(trade.market):
                continue
            row = positions.setdefault(
                trade.code,
                {
                    "codigo": trade.code,
                    "vencimento": trade.expiry,
                    "quantidade_vendida": Decimal("0"),
                    "quantidade_comprada": Decimal("0"),
                    "premio_vendido": Decimal("0"),
                    "premio_comprado": Decimal("0"),
                    "ultima_data": trade.dt,
                    "corretora": trade.broker,
                    "situacao": "Pendente: informar se foi exercida ou virou po",
                },
            )
            if trade.expiry and (row["vencimento"] is None or trade.expiry > row["vencimento"]):
                row["vencimento"] = trade.expiry
            row["ultima_data"] = max(row["ultima_data"], trade.dt)
            if trade.side == "venda":
                row["quantidade_vendida"] += trade.qty
                row["premio_vendido"] += trade.value
            else:
                row["quantidade_comprada"] += trade.qty
                row["premio_comprado"] += trade.value
        pending = []
        for row in positions.values():
            net_short = row["quantidade_vendida"] - row["quantidade_comprada"]
            base_code = normalize_option_exercise_code(row["codigo"])
            if net_short <= 0 or row["codigo"] in resolved_codes or base_code in resolved_base_codes:
                continue
            row["quantidade_aberta"] = q2(net_short)
            row["premio_liquido"] = q2(row["premio_vendido"] - row["premio_comprado"])
            if row["codigo"] in exercised_codes or base_code in exercised_base_codes:
                row["situacao"] = "Exercicio B3 encontrado; confirme se movimentou a acao ou se foi liquidacao financeira de trava"
            pending.append(row)
        return sorted(pending, key=lambda item: (item["vencimento"] or date.max, item["codigo"]))

    def _annual_tables(
        self,
        movements: list[Movement],
        positions: dict[str, Position],
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
        exempt: dict[tuple[str, str], Decimal] = defaultdict(Decimal)
        taxable: dict[tuple[str, str], Decimal] = defaultdict(Decimal)
        debts: dict[str, Decimal] = defaultdict(Decimal)

        for mov in movements:
            if mov.dt.year != self.year:
                continue
            kind = normalize_header(mov.kind)
            if "rendimento" in kind and mov.category == "fii":
                exempt[("26 - Outros", self._income_description("Rendimentos de FII/Fiagro", mov.code))] += mov.value
            elif "dividendo" in kind:
                exempt[("09 - Lucros e dividendos recebidos", self._income_description("Dividendos", mov.code))] += mov.value
            elif "juros_sobre_capital" in kind or "jcp" in kind:
                taxable[("10 - Juros sobre capital proprio", self._income_description("JCP", mov.code))] += mov.value
            elif "emprestimo" in kind and mov.value > 0:
                taxable[("06 - Rendimentos de aplicacoes financeiras", self._income_description("Rendimento de emprestimo de ativos", mov.code))] += mov.value
            if "emprestimo" in kind and mov.qty > 0:
                debts[mov.code] += mov.value if mov.value > 0 else Decimal("0")

        exempt_rows = [{"codigo": k[0], "descricao": k[1], "valor": q2(v)} for k, v in sorted(exempt.items())]
        taxable_rows = [{"codigo": k[0], "descricao": k[1], "valor": q2(v)} for k, v in sorted(taxable.items())]
        debt_rows = [
            {
                "codigo": "16 - Outras dividas e onus reais",
                "descricao": f"Emprestimo de Acoes - Ativo: {code}",
                "situacao_anterior": Decimal("0"),
                "situacao_atual": q2(value),
                "valor_pago": Decimal("0"),
            }
            for code, value in sorted(debts.items())
            if value > 0
        ]
        return exempt_rows, taxable_rows, debt_rows

    def _income_description(self, label: str, code: str) -> str:
        code = normalize_ticker(code)
        cnpj = self.asset_cnpjs.get(code)
        return f"{label} - {code} - CNPJ {cnpj}" if cnpj else f"{label} - {code}"


def read_b3_negotiation(path: Path) -> list[Trade]:
    if pd is None:
        raise RuntimeError("A biblioteca pandas nao esta instalada.")
    df = pd.read_excel(path, dtype=object)
    df.columns = [normalize_header(c) for c in df.columns]
    trades: list[Trade] = []
    for _, row in df.iterrows():
        dt = parse_date(row.get("data_do_negocio"))
        side_text = normalize_header(row.get("tipo_de_movimentacao"))
        if dt is None or side_text not in {"compra", "venda"}:
            continue
        code = normalize_ticker(row.get("codigo_de_negociacao"))
        market = str(row.get("mercado") or "")
        value = q2(row.get("valor"))
        qty = money(row.get("quantidade"))
        price = money(row.get("preco"))
        if not code or qty <= 0:
            continue
        trades.append(
            Trade(
                dt=dt,
                side=side_text,
                market=market,
                broker=str(row.get("instituicao") or ""),
                code=code,
                qty=qty,
                price=price,
                value=value if value else q2(qty * price),
                category=classify_asset(code, market),
                expiry=parse_date(row.get("prazo_vencimento")),
            )
        )
    return trades


def read_b3_movements(path: Path) -> list[Movement]:
    if pd is None:
        raise RuntimeError("A biblioteca pandas nao esta instalada.")
    df = pd.read_excel(path, dtype=object)
    df.columns = [normalize_header(c) for c in df.columns]
    movements: list[Movement] = []
    for _, row in df.iterrows():
        dt = parse_date(row.get("data"))
        if dt is None:
            continue
        product = str(row.get("produto") or "")
        code = asset_from_product(product)
        kind = str(row.get("movimentacao") or "")
        value = q2(row.get("valor_da_operacao"))
        qty = money(row.get("quantidade"))
        unit = money(row.get("preco_unitario"))
        movements.append(
            Movement(
                dt=dt,
                direction=str(row.get("entrada_saida") or ""),
                kind=kind,
                product=product,
                broker=str(row.get("instituicao") or ""),
                code=code,
                qty=qty,
                unit_price=unit,
                value=value,
                category=classify_asset(code, product=product),
            )
        )
    return movements


class EditableTable(ttk.Frame):
    def __init__(self, master: tk.Misc, columns: list[tuple[str, str, int]], height: int = 8) -> None:
        super().__init__(master)
        self.columns = columns
        self.tree = ttk.Treeview(self, columns=[c[0] for c in columns], show="headings", height=height)
        vsb = ttk.Scrollbar(self, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(self, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        for key, label, width in columns:
            self.tree.heading(key, text=label)
            self.tree.column(key, width=width, anchor="w")
        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        bar = ttk.Frame(self)
        bar.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(6, 0))
        ttk.Button(bar, text="Adicionar", command=self.add_row).pack(side="left")
        ttk.Button(bar, text="Editar", command=self.edit_selected).pack(side="left", padx=6)
        ttk.Button(bar, text="Remover", command=self.remove_selected).pack(side="left")
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)

    def add_row(self, values: Iterable[Any] | None = None) -> None:
        if values is None:
            values = ["" for _ in self.columns]
        self.tree.insert("", "end", values=list(values))

    def rows(self) -> list[dict[str, str]]:
        result = []
        for item in self.tree.get_children():
            values = self.tree.item(item, "values")
            result.append({self.columns[i][0]: str(values[i]) if i < len(values) else "" for i in range(len(self.columns))})
        return result

    def edit_selected(self) -> None:
        selected = self.tree.selection()
        if not selected:
            return
        item = selected[0]
        values = list(self.tree.item(item, "values"))
        dialog = RowEditor(self, self.columns, values)
        self.wait_window(dialog)
        if dialog.result is not None:
            self.tree.item(item, values=dialog.result)

    def remove_selected(self) -> None:
        for item in self.tree.selection():
            self.tree.delete(item)

    def clear(self) -> None:
        for item in self.tree.get_children():
            self.tree.delete(item)


class RowEditor(tk.Toplevel):
    def __init__(self, master: tk.Misc, columns: list[tuple[str, str, int]], values: list[Any]) -> None:
        super().__init__(master)
        self.title("Editar registro")
        self.resizable(False, False)
        self.result: list[str] | None = None
        self.vars: list[tk.StringVar] = []
        body = ttk.Frame(self, padding=12)
        body.pack(fill="both", expand=True)
        for idx, (_key, label, _width) in enumerate(columns):
            ttk.Label(body, text=label).grid(row=idx, column=0, sticky="w", pady=3)
            var = tk.StringVar(value=str(values[idx]) if idx < len(values) else "")
            entry = ttk.Entry(body, textvariable=var, width=42)
            entry.grid(row=idx, column=1, sticky="ew", pady=3, padx=(8, 0))
            self.vars.append(var)
        buttons = ttk.Frame(body)
        buttons.grid(row=len(columns), column=0, columnspan=2, sticky="e", pady=(10, 0))
        ttk.Button(buttons, text="Cancelar", command=self.destroy).pack(side="right")
        ttk.Button(buttons, text="OK", command=self.ok).pack(side="right", padx=(0, 8))
        self.bind("<Return>", lambda _e: self.ok())
        self.bind("<Escape>", lambda _e: self.destroy())
        self.grab_set()
        self.transient(master.winfo_toplevel())

    def ok(self) -> None:
        self.result = [var.get() for var in self.vars]
        self.destroy()


class IRSimpleApp:
    def __init__(self) -> None:
        self.root = tk.Tk()
        self.root.title(APP_TITLE)
        self.root.geometry("1320x820")
        self.trades: list[Trade] = []
        self.movements: list[Movement] = []
        self.result: CalculationResult | None = None
        self.results_by_year: dict[int, CalculationResult] = {}
        self.updating_history = False
        self.calculating = False

        default_year = date.today().year - 1
        self.username = DEFAULT_USER
        self.user_id = get_user_id(self.username)
        self.config: dict[str, Any] = {}
        self._load_config_from_db()
        if not self.config and self.username == "local":
            self.config.update(read_config())
        self.start_year_var = tk.IntVar(value=int(self.config.get("start_year") or default_year))
        self.end_year_var = tk.IntVar(value=int(self.config.get("end_year") or default_year))
        self.selected_year_var = tk.StringVar(value=str(self.config.get("selected_year") or self.end_year_var.get()))
        self.name_var = tk.StringVar(value=str(self.config.get("name") or ""))
        self.consolidated_files: list[str] = list(self.config.get("consolidated_files", []))
        self.benchmark_cache: dict[str, dict[str, str]] = self._load_benchmark_cache()
        default_neg = self.config.get("negociacao_file") or str(latest_matching_file(["negociacao-*.xlsx"], BASE_DIR / "negociacao-2026-05-17-16-41-46.xlsx"))
        default_mov = self.config.get("movimentacao_file") or str(latest_matching_file(["movimentacao-*.xlsx"], BASE_DIR / "movimentacao-2026-05-17-16-43-42.xlsx"))
        self.neg_file_var = tk.StringVar(value=default_neg)
        self.mov_file_var = tk.StringVar(value=default_mov)
        self.loss_vars = {
            key: tk.StringVar(value=str(self.config.get(f"loss_{key}") or "0,00"))
            for key in ["normal", "daytrade", "fii"]
        }
        self.loss_start_date_var = tk.StringVar(value=str(self.config.get("loss_start_date") or f"31/12/{int(self.start_year_var.get()) - 1}"))
        self.status_var = tk.StringVar(value="Importe as planilhas da B3 ou informe os dados manualmente.")
        self._build_ui()
        self._bind_persistent_header_fields()

    def _build_ui(self) -> None:
        root = ttk.Frame(self.root, padding=10)
        root.pack(fill="both", expand=True)
        top = ttk.Frame(root)
        top.pack(fill="x")
        ttk.Label(top, text="Nome").pack(side="left")
        ttk.Entry(top, textvariable=self.name_var, width=36).pack(side="left", padx=(6, 16))
        ttk.Label(top, text="Ano inicial").pack(side="left")
        ttk.Spinbox(top, from_=2000, to=2100, textvariable=self.start_year_var, width=8).pack(side="left", padx=(6, 10))
        ttk.Label(top, text="Ano-calendario final").pack(side="left")
        ttk.Spinbox(top, from_=2000, to=2100, textvariable=self.end_year_var, width=8).pack(side="left", padx=(6, 10))
        ttk.Label(top, text="Visualizar ano-calendario").pack(side="left")
        self.year_combo = ttk.Combobox(top, textvariable=self.selected_year_var, width=8, state="readonly", values=[str(self.end_year_var.get())])
        self.year_combo.pack(side="left", padx=(6, 16))
        self.year_combo.bind("<<ComboboxSelected>>", lambda _event: (self._save_config(), self._refresh_results()))
        self.calculate_button = ttk.Button(top, text="Calcular", command=self.calculate)
        self.calculate_button.pack(side="left")
        self.export_excel_button = ttk.Button(top, text="Exportar Excel", command=self.export_excel)
        self.export_excel_button.pack(side="left", padx=6)
        self.export_pdf_button = ttk.Button(top, text="Exportar PDF", command=self.export_pdf)
        self.export_pdf_button.pack(side="left")
        ttk.Label(root, textvariable=self.status_var, foreground="#334155").pack(fill="x", pady=(8, 8))
        self.exercise_var = tk.StringVar(value="")
        ttk.Label(root, textvariable=self.exercise_var, foreground="#475569").pack(fill="x", pady=(0, 6))

        notebook = ttk.Notebook(root)
        notebook.pack(fill="both", expand=True)
        self._build_import_tab(notebook)
        self._build_brokerage_notes_tab(notebook)
        self._build_manual_tab(notebook)
        self._build_results_tab(notebook)
        self._build_annual_tab(notebook)
        self._build_history_tab(notebook)
        self._build_monthly_portfolio_tab(notebook)
        self._build_wealth_tab(notebook)
        self._build_options_tab(notebook)
        self._build_b3_check_tab(notebook)
        self._build_consolidated_reports_tab(notebook)
        self._load_manual_data_from_db()
        self._load_option_events()
        self._load_cached_imported_data()
        self._refresh_consolidated_reports_view()

    def _build_import_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=12)
        notebook.add(tab, text="Importacao")
        tab.columnconfigure(1, weight=1)
        ttk.Label(tab, text="Arquivo negociacao B3").grid(row=0, column=0, sticky="w", pady=4)
        ttk.Entry(tab, textvariable=self.neg_file_var).grid(row=0, column=1, sticky="ew", padx=8)
        ttk.Button(tab, text="Selecionar", command=lambda: self.pick_file(self.neg_file_var)).grid(row=0, column=2)
        ttk.Label(tab, text="Arquivo movimentacao B3").grid(row=1, column=0, sticky="w", pady=4)
        ttk.Entry(tab, textvariable=self.mov_file_var).grid(row=1, column=1, sticky="ew", padx=8)
        ttk.Button(tab, text="Selecionar", command=lambda: self.pick_file(self.mov_file_var)).grid(row=1, column=2)
        ttk.Button(tab, text="Importar planilhas", command=lambda: self.import_files(show_message=True)).grid(row=2, column=0, sticky="w", pady=(12, 4))
        ttk.Button(tab, text="Atualizar CNPJs online", command=self.update_asset_cnpjs_online).grid(row=2, column=1, sticky="w", pady=(12, 4), padx=8)

        losses = ttk.LabelFrame(tab, text="Prejuizos a compensar na data inicial", padding=10)
        losses.grid(row=3, column=0, columnspan=3, sticky="ew", pady=(14, 0))
        ttk.Label(losses, text="Data-base do saldo").grid(row=0, column=0, sticky="w", padx=(0, 4))
        loss_date_entry = ttk.Entry(losses, textvariable=self.loss_start_date_var, width=12)
        loss_date_entry.grid(row=0, column=1, sticky="w", padx=(0, 6))
        loss_date_entry.bind("<FocusOut>", lambda _event: self._normalize_loss_date_field(show_error=False))
        ttk.Button(losses, text="31/12 ano anterior", command=self._set_loss_base_previous_year).grid(row=0, column=2, sticky="w", padx=(0, 10))
        ttk.Label(losses, text="Digite 31122024 ou 31/12/2024", foreground="#475569").grid(row=0, column=3, columnspan=4, sticky="w")
        labels = [("normal", "Operacoes normais"), ("daytrade", "Day trade"), ("fii", "FII/FIAGRO")]
        for idx, (key, label) in enumerate(labels):
            ttk.Label(losses, text=label).grid(row=1, column=idx * 2, sticky="w", padx=(0, 4), pady=(8, 0))
            ttk.Entry(losses, textvariable=self.loss_vars[key], width=14).grid(row=1, column=idx * 2 + 1, sticky="w", padx=(0, 12), pady=(8, 0))

        note = (
            "Regras usadas: lucro em acoes comuns com vendas mensais ate R$ 20.000,00 fica como rendimento isento; "
            "demais operacoes comuns usam 15%; day trade e FII usam 20%. Custos de corretagem/IRRF podem ser ajustados por eventos/notas."
        )
        ttk.Label(tab, text=note, wraplength=980, foreground="#475569").grid(row=4, column=0, columnspan=3, sticky="w", pady=(18, 0))

        preview = ttk.LabelFrame(tab, text="Previa da importacao", padding=8)
        preview.grid(row=5, column=0, columnspan=3, sticky="nsew", pady=(14, 0))
        tab.rowconfigure(5, weight=1)
        preview.columnconfigure(0, weight=1)
        preview.columnconfigure(1, weight=1)
        preview.rowconfigure(1, weight=1)

        self.trade_filter_col_var = tk.StringVar(value="todas")
        self.trade_filter_text_var = tk.StringVar(value="")
        self.trade_sort_date_var = tk.StringVar(value="crescente")
        self.movement_filter_col_var = tk.StringVar(value="todas")
        self.movement_filter_text_var = tk.StringVar(value="")
        self.movement_sort_date_var = tk.StringVar(value="crescente")
        trade_filter = ttk.Frame(preview)
        movement_filter = ttk.Frame(preview)
        trade_filter.grid(row=0, column=0, sticky="ew", padx=(0, 6), pady=(0, 4))
        movement_filter.grid(row=0, column=1, sticky="ew", padx=(6, 0), pady=(0, 4))
        for frame, col_var, text_var, columns in [
            (trade_filter, self.trade_filter_col_var, self.trade_filter_text_var, ["todas", "data", "tipo", "mercado", "codigo", "quantidade", "valor", "categoria"]),
            (movement_filter, self.movement_filter_col_var, self.movement_filter_text_var, ["todas", "data", "movimentacao", "produto", "quantidade", "valor", "categoria"]),
        ]:
            ttk.Label(frame, text="Filtrar").pack(side="left")
            ttk.Combobox(frame, textvariable=col_var, values=columns, width=14, state="readonly").pack(side="left", padx=4)
            entry = ttk.Entry(frame, textvariable=text_var, width=28)
            entry.pack(side="left")
            ttk.Button(frame, text="Aplicar", command=self._refresh_import_preview).pack(side="left", padx=4)
            ttk.Button(frame, text="Limpar", command=lambda v=text_var: (v.set(""), self._refresh_import_preview())).pack(side="left")
            text_var.trace_add("write", lambda *_args: self._refresh_import_preview())
        ttk.Label(trade_filter, text="Data").pack(side="left", padx=(10, 0))
        ttk.Combobox(trade_filter, textvariable=self.trade_sort_date_var, values=["crescente", "decrescente"], width=12, state="readonly").pack(side="left", padx=4)
        self.trade_sort_date_var.trace_add("write", lambda *_args: self._refresh_import_preview())
        ttk.Label(movement_filter, text="Data").pack(side="left", padx=(10, 0))
        ttk.Combobox(movement_filter, textvariable=self.movement_sort_date_var, values=["crescente", "decrescente"], width=12, state="readonly").pack(side="left", padx=4)
        self.movement_sort_date_var.trace_add("write", lambda *_args: self._refresh_import_preview())

        self.trade_preview = ttk.Treeview(
            preview,
            columns=["data", "tipo", "mercado", "codigo", "quantidade", "valor", "categoria"],
            show="headings",
            height=8,
        )
        self.movement_preview = ttk.Treeview(
            preview,
            columns=["data", "movimentacao", "produto", "quantidade", "valor", "categoria"],
            show="headings",
            height=8,
        )
        for tree in (self.trade_preview, self.movement_preview):
            for col in tree["columns"]:
                tree.heading(col, text=col)
                tree.column(col, width=105, anchor="w")
        self.trade_preview.grid(row=1, column=0, sticky="nsew", padx=(0, 6))
        self.movement_preview.grid(row=1, column=1, sticky="nsew", padx=(6, 0))

    def _build_brokerage_notes_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="Notas corretagem")
        bar = ttk.Frame(tab)
        bar.pack(fill="x", pady=(0, 8))
        ttk.Button(bar, text="Importar notas PDF", command=self.import_brokerage_notes).pack(side="left")
        ttk.Button(bar, text="Limpar notas importadas", command=self.clear_brokerage_notes).pack(side="left", padx=6)
        self.brokerage_notes_status_var = tk.StringVar(value="")
        ttk.Label(bar, textvariable=self.brokerage_notes_status_var, foreground="#475569").pack(side="left", padx=10)
        cols = ["data", "nota", "irrf_comum", "irrf_dt", "irrf_fii", "base", "arquivo"]
        labels = ["Data pregão", "Nota", "IRRF comum", "IRRF day trade", "IRRF FII", "Base IRRF", "Arquivo"]
        self.brokerage_notes_tree = ttk.Treeview(tab, columns=cols, show="headings", height=16)
        for col, label in zip(cols, labels):
            self.brokerage_notes_tree.heading(col, text=label)
            self.brokerage_notes_tree.column(col, width=130 if col != "arquivo" else 480, anchor="e" if col.startswith("irrf") or col == "base" else "w")
        self.brokerage_notes_tree.pack(fill="both", expand=True)

        compare_frame = ttk.LabelFrame(tab, text="Conferencia mensal: notas x calculo principal B3", padding=6)
        compare_frame.pack(fill="both", expand=True, pady=(8, 0))
        compare_cols = ["mes", "venda_normal_notas", "venda_normal_b3", "venda_fii_notas", "venda_fii_b3", "opcoes_notas", "opcoes_b3", "irrf_comum", "irrf_dt", "irrf_fii"]
        compare_labels = ["Mes", "Vendas normais notas", "Vendas normais B3", "Vendas FII notas", "Vendas FII B3", "Opcoes notas", "Opcoes B3", "IRRF comum", "IRRF DT", "IRRF FII"]
        self.brokerage_compare_tree = ttk.Treeview(compare_frame, columns=compare_cols, show="headings", height=8)
        for col, label in zip(compare_cols, compare_labels):
            self.brokerage_compare_tree.heading(col, text=label)
            self.brokerage_compare_tree.column(col, width=130, anchor="e" if col != "mes" else "w")
        self.brokerage_compare_tree.pack(fill="both", expand=True)
        self.refresh_brokerage_notes_view()

    def _build_manual_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="Dados manuais")
        panes = ttk.PanedWindow(tab, orient="vertical")
        panes.pack(fill="both", expand=True)

        pos_frame = ttk.LabelFrame(panes, text="Posicoes existentes no inicio do calculo", padding=8)
        self.position_table = EditableTable(
            pos_frame,
            [("ativo", "Ativo", 90), ("quantidade", "Quantidade", 90), ("custo", "Custo total", 100), ("categoria", "Categoria", 100), ("corretora", "Corretora", 180)],
            height=8,
        )
        self.position_table.pack(fill="both", expand=True)
        panes.add(pos_frame, weight=1)

        event_frame = ttk.LabelFrame(panes, text="Notas/eventos: exercicio de opcoes, migracao, desdobramento, bonificacao, subscricao", padding=8)
        self.event_table = EditableTable(
            event_frame,
            [("data", "Data", 90), ("tipo", "Tipo", 150), ("ativo", "Ativo", 90), ("quantidade", "Quantidade/Fator", 110), ("valor", "Valor/Custo", 100), ("categoria", "Categoria", 90), ("observacao", "Observacao", 300)],
            height=8,
        )
        self.event_table.pack(fill="both", expand=True)
        panes.add(event_frame, weight=1)

    def _build_results_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="Calculo mensal")
        cols = ["mes", "normal", "daytrade", "fii", "opcoes", "futuro", "base_normal", "base_dt", "base_fii", "imposto"]
        labels = ["Mes", "Resultado normal", "Day trade", "FII", "Opcoes", "Futuro", "Base normal", "Base DT", "Base FII", "Imposto a pagar"]
        self.month_tree = ttk.Treeview(tab, columns=cols, show="headings", height=14)
        for col, label in zip(cols, labels):
            self.month_tree.heading(col, text=label)
            self.month_tree.column(col, width=120, anchor="e" if col != "mes" else "w")
        self.month_tree.bind("<Double-1>", self._open_month_value_detail)
        self.month_tree.pack(fill="both", expand=True)
        ttk.Label(tab, text="Duplo clique em um valor diferente de zero para ver os registros que formaram o calculo.", foreground="#475569").pack(fill="x", pady=(4, 0))

        warn_frame = ttk.LabelFrame(tab, text="Alertas de conferencia", padding=8)
        warn_frame.pack(fill="both", expand=True, pady=(8, 0))
        self.warn_text = tk.Text(warn_frame, height=8, wrap="word")
        self.warn_text.pack(fill="both", expand=True)
        solution_frame = ttk.LabelFrame(warn_frame, text="Solucoes sugeridas para os alertas", padding=6)
        solution_frame.pack(fill="both", expand=True, pady=(8, 0))
        self.alert_solution_tree = ttk.Treeview(
            solution_frame,
            columns=["idx", "ativo", "data", "falta", "valor", "status", "sugestao", "evento"],
            show="headings",
            height=5,
        )
        for col, label, width in [
            ("idx", "#", 45),
            ("ativo", "Ativo", 90),
            ("data", "Data", 90),
            ("falta", "Qtd faltante", 100),
            ("valor", "Valor sugerido", 120),
            ("status", "Status", 140),
            ("sugestao", "Possivel solucao", 360),
            ("evento", "Evento de origem", 420),
        ]:
            self.alert_solution_tree.heading(col, text=label)
            self.alert_solution_tree.column(col, width=width, anchor="e" if col in {"falta", "valor"} else "w")
        self.alert_solution_tree.tag_configure("aplicado", background="#dcfce7")
        self.alert_solution_tree.tag_configure("pendente", background="#fff7ed")
        self.alert_solution_tree.pack(fill="both", expand=True)
        solution_buttons = ttk.Frame(solution_frame)
        solution_buttons.pack(fill="x", pady=(6, 0))
        ttk.Button(solution_buttons, text="Registrar posicao inicial", command=self._register_alert_initial_position).pack(side="left")
        ttk.Button(solution_buttons, text="Registrar migracao/transferencia", command=lambda: self._register_alert_event("migracao/transferencia")).pack(side="left", padx=6)
        ttk.Button(solution_buttons, text="Registrar split", command=lambda: self._register_corporate_alert_event("split/desdobramento")).pack(side="left")
        ttk.Button(solution_buttons, text="Registrar bonificacao", command=lambda: self._register_corporate_alert_event("bonificacao")).pack(side="left", padx=6)
        ttk.Button(solution_buttons, text="Registrar venda ativa", command=lambda: self._register_alert_event("venda descoberta ativa")).pack(side="left", padx=6)
        ttk.Button(solution_buttons, text="Recalcular agora", command=self.calculate).pack(side="left", padx=(12, 0))

    def _build_annual_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="Declaracao anual")
        panes = ttk.PanedWindow(tab, orient="vertical")
        panes.pack(fill="both", expand=True)
        self.annual_trees: dict[str, ttk.Treeview] = {}
        self.annual_total_vars: dict[str, tk.StringVar] = {}
        configs = [
            ("bens", "Bens e Direitos", ["codigo", "discriminacao", "qtd 31/12 anterior", "valor 31/12 anterior", "qtd 31/12 atual", "valor 31/12 atual"]),
            ("isentos", "Rendimentos Isentos", ["codigo", "descricao", "valor"]),
            ("sujeitos", "Rendimentos Sujeitos Exclusiva", ["codigo", "descricao", "valor"]),
            ("dividas", "Dividas e Onus", ["codigo", "descricao", "31/12 anterior", "31/12 atual", "valor pago"]),
        ]
        for key, title, cols in configs:
            frame = ttk.LabelFrame(panes, text=title, padding=6)
            tree = ttk.Treeview(frame, columns=cols, show="headings", height=5)
            for col in cols:
                tree.heading(col, text=col)
                tree.column(col, width=180, anchor="w")
            tree.pack(fill="both", expand=True)
            total_var = tk.StringVar(value="Totais: -")
            ttk.Label(frame, textvariable=total_var, anchor="e", foreground="#0f172a").pack(fill="x", pady=(4, 0))
            panes.add(frame, weight=1)
            self.annual_trees[key] = tree
            self.annual_total_vars[key] = total_var

    def _build_history_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="Historico")
        panes = ttk.PanedWindow(tab, orient="vertical")
        panes.pack(fill="both", expand=True)

        summary_frame = ttk.LabelFrame(panes, text="Resumo ano a ano", padding=6)
        summary_cols = [
            "ano",
            "exercicio",
            "ativos",
            "valor_carteira",
            "resultado_normal",
            "resultado_daytrade",
            "resultado_fii",
            "resultado_opcoes",
            "resultado_futuro",
            "imposto",
            "prejuizos_finais",
        ]
        summary_labels = [
            "Ano",
            "Exercicio",
            "Ativos carteira",
            "Valor carteira",
            "Lucro/Prej. normal",
            "Day trade",
            "FII",
            "Opcoes",
            "Futuro",
            "Imposto",
            "Prejuizos finais",
        ]
        self.history_summary_tree = ttk.Treeview(summary_frame, columns=summary_cols, show="headings", height=7)
        for col, label in zip(summary_cols, summary_labels):
            self.history_summary_tree.heading(col, text=label)
            self.history_summary_tree.column(col, width=120, anchor="e" if col not in {"ano", "exercicio"} else "center")
        self.history_summary_tree.pack(fill="both", expand=True)
        self.history_summary_tree.bind("<<TreeviewSelect>>", self._on_history_year_selected)
        panes.add(summary_frame, weight=1)

        detail_frame = ttk.Frame(panes)
        detail_panes = ttk.PanedWindow(detail_frame, orient="horizontal")
        detail_panes.pack(fill="both", expand=True)
        positions_frame = ttk.LabelFrame(detail_panes, text="Carteira no fim do ano selecionado", padding=6)
        self.history_positions_tree = ttk.Treeview(
            positions_frame,
            columns=["ativo", "categoria", "quantidade", "preco_medio", "custo_anterior", "custo_atual"],
            show="headings",
            height=12,
        )
        for col, label, width in [
            ("ativo", "Ativo", 90),
            ("categoria", "Categoria", 90),
            ("quantidade", "Quantidade", 100),
            ("preco_medio", "Preco medio", 110),
            ("custo_anterior", "31/12 anterior", 120),
            ("custo_atual", "31/12 atual", 120),
        ]:
            self.history_positions_tree.heading(col, text=label)
            self.history_positions_tree.column(col, width=width, anchor="e" if col != "ativo" and col != "categoria" else "w")
        self.history_positions_tree.pack(fill="both", expand=True)
        detail_panes.add(positions_frame, weight=1)

        monthly_frame = ttk.LabelFrame(detail_panes, text="Lucro, base e imposto por mes", padding=6)
        self.history_month_tree = ttk.Treeview(
            monthly_frame,
            columns=["mes", "normal", "daytrade", "fii", "opcoes", "futuro", "imposto", "prejuizos"],
            show="headings",
            height=12,
        )
        for col, label, width in [
            ("mes", "Mes", 60),
            ("normal", "Normal", 100),
            ("daytrade", "Day trade", 100),
            ("fii", "FII", 100),
            ("opcoes", "Opcoes", 100),
            ("futuro", "Futuro", 100),
            ("imposto", "Imposto", 100),
            ("prejuizos", "Prejuizos fim", 130),
        ]:
            self.history_month_tree.heading(col, text=label)
            self.history_month_tree.column(col, width=width, anchor="e" if col != "mes" else "w")
        self.history_month_tree.pack(fill="both", expand=True)
        detail_panes.add(monthly_frame, weight=1)

        loans_frame = ttk.LabelFrame(detail_panes, text="Acoes em emprestimo por mes", padding=6)
        self.history_loans_tree = ttk.Treeview(
            loans_frame,
            columns=["mes", "ativo", "quantidade"],
            show="headings",
            height=12,
        )
        for col, label, width in [("mes", "Mes", 70), ("ativo", "Ativo", 90), ("quantidade", "Quantidade", 110)]:
            self.history_loans_tree.heading(col, text=label)
            self.history_loans_tree.column(col, width=width, anchor="e" if col == "quantidade" else "w")
        self.history_loans_tree.pack(fill="both", expand=True)
        detail_panes.add(loans_frame, weight=1)
        panes.add(detail_frame, weight=2)

    def _build_options_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="Opcoes a conferir")
        self.pending_options_tree = ttk.Treeview(
            tab,
            columns=["ano", "codigo", "vencimento", "quantidade", "premio", "ultima_data", "situacao"],
            show="headings",
            height=16,
        )
        for col, label, width in [
            ("ano", "Ano", 70),
            ("codigo", "Opcao", 110),
            ("vencimento", "Vencimento", 100),
            ("quantidade", "Qtd aberta (+compra/-venda)", 150),
            ("premio", "Premio liquido", 120),
            ("ultima_data", "Ultima operacao", 110),
            ("situacao", "Conferencia necessaria", 360),
        ]:
            self.pending_options_tree.heading(col, text=label)
            self.pending_options_tree.column(col, width=width, anchor="e" if col in {"quantidade", "premio"} else "w")
        self.pending_options_tree.tag_configure("registered", background="#dcfce7")
        self.pending_options_tree.pack(fill="both", expand=True)
        buttons = ttk.Frame(tab)
        buttons.pack(fill="x", pady=(8, 0))
        ttk.Button(buttons, text="Marcar como virou po", command=lambda: self._add_option_resolution_event("opcao virou po")).pack(side="left")
        ttk.Button(buttons, text="Marcar como exercida", command=lambda: self._add_option_resolution_event("exercicio de opcoes")).pack(side="left", padx=8)
        ttk.Button(buttons, text="Informar recompra", command=self._add_option_repurchase_event).pack(side="left")
        ttk.Button(buttons, text="Liquidacao trava sem acao", command=self._add_option_spread_settlement_event).pack(side="left", padx=8)
        ttk.Button(buttons, text="Marcar como ainda ativa", command=lambda: self._add_option_resolution_event("opcao vendida ainda ativa")).pack(side="left")
        ttk.Button(buttons, text="Recalcular agora", command=self.calculate).pack(side="left", padx=(12, 0))
        ttk.Label(
            tab,
            text="As marcacoes ficam registradas abaixo e tambem na aba Dados manuais. Depois de marcar, clique em Recalcular agora para atualizar a lista de pendencias e os calculos.",
            foreground="#475569",
            wraplength=1100,
        ).pack(fill="x", pady=(8, 0))
        registered_frame = ttk.LabelFrame(tab, text="Opcoes ja registradas para conferencia", padding=6)
        registered_frame.pack(fill="both", expand=False, pady=(8, 0))
        self.option_events_tree = ttk.Treeview(
            registered_frame,
            columns=["data", "tipo", "ativo", "quantidade", "valor", "observacao"],
            show="headings",
            height=5,
        )
        for col, label, width in [
            ("data", "Data", 90),
            ("tipo", "Tipo", 140),
            ("ativo", "Opcao", 100),
            ("quantidade", "Quantidade", 100),
            ("valor", "Valor", 100),
            ("observacao", "Observacao", 520),
        ]:
            self.option_events_tree.heading(col, text=label)
            self.option_events_tree.column(col, width=width, anchor="e" if col in {"quantidade", "valor"} else "w")
        self.option_events_tree.pack(fill="both", expand=True)
        event_buttons = ttk.Frame(registered_frame)
        event_buttons.pack(fill="x", pady=(6, 0))
        ttk.Button(event_buttons, text="Salvar conferencias", command=self._save_option_events).pack(side="left")
        ttk.Button(event_buttons, text="Editar selecionada", command=self._edit_option_event).pack(side="left", padx=6)
        ttk.Button(event_buttons, text="Excluir selecionada", command=self._delete_option_event).pack(side="left")

    def _build_monthly_portfolio_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="Carteira mensal")
        self.monthly_portfolio_tree = ttk.Treeview(
            tab,
            columns=["ano", "mes", "ativo", "categoria", "quantidade", "preco_medio", "custo", "emprestado"],
            show="headings",
            height=20,
        )
        for col, label, width in [
            ("ano", "Ano", 60),
            ("mes", "Mes", 60),
            ("ativo", "Ativo", 90),
            ("categoria", "Categoria", 90),
            ("quantidade", "Quantidade", 100),
            ("preco_medio", "Preco medio", 110),
            ("custo", "Custo total", 120),
            ("emprestado", "Em emprestimo", 120),
        ]:
            self.monthly_portfolio_tree.heading(col, text=label)
            self.monthly_portfolio_tree.column(col, width=width, anchor="e" if col not in {"ativo", "categoria", "mes"} else "w")
        self.monthly_portfolio_tree.pack(fill="both", expand=True)

    def _build_wealth_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="Patrimonio")
        top = ttk.Frame(tab)
        top.pack(fill="x", pady=(0, 8))
        ttk.Label(top, text="Visualizar").pack(side="left")
        self.wealth_view_var = tk.StringVar(value="anual")
        view_combo = ttk.Combobox(top, textvariable=self.wealth_view_var, values=["anual", "mensal"], width=10, state="readonly")
        view_combo.pack(side="left", padx=(6, 16))
        view_combo.bind("<<ComboboxSelected>>", lambda _event: self._refresh_wealth_view())
        self.show_ibov_var = tk.BooleanVar(value=True)
        self.show_selic_var = tk.BooleanVar(value=True)
        self.show_poupanca_var = tk.BooleanVar(value=True)
        for text, var in [("Ibovespa", self.show_ibov_var), ("Selic", self.show_selic_var), ("Poupanca", self.show_poupanca_var)]:
            ttk.Checkbutton(top, text=text, variable=var, command=self._draw_wealth_chart).pack(side="left", padx=(0, 8))
        ttk.Button(top, text="Atualizar indices", command=self._update_benchmarks_from_web).pack(side="left", padx=(4, 12))
        ttk.Label(
            top,
            text="Comparacao normalizada pela base do primeiro ponto de patrimonio.",
            foreground="#475569",
        ).pack(side="left")

        body = ttk.PanedWindow(tab, orient="horizontal")
        body.pack(fill="both", expand=True)
        table_frame = ttk.LabelFrame(body, text="Dados do grafico", padding=6)
        self.wealth_tree = ttk.Treeview(
            table_frame,
            columns=["periodo", "patrimonio", "ibovespa", "selic", "poupanca", "ganho_acoes", "ganho_fundos", "outros", "total", "imposto", "prejuizos"],
            show="headings",
            height=18,
        )
        for col, label, width in [
            ("periodo", "Periodo", 110),
            ("patrimonio", "Patrimonio", 130),
            ("ibovespa", "Ibovespa", 120),
            ("selic", "Selic", 120),
            ("poupanca", "Poupanca", 120),
            ("ganho_acoes", "Ganho acoes/ETF", 130),
            ("ganho_fundos", "Ganho fundos/FII", 130),
            ("outros", "Opcoes/Futuro/DT", 130),
            ("total", "Total", 120),
            ("imposto", "Imposto", 110),
            ("prejuizos", "Prejuizos finais", 180),
        ]:
            self.wealth_tree.heading(col, text=label)
            self.wealth_tree.column(col, width=width, anchor="e" if col != "periodo" else "w")
        self.wealth_tree.pack(fill="both", expand=True)
        body.add(table_frame, weight=1)

        chart_frame = ttk.LabelFrame(body, text="Evolucao do patrimonio", padding=6)
        self.wealth_canvas = tk.Canvas(chart_frame, background="white", highlightthickness=1, highlightbackground="#cbd5e1")
        self.wealth_canvas.pack(fill="both", expand=True)
        self.wealth_canvas.bind("<Configure>", lambda _event: self._draw_wealth_chart())
        body.add(chart_frame, weight=2)
        self.wealth_points: list[tuple[str, Decimal, Decimal, Decimal]] = []
        self.wealth_benchmark_points: dict[str, list[tuple[str, Decimal]]] = {}

    def _build_b3_check_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="Conferencia B3")
        bar = ttk.Frame(tab)
        bar.pack(fill="x", pady=(0, 8))
        ttk.Button(bar, text="Conferir relatorios B3", command=self._refresh_b3_check).pack(side="left")
        self.b3_check_status_var = tk.StringVar(value="Clique em Calcular e depois em Conferir relatorios B3.")
        ttk.Label(bar, textvariable=self.b3_check_status_var, foreground="#475569").pack(side="left", padx=10)
        self.b3_check_tree = ttk.Treeview(
            tab,
            columns=["ano", "mes", "tipo", "ativo", "b3", "calculado", "diferenca", "arquivo"],
            show="headings",
            height=22,
        )
        for col, label, width in [
            ("ano", "Ano", 60),
            ("mes", "Mes", 60),
            ("tipo", "Tipo", 100),
            ("ativo", "Ativo", 90),
            ("b3", "B3", 100),
            ("calculado", "Calculado", 100),
            ("diferenca", "Diferenca", 100),
            ("arquivo", "Arquivo", 280),
        ]:
            self.b3_check_tree.heading(col, text=label)
            self.b3_check_tree.column(col, width=width, anchor="e" if col in {"b3", "calculado", "diferenca"} else "w")
        self.b3_check_tree.tag_configure("diff", background="#fee2e2")
        self.b3_check_tree.tag_configure("ok", background="#ecfdf5")
        self.b3_check_tree.pack(fill="both", expand=True)

    def _build_consolidated_reports_tab(self, notebook: ttk.Notebook) -> None:
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="Relatorios B3")
        bar = ttk.Frame(tab)
        bar.pack(fill="x", pady=(0, 8))
        ttk.Button(bar, text="Importar relatorios consolidados", command=self._import_consolidated_reports).pack(side="left")
        ttk.Button(bar, text="Remover selecionado", command=self._remove_consolidated_report).pack(side="left", padx=6)
        ttk.Button(bar, text="Atualizar visualizacao", command=self._refresh_consolidated_reports_view).pack(side="left")
        self.consolidated_status_var = tk.StringVar(value="")
        ttk.Label(bar, textvariable=self.consolidated_status_var, foreground="#475569").pack(side="left", padx=10)

        panes = ttk.PanedWindow(tab, orient="horizontal")
        panes.pack(fill="both", expand=True)
        files_frame = ttk.LabelFrame(panes, text="Arquivos importados", padding=6)
        self.consolidated_files_tree = ttk.Treeview(files_frame, columns=["periodo", "arquivo"], show="headings", height=18)
        for col, label, width in [("periodo", "Periodo", 100), ("arquivo", "Arquivo", 420)]:
            self.consolidated_files_tree.heading(col, text=label)
            self.consolidated_files_tree.column(col, width=width, anchor="w")
        self.consolidated_files_tree.pack(fill="both", expand=True)
        self.consolidated_files_tree.bind("<<TreeviewSelect>>", lambda _event: self._show_selected_consolidated_report())
        panes.add(files_frame, weight=1)

        data_frame = ttk.LabelFrame(panes, text="Sheets do relatorio selecionado", padding=6)
        self.consolidated_sheet_notebook = ttk.Notebook(data_frame)
        self.consolidated_sheet_notebook.pack(fill="both", expand=True)
        panes.add(data_frame, weight=2)

    def _import_consolidated_reports(self) -> None:
        paths = filedialog.askopenfilenames(
            title="Selecionar relatorios consolidados B3",
            filetypes=[("Excel", "*.xlsx *.xls"), ("Todos", "*.*")],
            initialdir=str(Path.home() / "Downloads"),
        )
        if not paths:
            return
        known = set(self.consolidated_files)
        for path in paths:
            if path not in known:
                self.consolidated_files.append(path)
        self._save_config()
        self._refresh_consolidated_reports_view()

    def _remove_consolidated_report(self) -> None:
        selected = self.consolidated_files_tree.selection()
        if not selected:
            return
        remove_paths = {self.consolidated_files_tree.item(item, "values")[1] for item in selected}
        self.consolidated_files = [path for path in self.consolidated_files if path not in remove_paths]
        self._save_config()
        self._refresh_consolidated_reports_view()

    def _refresh_consolidated_reports_view(self) -> None:
        if not hasattr(self, "consolidated_files_tree"):
            return
        for item in self.consolidated_files_tree.get_children():
            self.consolidated_files_tree.delete(item)
        self._clear_consolidated_sheet_tabs()
        valid_files = []
        for file in self.consolidated_files:
            path = Path(file)
            if not path.exists():
                continue
            period = parse_consolidated_period(path)
            label = f"{period[0]}-{period[1]:02d}" if period else "periodo?"
            self.consolidated_files_tree.insert("", "end", values=[label, str(path)])
            valid_files.append(str(path))
        self.consolidated_files = valid_files
        self.consolidated_status_var.set(f"{len(valid_files)} relatorios consolidados carregados.")
        self._save_config()

    def _clear_consolidated_sheet_tabs(self) -> None:
        if not hasattr(self, "consolidated_sheet_notebook"):
            return
        for tab_id in self.consolidated_sheet_notebook.tabs():
            self.consolidated_sheet_notebook.forget(tab_id)

    def _add_consolidated_sheet_tab(self, sheet: dict[str, Any]) -> None:
        frame = ttk.Frame(self.consolidated_sheet_notebook, padding=4)
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(0, weight=1)
        raw_columns = sheet.get("columns") or ["Sem dados"]
        columns = [f"c{idx}" for idx, _col in enumerate(raw_columns)]
        tree = ttk.Treeview(frame, columns=columns, show="headings", height=18)
        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        for key, label in zip(columns, raw_columns):
            tree.heading(key, text=str(label))
            tree.column(key, width=max(110, min(260, len(str(label)) * 10 + 40)), anchor="w")
        for row in sheet.get("rows", []):
            values = list(row[: len(columns)])
            if len(values) < len(columns):
                values.extend([""] * (len(columns) - len(values)))
            tree.insert("", "end", values=values)
        tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        self.consolidated_sheet_notebook.add(frame, text=str(sheet.get("name") or "Sheet")[:28])

    def _show_selected_consolidated_report(self) -> None:
        self._clear_consolidated_sheet_tabs()
        selected = self.consolidated_files_tree.selection()
        if not selected:
            return
        path = Path(self.consolidated_files_tree.item(selected[0], "values")[1])
        try:
            sheets = read_consolidated_sheets(path)
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"Falha ao ler relatorio consolidado:\n{exc}")
            return
        for sheet in sheets:
            self._add_consolidated_sheet_tab(sheet)
        self.consolidated_status_var.set(f"{path.name}: {len(sheets)} sheets exibidos.")

    def pick_file(self, var: tk.StringVar) -> None:
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xls"), ("Todos", "*.*")], initialdir=str(BASE_DIR))
        if path:
            var.set(path)
            self._save_config()

    def _load_config_from_db(self) -> None:
        try:
            with db_connect() as conn:
                rows = conn.execute("SELECT key, value FROM app_config WHERE user_id = ?", (self.user_id,)).fetchall()
                for row in rows:
                    self.config[row["key"]] = row["value"]
                files = conn.execute("SELECT file_type, path FROM source_files WHERE user_id = ? AND active = 1", (self.user_id,)).fetchall()
                consolidated = []
                for row in files:
                    if row["file_type"] == "negociacao":
                        self.config["negociacao_file"] = row["path"]
                    elif row["file_type"] == "movimentacao":
                        self.config["movimentacao_file"] = row["path"]
                    elif row["file_type"] == "consolidado":
                        consolidated.append(row["path"])
                if consolidated:
                    self.config["consolidated_files"] = consolidated
        except Exception as exc:
            debug_log(f"Falha ao carregar configuracao do banco: {exc}")

    def _load_benchmark_cache(self) -> dict[str, dict[str, str]]:
        raw = self.config.get("benchmark_cache")
        if not raw:
            return {}
        try:
            data = json.loads(raw) if isinstance(raw, str) else raw
            if isinstance(data, dict):
                return {
                    str(name): {str(period): str(value) for period, value in values.items()}
                    for name, values in data.items()
                    if isinstance(values, dict)
                }
        except Exception as exc:
            debug_log(f"Falha ao carregar cache de indices: {exc}")
        return {}

    def _save_config(self) -> None:
        self.config["name"] = self.name_var.get()
        self.config["start_year"] = str(self.start_year_var.get())
        self.config["end_year"] = str(self.end_year_var.get())
        self.config["selected_year"] = self.selected_year_var.get()
        self.config["loss_start_date"] = self.loss_start_date_var.get()
        for key, var in self.loss_vars.items():
            self.config[f"loss_{key}"] = var.get()
        self.config["negociacao_file"] = self.neg_file_var.get()
        self.config["movimentacao_file"] = self.mov_file_var.get()
        self.config["consolidated_files"] = self.consolidated_files
        self.config["benchmark_cache"] = json.dumps(self.benchmark_cache, ensure_ascii=False, sort_keys=True)
        if self.username == "local":
            write_config(self.config)
        try:
            with db_connect() as conn:
                for key, value in self.config.items():
                    if isinstance(value, list):
                        continue
                    conn.execute(
                        "INSERT INTO app_config (user_id, key, value) VALUES (?, ?, ?) ON CONFLICT(user_id, key) DO UPDATE SET value = excluded.value",
                        (self.user_id, key, str(value)),
                    )
                for file_type, path in [("negociacao", self.neg_file_var.get()), ("movimentacao", self.mov_file_var.get())]:
                    if path:
                        conn.execute("UPDATE source_files SET active = 0 WHERE user_id = ? AND file_type = ?", (self.user_id, file_type))
                        conn.execute(
                            "INSERT OR IGNORE INTO source_files (user_id, file_type, path, active) VALUES (?, ?, ?, 1)",
                            (self.user_id, file_type, path),
                        )
                        conn.execute("UPDATE source_files SET active = 1 WHERE user_id = ? AND file_type = ? AND path = ?", (self.user_id, file_type, path))
                conn.execute("UPDATE source_files SET active = 0 WHERE user_id = ? AND file_type = 'consolidado'", (self.user_id,))
                for path in self.consolidated_files:
                    conn.execute(
                        "INSERT OR IGNORE INTO source_files (user_id, file_type, path, active) VALUES (?, 'consolidado', ?, 1)",
                        (self.user_id, path),
                    )
                    conn.execute("UPDATE source_files SET active = 1 WHERE user_id = ? AND file_type = 'consolidado' AND path = ?", (self.user_id, path))
        except Exception as exc:
            debug_log(f"Falha ao salvar configuracao no banco: {exc}")

    def _bind_persistent_header_fields(self) -> None:
        for var in [self.name_var, self.selected_year_var, self.start_year_var, self.end_year_var, self.loss_start_date_var, *self.loss_vars.values()]:
            var.trace_add("write", lambda *_args: self._save_config())

    def _normalize_loss_date_field(self, show_error: bool = True) -> date | None:
        raw = self.loss_start_date_var.get().strip()
        if not raw:
            return None
        parsed = parse_date(raw)
        if parsed is None:
            if show_error:
                messagebox.showerror(APP_TITLE, "Informe a data-base do saldo no formato 31122024 ou 31/12/2024.")
            return None
        formatted = parsed.strftime("%d/%m/%Y")
        if raw != formatted:
            self.loss_start_date_var.set(formatted)
        return parsed

    def _set_loss_base_previous_year(self) -> None:
        try:
            year = int(self.start_year_var.get()) - 1
        except Exception:
            year = date.today().year - 1
        self.loss_start_date_var.set(f"31/12/{year}")

    def _save_imported_data_to_db(self, neg_file: Path, mov_file: Path) -> None:
        try:
            with db_connect() as conn:
                conn.execute("DELETE FROM trades WHERE user_id = ?", (self.user_id,))
                conn.execute("DELETE FROM movements WHERE user_id = ?", (self.user_id,))
                conn.executemany(
                    """
                    INSERT INTO trades (user_id, dt, side, market, broker, code, qty, price, value, category, expiry, source_file)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            self.user_id,
                            trade.dt.isoformat(),
                            trade.side,
                            trade.market,
                            trade.broker,
                            trade.code,
                            str(trade.qty),
                            str(trade.price),
                            str(trade.value),
                            trade.category,
                            trade.expiry.isoformat() if trade.expiry else "",
                            str(neg_file),
                        )
                        for trade in self.trades
                    ],
                )
                conn.executemany(
                    """
                    INSERT INTO movements (user_id, dt, direction, kind, product, broker, code, qty, unit_price, value, category, source_file)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            self.user_id,
                            mov.dt.isoformat(),
                            mov.direction,
                            mov.kind,
                            mov.product,
                            mov.broker,
                            mov.code,
                            str(mov.qty),
                            str(mov.unit_price),
                            str(mov.value),
                            mov.category,
                            str(mov_file),
                        )
                        for mov in self.movements
                    ],
                )
        except Exception as exc:
            debug_log(f"Falha ao salvar planilhas importadas no banco: {exc}")

    def _load_cached_imported_data(self) -> None:
        try:
            with db_connect() as conn:
                trade_rows = conn.execute(
                    """
                    SELECT dt, side, market, broker, code, qty, price, value, category, expiry
                    FROM trades
                    WHERE user_id = ?
                    ORDER BY dt, id
                    """,
                    (self.user_id,),
                ).fetchall()
                movement_rows = conn.execute(
                    """
                    SELECT dt, direction, kind, product, broker, code, qty, unit_price, value, category
                    FROM movements
                    WHERE user_id = ?
                    ORDER BY dt, id
                    """,
                    (self.user_id,),
                ).fetchall()
            if not trade_rows and not movement_rows:
                return
            self.trades = [
                Trade(
                    dt=parse_date(row["dt"]) or date.today(),
                    side=row["side"] or "",
                    market=row["market"] or "",
                    broker=row["broker"] or "",
                    code=normalize_ticker(row["code"]),
                    qty=money(row["qty"]),
                    price=money(row["price"]),
                    value=q2(row["value"]),
                    category=refine_asset_category(normalize_ticker(row["code"]), row["category"], row["market"] or ""),
                    expiry=parse_date(row["expiry"]),
                )
                for row in trade_rows
            ]
            self.movements = [
                Movement(
                    dt=parse_date(row["dt"]) or date.today(),
                    direction=row["direction"] or "",
                    kind=row["kind"] or "",
                    product=row["product"] or "",
                    broker=row["broker"] or "",
                    code=normalize_ticker(row["code"]),
                    qty=money(row["qty"]),
                    unit_price=money(row["unit_price"]),
                    value=q2(row["value"]),
                    category=refine_asset_category(normalize_ticker(row["code"]), row["category"], product=row["product"] or ""),
                )
                for row in movement_rows
            ]
            self._refresh_import_preview()
            self.status_var.set(f"Dados carregados do banco: {len(self.trades)} negocios e {len(self.movements)} movimentacoes.")
        except Exception as exc:
            debug_log(f"Falha ao carregar planilhas do banco: {exc}")

    def update_asset_cnpjs_online(self, show_message: bool = True) -> int:
        self.status_var.set("Atualizando tabela de CNPJs pela internet...")
        self.root.configure(cursor="watch")
        self.root.update_idletasks()
        try:
            rows = parse_asset_cnpj_sources()
            with db_connect() as conn:
                conn.execute("DELETE FROM asset_cnpjs WHERE user_id = ?", (self.user_id,))
                conn.executemany(
                    """
                    INSERT INTO asset_cnpjs (user_id, ticker, cnpj, nome, tipo, fonte, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                    """,
                    [(self.user_id, row["ticker"], row["cnpj"], row["nome"], row["tipo"], row["fonte"]) for row in rows],
                )
            msg = f"Tabela de CNPJs atualizada: {len(rows)} ativos."
            self.status_var.set(msg)
            if show_message:
                messagebox.showinfo(APP_TITLE, msg)
            return len(rows)
        except Exception as exc:
            self.status_var.set("Falha ao atualizar CNPJs.")
            debug_log(f"Falha ao atualizar CNPJs: {exc}")
            if show_message:
                messagebox.showerror(APP_TITLE, f"Falha ao atualizar CNPJs:\n{exc}")
            return 0
        finally:
            self.root.configure(cursor="")

    def _asset_cnpj_map(self) -> dict[str, str]:
        try:
            with db_connect() as conn:
                rows = conn.execute("SELECT ticker, cnpj FROM asset_cnpjs WHERE user_id = ?", (self.user_id,)).fetchall()
            return {normalize_ticker(row["ticker"]): row["cnpj"] for row in rows}
        except Exception as exc:
            debug_log(f"Falha ao carregar CNPJs do banco: {exc}")
            return {}

    def import_brokerage_notes(self) -> None:
        paths = filedialog.askopenfilenames(
            title="Selecionar notas de corretagem PDF",
            filetypes=[("PDF", "*.pdf"), ("Todos", "*.*")],
            initialdir=str(BASE_DIR / "notas" if (BASE_DIR / "notas").exists() else Path.home() / "Downloads"),
        )
        if not paths:
            return
        try:
            notes = parse_brokerage_note_pdfs([Path(path) for path in paths])
            if not notes:
                messagebox.showwarning(APP_TITLE, "Nenhuma nota de corretagem reconhecida nos PDFs selecionados.")
                return
            with db_connect() as conn:
                existing_rows = conn.execute(
                    "SELECT id, note_number, trade_date FROM brokerage_note_taxes WHERE user_id = ?",
                    (self.user_id,),
                ).fetchall()
                existing = {(str(row["note_number"] or ""), str(row["trade_date"] or "")): row["id"] for row in existing_rows}
                inserted = 0
                updated = 0
                skipped = 0
                seen_batch: set[tuple[str, str]] = set()
                for note in notes:
                    key = (str(note["note_number"]), note["trade_date"].isoformat())
                    if key in seen_batch:
                        skipped += 1
                        continue
                    seen_batch.add(key)
                    values = (
                        note["broker"],
                        note["source_file"],
                        note["page"],
                        str(note["irrf_common"]),
                        str(note["irrf_daytrade"]),
                        str(note["irrf_fii"]),
                        str(note["irrf_base"]),
                        str(note["buy_normal"]),
                        str(note["sell_normal"]),
                        str(note["buy_fii"]),
                        str(note["sell_fii"]),
                        str(note["buy_options"]),
                        str(note["sell_options"]),
                    )
                    if key in existing:
                        conn.execute(
                            """
                            UPDATE brokerage_note_taxes
                            SET broker = ?, source_file = ?, page = ?, irrf_common = ?, irrf_daytrade = ?, irrf_fii = ?, irrf_base = ?,
                                buy_normal = ?, sell_normal = ?, buy_fii = ?, sell_fii = ?, buy_options = ?, sell_options = ?
                            WHERE id = ? AND user_id = ?
                            """,
                            values + (existing[key], self.user_id),
                        )
                        updated += 1
                    else:
                        conn.execute(
                            """
                            INSERT INTO brokerage_note_taxes
                            (user_id, note_number, trade_date, broker, source_file, page, irrf_common, irrf_daytrade, irrf_fii, irrf_base, buy_normal, sell_normal, buy_fii, sell_fii, buy_options, sell_options)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                            """,
                            (
                                self.user_id,
                                note["note_number"],
                                note["trade_date"].isoformat(),
                                *values,
                            ),
                        )
                        inserted += 1
            self.refresh_brokerage_notes_view()
            msg = f"Notas processadas: {len(notes)}. Novas: {inserted}. Atualizadas: {updated}. Duplicadas no lote: {skipped}."
            self.status_var.set(f"{msg} O calculo principal continua usando as planilhas B3.")
            messagebox.showinfo(APP_TITLE, f"{msg}\nOperacoes nao serao duplicadas; notas servem para IRRF e conferencia.")
        except Exception as exc:
            debug_log(f"Falha ao importar notas de corretagem: {exc}")
            messagebox.showerror(APP_TITLE, f"Falha ao importar notas de corretagem:\n{exc}")

    def clear_brokerage_notes(self) -> None:
        if not messagebox.askyesno(APP_TITLE, "Remover todas as notas de corretagem importadas deste usuario?"):
            return
        with db_connect() as conn:
            conn.execute("DELETE FROM brokerage_note_taxes WHERE user_id = ?", (self.user_id,))
        self.refresh_brokerage_notes_view()
        self.status_var.set("Notas de corretagem removidas.")
        self.refresh_brokerage_compare_view()

    def refresh_brokerage_notes_view(self) -> None:
        if not hasattr(self, "brokerage_notes_tree"):
            return
        for item in self.brokerage_notes_tree.get_children():
            self.brokerage_notes_tree.delete(item)
        try:
            with db_connect() as conn:
                rows = conn.execute(
                    """
                    SELECT trade_date, note_number, irrf_common, irrf_daytrade, irrf_fii, irrf_base, source_file
                    FROM brokerage_note_taxes
                    WHERE user_id = ?
                    ORDER BY trade_date, note_number
                    """,
                    (self.user_id,),
                ).fetchall()
        except Exception as exc:
            debug_log(f"Falha ao listar notas de corretagem: {exc}")
            rows = []
        for row in rows:
            dt = parse_date(row["trade_date"])
            self.brokerage_notes_tree.insert(
                "",
                "end",
                values=[
                    dt.strftime("%d/%m/%Y") if dt else row["trade_date"],
                    row["note_number"],
                    fmt_money(row["irrf_common"]),
                    fmt_money(row["irrf_daytrade"]),
                    fmt_money(row["irrf_fii"]),
                    fmt_money(row["irrf_base"]),
                    Path(row["source_file"]).name,
                ],
            )
        if hasattr(self, "brokerage_notes_status_var"):
            self.brokerage_notes_status_var.set(f"{len(rows)} notas carregadas.")
        self.refresh_brokerage_compare_view()

    def refresh_brokerage_compare_view(self) -> None:
        if not hasattr(self, "brokerage_compare_tree"):
            return
        for item in self.brokerage_compare_tree.get_children():
            self.brokerage_compare_tree.delete(item)
        try:
            selected_year = int(self.selected_year_var.get())
        except Exception:
            selected_year = date.today().year
        note_rows = self._brokerage_monthly_summary(selected_year)
        result = self.results_by_year.get(selected_year)
        for month in range(1, 13):
            note = note_rows.get(month, defaultdict(Decimal))
            b3 = result.monthly[month - 1] if result else MonthlyTax(month=month)
            self.brokerage_compare_tree.insert(
                "",
                "end",
                values=[
                    MONTHS[month - 1],
                    fmt_money(note["sell_normal"]),
                    fmt_money(b3.normal_sales),
                    fmt_money(note["sell_fii"]),
                    fmt_money(b3.fii_sales),
                    fmt_money(note["buy_options"] + note["sell_options"]),
                    fmt_money(abs(b3.options_result)),
                    fmt_money(note["irrf_common"]),
                    fmt_money(note["irrf_daytrade"]),
                    fmt_money(note["irrf_fii"]),
                ],
            )

    def _brokerage_monthly_summary(self, year: int) -> dict[int, defaultdict[str, Decimal]]:
        summary: dict[int, defaultdict[str, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
        try:
            with db_connect() as conn:
                rows = conn.execute(
                    """
                    SELECT trade_date, irrf_common, irrf_daytrade, irrf_fii, buy_normal, sell_normal, buy_fii, sell_fii, buy_options, sell_options
                    FROM brokerage_note_taxes
                    WHERE user_id = ?
                    """,
                    (self.user_id,),
                ).fetchall()
        except Exception as exc:
            debug_log(f"Falha ao resumir notas de corretagem: {exc}")
            return summary
        for row in rows:
            dt = parse_date(row["trade_date"])
            if not dt or dt.year != year:
                continue
            bucket = summary[dt.month]
            for key in ["irrf_common", "irrf_daytrade", "irrf_fii", "buy_normal", "sell_normal", "buy_fii", "sell_fii", "buy_options", "sell_options"]:
                bucket[key] += money(row[key])
        return summary

    def _brokerage_note_taxes(self) -> list[BrokerageNoteTax]:
        try:
            with db_connect() as conn:
                rows = conn.execute(
                    "SELECT trade_date, note_number, source_file, irrf_common, irrf_daytrade, irrf_fii FROM brokerage_note_taxes WHERE user_id = ?",
                    (self.user_id,),
                ).fetchall()
            result: list[BrokerageNoteTax] = []
            for row in rows:
                dt = parse_date(row["trade_date"])
                if dt is None:
                    continue
                result.append(
                    BrokerageNoteTax(
                        dt=dt,
                        irrf_common=money(row["irrf_common"]),
                        irrf_daytrade=money(row["irrf_daytrade"]),
                        irrf_fii=money(row["irrf_fii"]),
                        note_number=str(row["note_number"] or ""),
                        source_file=str(row["source_file"] or ""),
                    )
                )
            return result
        except Exception as exc:
            debug_log(f"Falha ao carregar notas de corretagem: {exc}")
            return []

    def _brokerage_note_files_for_month(self, year: int, month: int) -> list[Path]:
        paths: list[Path] = []
        try:
            with db_connect() as conn:
                rows = conn.execute(
                    """
                    SELECT DISTINCT source_file
                    FROM brokerage_note_taxes
                    WHERE user_id = ? AND trade_date >= ? AND trade_date <= ?
                    """,
                    (self.user_id, f"{year}-{month:02d}-01", f"{year}-{month:02d}-31"),
                ).fetchall()
            for row in rows:
                if not row["source_file"]:
                    continue
                path = Path(row["source_file"])
                if not path.is_absolute():
                    path = BASE_DIR / path
                if path.exists() and path not in paths:
                    paths.append(path)
        except Exception as exc:
            debug_log(f"Falha ao localizar notas do mes: {exc}")
        patterns = [
            f"*{year}_{month:02d}_*.pdf",
            f"*{year}{month:02d}*.pdf",
            f"*{month:02d}-{str(year)[2:]}*.pdf",
        ]
        for folder in [BASE_DIR / "notas", Path.home() / "Downloads"]:
            if not folder.exists():
                continue
            for pattern in patterns:
                for path in folder.glob(pattern):
                    if path.exists() and path not in paths:
                        paths.append(path)
        return paths


    def import_files(self, show_message: bool = False) -> bool:
        debug_log("Inicio da importacao")
        self.status_var.set("Importando planilhas da B3...")
        self.root.configure(cursor="watch")
        self.root.update_idletasks()
        try:
            neg = Path(self.neg_file_var.get())
            mov = Path(self.mov_file_var.get())
            missing = [str(path) for path in (neg, mov) if not path.exists()]
            if missing:
                raise FileNotFoundError("Arquivo nao encontrado:\n" + "\n".join(missing))
            self.trades = read_b3_negotiation(neg) if neg.exists() else []
            self.movements = read_b3_movements(mov) if mov.exists() else []
            category_map = infer_categories_from_movements(self.movements)
            for trade in self.trades:
                if trade.code in category_map and trade.category == "normal":
                    trade.category = category_map[trade.code]
            self._save_imported_data_to_db(neg, mov)
            self._refresh_import_preview()
            msg = f"Importados {len(self.trades)} negocios e {len(self.movements)} movimentacoes."
            self.status_var.set(msg)
            debug_log(msg)
            self._save_config()
            if show_message:
                messagebox.showinfo(APP_TITLE, msg)
            return True
        except Exception as exc:
            self.status_var.set("Falha ao importar planilhas.")
            debug_log(f"Falha ao importar: {exc}")
            messagebox.showerror(APP_TITLE, f"Falha ao importar:\n{exc}")
            return False
        finally:
            self.root.configure(cursor="")

    def _refresh_import_preview(self) -> None:
        for tree in (self.trade_preview, self.movement_preview):
            for item in tree.get_children():
                tree.delete(item)
        for trade in self._filtered_trade_preview_rows()[:300]:
            self.trade_preview.insert(
                "",
                "end",
                values=[
                    trade.dt.strftime("%d/%m/%Y"),
                    trade.side,
                    trade.market,
                    trade.code,
                    fmt_decimal(trade.qty),
                    fmt_money(trade.value),
                    trade.category,
                ],
            )
        for mov in self._filtered_movement_preview_rows()[:300]:
            self.movement_preview.insert(
                "",
                "end",
                values=[
                    mov.dt.strftime("%d/%m/%Y"),
                    mov.kind,
                    mov.code,
                    fmt_decimal(mov.qty),
                    fmt_money(mov.value),
                    mov.category,
                ],
            )

    def _row_matches_filter(self, values: dict[str, str], column: str, text: str) -> bool:
        needle = normalize_header(text)
        if not needle:
            return True
        if column == "todas":
            return any(needle in normalize_header(value) for value in values.values())
        return needle in normalize_header(values.get(column, ""))

    def _filtered_trade_preview_rows(self) -> list[Trade]:
        column = self.trade_filter_col_var.get() if hasattr(self, "trade_filter_col_var") else "todas"
        text = self.trade_filter_text_var.get() if hasattr(self, "trade_filter_text_var") else ""
        reverse = getattr(self, "trade_sort_date_var", tk.StringVar(value="crescente")).get() == "decrescente"
        rows = []
        for trade in self.trades:
            values = {
                "data": trade.dt.strftime("%d/%m/%Y"),
                "tipo": trade.side,
                "mercado": trade.market,
                "codigo": trade.code,
                "quantidade": fmt_decimal(trade.qty),
                "valor": fmt_money(trade.value),
                "categoria": trade.category,
            }
            if self._row_matches_filter(values, column, text):
                rows.append(trade)
        return sorted(rows, key=lambda trade: (trade.dt, trade.code, trade.side), reverse=reverse)

    def _filtered_movement_preview_rows(self) -> list[Movement]:
        column = self.movement_filter_col_var.get() if hasattr(self, "movement_filter_col_var") else "todas"
        text = self.movement_filter_text_var.get() if hasattr(self, "movement_filter_text_var") else ""
        reverse = getattr(self, "movement_sort_date_var", tk.StringVar(value="crescente")).get() == "decrescente"
        rows = []
        for mov in self.movements:
            values = {
                "data": mov.dt.strftime("%d/%m/%Y"),
                "movimentacao": mov.kind,
                "produto": mov.code,
                "quantidade": fmt_decimal(mov.qty),
                "valor": fmt_money(mov.value),
                "categoria": mov.category,
            }
            if self._row_matches_filter(values, column, text):
                rows.append(mov)
        return sorted(rows, key=lambda mov: (mov.dt, mov.code, mov.kind), reverse=reverse)

    def _manual_positions(self) -> list[Position]:
        rows = []
        for row in self.position_table.rows():
            code = normalize_ticker(row["ativo"])
            if not code:
                continue
            rows.append(Position(code=code, qty=money(row["quantidade"]), cost=q2(row["custo"]), category=row["categoria"].strip() or classify_asset(code), broker=row["corretora"]))
        return rows

    def _manual_events(self) -> list[dict[str, Any]]:
        return self.event_table.rows()

    def _manual_data_signature(self) -> str:
        payload = {
            "positions": self.position_table.rows(),
            "events": self.event_table.rows(),
            "losses": {key: var.get() for key, var in self.loss_vars.items()},
            "loss_start_date": self.loss_start_date_var.get(),
        }
        return json.dumps(payload, ensure_ascii=False, sort_keys=True)

    def _event_row_key(self, row: dict[str, Any]) -> tuple[str, str, str, str, str]:
        return (
            str(row.get("data", "")),
            str(row.get("tipo", "")),
            normalize_ticker(row.get("ativo", "")),
            str(row.get("quantidade", "")),
            str(row.get("valor", "")),
        )

    def _load_manual_data_from_db(self) -> None:
        try:
            with db_connect() as conn:
                positions = conn.execute(
                    "SELECT ativo, quantidade, custo, categoria, corretora FROM manual_positions WHERE user_id = ? ORDER BY id",
                    (self.user_id,),
                ).fetchall()
                events = conn.execute(
                    "SELECT data, tipo, ativo, quantidade, valor, categoria, observacao FROM manual_events WHERE user_id = ? ORDER BY id",
                    (self.user_id,),
                ).fetchall()
            for row in positions:
                self.position_table.add_row([row["ativo"], row["quantidade"], row["custo"], row["categoria"], row["corretora"]])
            for row in events:
                self.event_table.add_row([row["data"], row["tipo"], row["ativo"], row["quantidade"], row["valor"], row["categoria"], row["observacao"]])
        except Exception as exc:
            debug_log(f"Falha ao carregar dados manuais do banco: {exc}")

    def _save_manual_data_to_db(self) -> None:
        try:
            with db_connect() as conn:
                conn.execute("DELETE FROM manual_positions WHERE user_id = ?", (self.user_id,))
                conn.execute("DELETE FROM manual_events WHERE user_id = ?", (self.user_id,))
                for row in self.position_table.rows():
                    conn.execute(
                        """
                        INSERT INTO manual_positions (user_id, ativo, quantidade, custo, categoria, corretora)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (self.user_id, row.get("ativo", ""), row.get("quantidade", ""), row.get("custo", ""), row.get("categoria", ""), row.get("corretora", "")),
                    )
                for row in self.event_table.rows():
                    conn.execute(
                        """
                        INSERT INTO manual_events (user_id, data, tipo, ativo, quantidade, valor, categoria, observacao)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            self.user_id,
                            row.get("data", ""),
                            row.get("tipo", ""),
                            row.get("ativo", ""),
                            row.get("quantidade", ""),
                            row.get("valor", ""),
                            row.get("categoria", ""),
                            row.get("observacao", ""),
                        ),
                    )
        except Exception as exc:
            debug_log(f"Falha ao salvar dados manuais no banco: {exc}")

    def _load_option_events(self) -> None:
        existing_keys = {self._event_row_key(row) for row in self.event_table.rows()}
        loaded_from_db = False
        try:
            with db_connect() as conn:
                rows = conn.execute(
                    "SELECT data, tipo, ativo, quantidade, valor, categoria, observacao FROM option_conferences WHERE user_id = ? ORDER BY id",
                    (self.user_id,),
                ).fetchall()
                for row in rows:
                    row_dict = dict(row)
                    if self._event_row_key(row_dict) not in existing_keys:
                        self.event_table.add_row([row["data"], row["tipo"], row["ativo"], row["quantidade"], row["valor"], row["categoria"], row["observacao"]])
                        existing_keys.add(self._event_row_key(row_dict))
                loaded_from_db = bool(rows)
        except Exception as exc:
            debug_log(f"Falha ao carregar conferencias do banco: {exc}")
        if loaded_from_db:
            self._refresh_option_events()
            return
        if self.username != "local":
            return
        if not OPTION_EVENTS_FILE.exists():
            return
        try:
            rows = json.loads(OPTION_EVENTS_FILE.read_text(encoding="utf-8"))
        except Exception as exc:
            debug_log(f"Falha ao carregar conferencias de opcoes: {exc}")
            return
        for row in rows:
            if self._event_row_key(row) in existing_keys:
                continue
            self.event_table.add_row([row.get("data", ""), row.get("tipo", ""), row.get("ativo", ""), row.get("quantidade", ""), row.get("valor", ""), row.get("categoria", "opcoes"), row.get("observacao", "Carregado de conferencias_opcoes.json")])
            existing_keys.add(self._event_row_key(row))
        self._refresh_option_events()
        self._save_option_events(suppress_status=True)

    def _save_option_events(self, suppress_status: bool = False) -> None:
        rows = []
        for row in self.event_table.rows():
            event_type = normalize_header(row.get("tipo"))
            if any(marker in event_type for marker in ["exercicio", "virou_po", "po", "ainda_ativa", "vendida_ativa", "ainda_em_aberto", "em_aberto", "recompra", "recomprada", "liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao"]):
                rows.append(row)
        if self.username == "local":
            OPTION_EVENTS_FILE.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
        try:
            with db_connect() as conn:
                conn.execute("DELETE FROM option_conferences WHERE user_id = ?", (self.user_id,))
                for row in rows:
                    conn.execute(
                        """
                        INSERT INTO option_conferences (user_id, data, tipo, ativo, quantidade, valor, categoria, observacao)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            self.user_id,
                            row.get("data", ""),
                            row.get("tipo", ""),
                            row.get("ativo", ""),
                            row.get("quantidade", ""),
                            row.get("valor", ""),
                            row.get("categoria", "opcoes"),
                            row.get("observacao", ""),
                        ),
                    )
        except Exception as exc:
            debug_log(f"Falha ao salvar conferencias no banco: {exc}")
        self._save_manual_data_to_db()
        if not suppress_status and self.username == "local":
            self.status_var.set(f"Conferencias salvas em {OPTION_EVENTS_FILE}.")
        elif not suppress_status:
            self.status_var.set("Conferencias salvas no banco de dados.")

    def _sync_option_events_to_manual_table(self) -> None:
        option_keys = set()
        for item in self.option_events_tree.get_children():
            values = self.option_events_tree.item(item, "values")
            option_keys.add(tuple(values[:5]))
        keep_rows = []
        for row in self.event_table.rows():
            event_type = normalize_header(row.get("tipo"))
            is_option_resolution = any(marker in event_type for marker in ["exercicio", "virou_po", "po", "ainda_ativa", "vendida_ativa", "ainda_em_aberto", "em_aberto", "recompra", "recomprada", "liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao"])
            key = (row.get("data", ""), row.get("tipo", ""), row.get("ativo", ""), row.get("quantidade", ""), row.get("valor", ""))
            if not is_option_resolution or key in option_keys:
                keep_rows.append(row)
        self.event_table.clear()
        for row in keep_rows:
            self.event_table.add_row([row.get("data", ""), row.get("tipo", ""), row.get("ativo", ""), row.get("quantidade", ""), row.get("valor", ""), row.get("categoria", ""), row.get("observacao", "")])

    def _edit_option_event(self) -> None:
        selected = self.option_events_tree.selection()
        if not selected:
            messagebox.showinfo(APP_TITLE, "Selecione uma conferencia para editar.")
            return
        item = selected[0]
        values = list(self.option_events_tree.item(item, "values"))
        columns = [("data", "Data", 90), ("tipo", "Tipo", 150), ("ativo", "Opcao", 100), ("quantidade", "Quantidade", 110), ("valor", "Valor", 100), ("observacao", "Observacao", 300)]
        dialog = RowEditor(self.root, columns, values)
        self.root.wait_window(dialog)
        if dialog.result is None:
            return
        self.option_events_tree.item(item, values=dialog.result)
        self._sync_option_events_to_manual_table()
        self._save_option_events()

    def _delete_option_event(self) -> None:
        selected = self.option_events_tree.selection()
        if not selected:
            messagebox.showinfo(APP_TITLE, "Selecione uma conferencia para excluir.")
            return
        for item in selected:
            self.option_events_tree.delete(item)
        self._sync_option_events_to_manual_table()
        self._save_option_events()

    def _add_option_resolution_event(self, event_type: str) -> None:
        selected = self.pending_options_tree.selection()
        if not selected:
            messagebox.showinfo(APP_TITLE, "Selecione uma opcao pendente na aba Opcoes a conferir.")
            return
        values = self.pending_options_tree.item(selected[0], "values")
        if len(values) < 5:
            return
        year, code, expiry, qty, premium = values[:5]
        qty = fmt_decimal(abs(money(qty)))
        if "ainda ativa" in event_type or "em aberto" in event_type:
            event_date = f"31/12/{year}" if str(year).isdigit() else date.today().strftime("%d/%m/%Y")
        else:
            event_date = expiry or date.today().strftime("%d/%m/%Y")
        event_value = premium if event_type == "opcao virou po" else "0,00"
        if "ainda ativa" in event_type or "em aberto" in event_type:
            event_value = premium
        self.event_table.add_row([event_date, event_type, code, qty, event_value, "opcoes", "Lancado pela aba Opcoes a conferir; recalcule para atualizar."])
        self.pending_options_tree.set(selected[0], "situacao", f"Registrada como '{event_type}'. Recalcule para atualizar.")
        self.pending_options_tree.item(selected[0], tags=("registered",))
        self._refresh_option_events()
        self._save_option_events()
        self.status_var.set(f"Evento '{event_type}' incluido para {code}. Clique em Recalcular agora.")

    def _add_option_repurchase_event(self) -> None:
        selected = self.pending_options_tree.selection()
        if not selected:
            messagebox.showinfo(APP_TITLE, "Selecione uma opcao pendente na aba Opcoes a conferir.")
            return
        values = self.pending_options_tree.item(selected[0], "values")
        if len(values) < 5:
            return
        year, code, expiry, qty, _premium = values[:5]
        qty = fmt_decimal(abs(money(qty)))
        columns = [
            ("data", "Data da recompra", 90),
            ("tipo", "Tipo", 150),
            ("ativo", "Opcao", 100),
            ("quantidade", "Quantidade recomprada", 120),
            ("valor", "Valor pago na recompra", 130),
            ("observacao", "Observacao", 300),
        ]
        defaults = [expiry or f"31/12/{year}", "opcao recomprada", code, qty, "0,00", "Recompra informada pela aba Opcoes a conferir; recalcule para atualizar."]
        dialog = RowEditor(self.root, columns, defaults)
        self.root.wait_window(dialog)
        if dialog.result is None:
            return
        data, tipo, ativo, quantidade, valor, observacao = dialog.result
        self.event_table.add_row([data, tipo or "opcao recomprada", ativo or code, quantidade or qty, valor or "0,00", "opcoes", observacao])
        self.pending_options_tree.set(selected[0], "situacao", "Registrada como recomprada. Recalcule para atualizar.")
        self.pending_options_tree.item(selected[0], tags=("registered",))
        self._refresh_option_events()
        self._save_option_events()
        self.status_var.set(f"Recompra registrada para {ativo or code}. Clique em Recalcular agora.")

    def _add_option_spread_settlement_event(self) -> None:
        selected = self.pending_options_tree.selection()
        values = self.pending_options_tree.item(selected[0], "values") if selected else []
        year = values[0] if values else self.selected_year_var.get()
        code = values[1] if values else ""
        qty = values[3] if len(values) > 3 else ""
        expiry = values[2] if len(values) > 2 else ""
        columns = [
            ("data", "Data da liquidacao", 90),
            ("tipo", "Tipo", 190),
            ("ativo", "Opcao exercida/linha B3", 140),
            ("quantidade", "Quantidade", 120),
            ("valor", "Valor financeiro", 130),
            ("observacao", "Observacao", 420),
        ]
        defaults = [
            expiry or f"31/12/{year}",
            "liquidacao trava sem exercicio acao",
            normalize_option_exercise_code(code),
            fmt_decimal(abs(money(qty))) if qty else "",
            "0,00",
            "Exercicio informado pela B3 foi liquidacao financeira/recompra de trava; nao movimentar a acao-base.",
        ]
        dialog = RowEditor(self.root, columns, defaults)
        self.root.wait_window(dialog)
        if dialog.result is None:
            return
        data, tipo, ativo, quantidade, valor, observacao = dialog.result
        self.event_table.add_row([data, tipo or "liquidacao trava sem exercicio acao", ativo or code, quantidade or qty, valor or "0,00", "opcoes", observacao])
        if selected:
            self.pending_options_tree.set(selected[0], "situacao", "Registrada como liquidacao de trava sem acao. Recalcule para atualizar.")
            self.pending_options_tree.item(selected[0], tags=("registered",))
        self._refresh_option_events()
        self._save_option_events()
        self.status_var.set(f"Liquidacao de trava sem acao registrada para {ativo or code}. Clique em Recalcular agora.")

    def calculate(self) -> None:
        if self.calculating:
            return
        debug_log("Botao Calcular acionado")
        if not self.trades and not self.movements:
            if not self.import_files():
                return
        try:
            start_year = int(self.start_year_var.get())
            end_year = int(self.end_year_var.get())
        except Exception:
            messagebox.showerror(APP_TITLE, "Informe anos validos.")
            return
        if end_year < start_year:
            messagebox.showerror(APP_TITLE, "O ano-calendario final deve ser maior ou igual ao ano inicial.")
            return
        if end_year - start_year > 80:
            messagebox.showerror(APP_TITLE, "Intervalo de anos muito grande. Confira o ano inicial e o ano final.")
            return
        losses = {key: money(var.get()) for key, var in self.loss_vars.items()}
        positions = self._manual_positions()
        events = self._manual_events()
        parsed_loss_base_date = self._normalize_loss_date_field(show_error=True)
        if self.loss_start_date_var.get().strip() and parsed_loss_base_date is None:
            return
        self._set_calculating(True)
        try:
            debug_log(f"Inicio do calculo: {start_year}-{end_year}; negocios={len(self.trades)}; movimentacoes={len(self.movements)}")
            results: dict[int, CalculationResult] = {}
            current_positions = positions
            current_losses = {key: Decimal("0") for key in ["normal", "daytrade", "fii"]}
            loss_start_date = self._loss_effect_date(parsed_loss_base_date)
            manual_losses_applied = False
            asset_cnpjs = self._asset_cnpj_map()
            brokerage_taxes = self._brokerage_note_taxes()
            for year in range(start_year, end_year + 1):
                self.status_var.set(f"Calculando ano-calendario {year}...")
                self.root.update_idletasks()
                debug_log(f"Calculando ano {year}")
                initial_losses = dict(current_losses)
                dated_losses: dict[str, Decimal] = {}
                dated_loss_start: date | None = None
                if not manual_losses_applied:
                    if loss_start_date is None or loss_start_date < date(year, 1, 1):
                        initial_losses = self._declared_loss_dict(losses)
                        manual_losses_applied = True
                    elif loss_start_date.year == year:
                        dated_losses = self._declared_loss_dict(losses)
                        dated_loss_start = loss_start_date
                        manual_losses_applied = True
                engine = IRSimpleEngine(year, asset_cnpjs)
                result = engine.calculate(
                    self.trades,
                    self.movements,
                    current_positions,
                    initial_losses,
                    self._events_for_year(events, year, start_year),
                    dated_losses,
                    dated_loss_start,
                    brokerage_taxes,
                )
                debug_log(f"Ano {year} calculado: posicoes={len(result.positions)}; alertas={len(result.warnings)}")
                results[year] = result
                current_positions = self._carry_positions(result)
                current_losses = self._carry_losses(result)
            self._finish_calculation(results, start_year, end_year)
        except Exception as exc:
            debug_log(f"Falha no calculo: {exc}")
            self._fail_calculation(exc)

    def _declared_loss_dict(self, values: dict[str, Decimal]) -> dict[str, Decimal]:
        return {key: q2(abs(money(values.get(key)))) for key in ["normal", "daytrade", "fii"]}

    def _loss_effect_date(self, base_date: date | None) -> date | None:
        if base_date is None:
            return None
        if base_date.day == 1:
            return base_date
        year = base_date.year + 1 if base_date.month == 12 else base_date.year
        month = 1 if base_date.month == 12 else base_date.month + 1
        return date(year, month, 1)

    def _set_calculating(self, calculating: bool) -> None:
        self.calculating = calculating
        state = "disabled" if calculating else "normal"
        self.calculate_button.configure(state=state)
        self.export_excel_button.configure(state=state)
        self.export_pdf_button.configure(state=state)
        self.root.configure(cursor="watch" if calculating else "")
        if calculating:
            self.status_var.set("Iniciando calculo...")
        self.root.update_idletasks()

    def _finish_calculation(self, results: dict[int, CalculationResult], start_year: int, end_year: int) -> None:
        debug_log("Atualizando telas com resultados")
        self.results_by_year = results
        self.last_calculation_manual_signature = self._manual_data_signature()
        self.result = self.results_by_year[end_year]
        years = [str(year) for year in self.results_by_year]
        self.year_combo.configure(values=years)
        if self.selected_year_var.get() not in years:
            self.selected_year_var.set(str(end_year))
        self._save_config()
        self._refresh_results()
        self.refresh_brokerage_compare_view()
        self.status_var.set(f"Calculo concluido de {start_year} ate {end_year} (ultima declaracao: exercicio {end_year + 1}).")
        self._set_calculating(False)
        self.root.update_idletasks()
        debug_log("Calculo concluido")

    def _fail_calculation(self, exc: Exception) -> None:
        self._set_calculating(False)
        self.status_var.set("Falha no calculo.")
        messagebox.showerror(APP_TITLE, f"Falha no calculo:\n{exc}")

    def _events_for_year(self, events: list[dict[str, Any]], year: int, start_year: int) -> list[dict[str, Any]]:
        selected = []
        for event in events:
            event_date = parse_date(event.get("data"))
            event_type = normalize_header(event.get("tipo"))
            is_position_adjustment = any(
                marker in event_type
                for marker in ["transferencia", "migracao", "bonificacao", "subscricao", "desdobramento", "grupamento", "split"]
            )
            if event_date is None and year == start_year:
                selected.append(event)
            elif event_date is not None and event_date.year == year:
                selected.append(event)
        return selected

    def _carry_positions(self, result: CalculationResult) -> list[Position]:
        positions = []
        for pos in result.positions.values():
            if pos.qty <= 0:
                continue
            positions.append(
                Position(
                    code=pos.code,
                    qty=pos.qty,
                    cost=q2(pos.cost),
                    previous_qty=pos.qty,
                    previous_cost=q2(pos.cost),
                    category=pos.category,
                    broker=pos.broker,
                )
            )
        return positions

    def _carry_losses(self, result: CalculationResult) -> dict[str, Decimal]:
        last = result.monthly[-1]
        return {
            "normal": last.normal_loss_after,
            "daytrade": last.daytrade_loss_after,
            "fii": last.fii_loss_after,
        }

    def _portfolio_cost(self, positions: dict[str, Position]) -> Decimal:
        return q2(
            sum(
                (pos.cost for pos in positions.values() if pos.qty > 0 and pos.category not in {"opcoes", "futuro"}),
                Decimal("0"),
            )
        )

    def _month_profit(self, item: MonthlyTax) -> Decimal:
        return q2(item.normal_result + item.daytrade_result + item.fii_result + item.options_result + item.future_result)

    def _month_stock_profit(self, item: MonthlyTax) -> Decimal:
        return q2(item.normal_result)

    def _month_fund_profit(self, item: MonthlyTax) -> Decimal:
        return q2(item.fii_result)

    def _month_other_profit(self, item: MonthlyTax) -> Decimal:
        return q2(item.daytrade_result + item.options_result + item.future_result)

    def _benchmark_period_range(self) -> tuple[date, date] | None:
        if not self.results_by_year:
            return None
        start_year = min(self.results_by_year)
        end_year = max(self.results_by_year)
        return date(start_year, 1, 1), date(end_year, 12, 31)

    def _ensure_benchmarks(self, force: bool = False) -> None:
        period_range = self._benchmark_period_range()
        if period_range is None:
            return
        start, end = period_range
        start_key = month_key(start)
        end_key = month_key(date(end.year, end.month, 1))
        current = set(self.benchmark_cache.get("selic", {})) & set(self.benchmark_cache.get("poupanca", {})) & set(self.benchmark_cache.get("ibovespa", {}))
        if not force and start_key in current and current:
            return
        try:
            selic = fetch_bcb_monthly_percent(4390, start, end)
            poupanca = fetch_bcb_monthly_percent(195, start, end)
            ibov = fetch_ibovespa_monthly_close(start, end)
            self.benchmark_cache = {
                "selic": {key: str(q2(value)) for key, value in selic.items()},
                "poupanca": {key: str(q2(value)) for key, value in poupanca.items()},
                "ibovespa": {key: str(q2(value)) for key, value in ibov.items()},
            }
            self._save_config()
        except Exception as exc:
            debug_log(f"Falha ao atualizar indices de comparacao: {exc}")
            if force:
                messagebox.showerror(APP_TITLE, f"Falha ao buscar indices:\n{exc}")

    def _update_benchmarks_from_web(self) -> None:
        if not self.results_by_year:
            messagebox.showinfo(APP_TITLE, "Clique em Calcular antes de atualizar os indices.")
            return
        self.status_var.set("Atualizando Ibovespa, Selic e Poupanca...")
        self.root.update_idletasks()
        self._ensure_benchmarks(force=True)
        self._refresh_wealth_view()
        self.status_var.set("Indices atualizados para comparacao patrimonial.")

    def _benchmark_values_for_rows(self, rows: list[tuple[str, Decimal, Decimal, Decimal, Decimal, Decimal, Decimal, str]]) -> dict[str, dict[str, Decimal]]:
        if not rows:
            return {"ibovespa": {}, "selic": {}, "poupanca": {}}
        self._ensure_benchmarks(force=False)
        base = next((row[1] for row in rows if row[1] > 0), Decimal("100"))
        if base <= 0:
            base = Decimal("100")
        monthly_keys = [f"{year}-{month:02d}" for year in self.results_by_year for month in range(1, 13)]
        if not monthly_keys:
            return {"ibovespa": {}, "selic": {}, "poupanca": {}}
        first_key = min(monthly_keys)

        selic_monthly = {key: money(value) for key, value in self.benchmark_cache.get("selic", {}).items()}
        poupanca_monthly = {key: money(value) for key, value in self.benchmark_cache.get("poupanca", {}).items()}
        ibov_monthly = {key: money(value) for key, value in self.benchmark_cache.get("ibovespa", {}).items()}

        factors: dict[str, dict[str, Decimal]] = {"selic": {}, "poupanca": {}, "ibovespa": {}}
        selic_value = base
        poupanca_value = base
        first_ibov = next((ibov_monthly[key] for key in sorted(ibov_monthly) if key >= first_key and ibov_monthly[key] > 0), Decimal("0"))
        for key in sorted(monthly_keys):
            if key in selic_monthly:
                selic_value = q2(selic_value * (Decimal("1") + selic_monthly[key] / Decimal("100")))
            if key in poupanca_monthly:
                poupanca_value = q2(poupanca_value * (Decimal("1") + poupanca_monthly[key] / Decimal("100")))
            factors["selic"][key] = selic_value
            factors["poupanca"][key] = poupanca_value
            if first_ibov > 0 and key in ibov_monthly and ibov_monthly[key] > 0:
                factors["ibovespa"][key] = q2(base * ibov_monthly[key] / first_ibov)

        selected: dict[str, dict[str, Decimal]] = {"ibovespa": {}, "selic": {}, "poupanca": {}}
        for period, *_rest in rows:
            parsed = parse_month_key(period)
            if parsed is None:
                continue
            year, month = parsed
            key = f"{year}-{month:02d}"
            for name in selected:
                selected[name][period] = factors[name].get(key, Decimal("0"))
        return selected

    def _selected_result(self) -> CalculationResult | None:
        try:
            year = int(self.selected_year_var.get())
        except ValueError:
            return self.result
        return self.results_by_year.get(year, self.result)

    def _on_history_year_selected(self, _event: tk.Event) -> None:
        if self.updating_history:
            return
        selected = self.history_summary_tree.selection()
        if not selected:
            return
        year = selected[0]
        if year == self.selected_year_var.get():
            return
        self.selected_year_var.set(year)
        self._refresh_results()

    def _result_summary(self, result: CalculationResult) -> dict[str, Decimal]:
        return {
            "normal": q2(sum((m.normal_result for m in result.monthly), Decimal("0"))),
            "daytrade": q2(sum((m.daytrade_result for m in result.monthly), Decimal("0"))),
            "fii": q2(sum((m.fii_result for m in result.monthly), Decimal("0"))),
            "opcoes": q2(sum((m.options_result for m in result.monthly), Decimal("0"))),
            "futuro": q2(sum((m.future_result for m in result.monthly), Decimal("0"))),
            "imposto": q2(sum((m.tax_due for m in result.monthly), Decimal("0"))),
        }

    def _losses_text(self, monthly: MonthlyTax) -> str:
        values = [
            ("N", monthly.normal_loss_after),
            ("DT", monthly.daytrade_loss_after),
            ("FII", monthly.fii_loss_after),
        ]
        return " | ".join(f"{label}: {fmt_money(value)}" for label, value in values if value > 0) or "0,00"

    def _refresh_history(self, selected_year: int) -> None:
        self.updating_history = True
        try:
            for tree in (self.history_positions_tree, self.history_month_tree, self.history_loans_tree):
                for item in tree.get_children():
                    tree.delete(item)
            existing_years = set(self.history_summary_tree.get_children())
            target_years = {str(year) for year in self.results_by_year}
            if existing_years != target_years:
                for item in self.history_summary_tree.get_children():
                    self.history_summary_tree.delete(item)
                for year, result in self.results_by_year.items():
                    summary = self._result_summary(result)
                    active_positions = [pos for pos in result.positions.values() if pos.qty > 0]
                    portfolio_value = q2(sum((pos.cost for pos in active_positions), Decimal("0")))
                    final_losses = self._losses_text(result.monthly[-1])
                    self.history_summary_tree.insert(
                        "",
                        "end",
                        iid=str(year),
                        values=[
                            year,
                            year + 1,
                            len(active_positions),
                            fmt_money(portfolio_value),
                            fmt_money(summary["normal"]),
                            fmt_money(summary["daytrade"]),
                            fmt_money(summary["fii"]),
                            fmt_money(summary["opcoes"]),
                            fmt_money(summary["futuro"]),
                            fmt_money(summary["imposto"]),
                            final_losses,
                        ],
                    )
            if str(selected_year) in self.history_summary_tree.get_children():
                selected_key = str(selected_year)
                if self.history_summary_tree.selection() != (selected_key,):
                    self.history_summary_tree.selection_set(selected_key)
                self.history_summary_tree.focus(selected_key)
                self.history_summary_tree.see(str(selected_year))

            result = self.results_by_year.get(selected_year)
            if result is None:
                return
            for code, pos in sorted(result.positions.items()):
                if pos.qty <= 0:
                    continue
                self.history_positions_tree.insert(
                    "",
                    "end",
                    values=[
                        code,
                        pos.category,
                        fmt_decimal(pos.qty),
                        fmt_money(pos.avg_price),
                        fmt_money(pos.previous_cost),
                        fmt_money(pos.cost),
                    ],
                )
            for item in result.monthly:
                self.history_month_tree.insert(
                    "",
                    "end",
                    values=[
                        MONTHS[item.month - 1],
                        fmt_money(item.normal_result),
                        fmt_money(item.daytrade_result),
                        fmt_money(item.fii_result),
                        fmt_money(item.options_result),
                        fmt_money(item.future_result),
                        fmt_money(item.tax_payable),
                        self._losses_text(item),
                    ],
                )
            for month in range(1, 13):
                loans = result.monthly_loans.get(month, {})
                for code, qty in sorted(loans.items()):
                    self.history_loans_tree.insert("", "end", values=[MONTHS[month - 1], code, fmt_decimal(qty)])
        finally:
            self.updating_history = False

    def _wealth_rows(self) -> list[tuple[str, Decimal, Decimal, Decimal, Decimal, Decimal, Decimal, str]]:
        rows: list[tuple[str, Decimal, Decimal, Decimal, Decimal, Decimal, Decimal, str]] = []
        if self.wealth_view_var.get() == "mensal":
            for year, result in sorted(self.results_by_year.items()):
                last_value = Decimal("0")
                for item in result.monthly:
                    positions = result.monthly_positions.get(item.month, {})
                    if positions:
                        last_value = self._portfolio_cost(positions)
                    rows.append(
                        (
                            f"{year}-{item.month:02d}",
                            last_value,
                            self._month_stock_profit(item),
                            self._month_fund_profit(item),
                            self._month_other_profit(item),
                            self._month_profit(item),
                            q2(item.tax_payable),
                            self._losses_text(item),
                        )
                    )
            return rows
        for year, result in sorted(self.results_by_year.items()):
            summary = self._result_summary(result)
            rows.append(
                (
                    str(year),
                    self._portfolio_cost(result.positions),
                    q2(summary["normal"]),
                    q2(summary["fii"]),
                    q2(summary["daytrade"] + summary["opcoes"] + summary["futuro"]),
                    q2(summary["normal"] + summary["daytrade"] + summary["fii"] + summary["opcoes"] + summary["futuro"]),
                    q2(sum((m.tax_payable for m in result.monthly), Decimal("0"))),
                    self._losses_text(result.monthly[-1]),
                )
            )
        return rows

    def _refresh_wealth_view(self) -> None:
        if not hasattr(self, "wealth_tree"):
            return
        for item in self.wealth_tree.get_children():
            self.wealth_tree.delete(item)
        rows = self._wealth_rows()
        benchmarks = self._benchmark_values_for_rows(rows)
        self.wealth_points = [(period, value, total, tax) for period, value, _stock, _fund, _other, total, tax, _losses in rows]
        self.wealth_benchmark_points = {
            name: [(period, values.get(period, Decimal("0"))) for period, *_rest in rows]
            for name, values in benchmarks.items()
        }
        for period, value, stock, fund, other, total, tax, losses in rows:
            self.wealth_tree.insert(
                "",
                "end",
                values=[
                    period,
                    fmt_money(value),
                    fmt_money(benchmarks["ibovespa"].get(period, 0)) if benchmarks["ibovespa"].get(period, 0) else "",
                    fmt_money(benchmarks["selic"].get(period, 0)) if benchmarks["selic"].get(period, 0) else "",
                    fmt_money(benchmarks["poupanca"].get(period, 0)) if benchmarks["poupanca"].get(period, 0) else "",
                    fmt_money(stock),
                    fmt_money(fund),
                    fmt_money(other),
                    fmt_money(total),
                    fmt_money(tax),
                    losses,
                ],
            )
        self._draw_wealth_chart()

    def _detail_subtotals(self, rows: list[list[str]]) -> dict[str, Decimal]:
        totals: dict[str, Decimal] = defaultdict(Decimal)
        for row in rows:
            if len(row) < 13:
                continue
            source = str(row[0])
            kind = str(row[2])
            profit = money(row[12])
            if not profit:
                continue
            if "conferencia" in source:
                totals["diferenca/conferencia"] += profit
            elif "vencimento" in kind or "virou po" in kind:
                totals["suposicoes de vencimento"] += profit
            elif "recompra" in kind:
                totals["recompras/fechamentos"] += profit
            elif "venda de opcao comprada" in kind:
                totals["vendas de opcoes compradas"] += profit
            elif "compra/venda" in kind or "day trade" in source:
                totals["day trade"] += profit
            else:
                totals["operacoes realizadas"] += profit
        return {key: q2(value) for key, value in totals.items()}

    def _draw_wealth_chart(self) -> None:
        if not hasattr(self, "wealth_canvas"):
            return
        canvas = self.wealth_canvas
        canvas.delete("all")
        width = max(canvas.winfo_width(), 300)
        height = max(canvas.winfo_height(), 220)
        margin_left = 82
        margin_right = 24
        margin_top = 24
        margin_bottom = 54
        plot_w = width - margin_left - margin_right
        plot_h = height - margin_top - margin_bottom
        canvas.create_rectangle(margin_left, margin_top, width - margin_right, height - margin_bottom, outline="#cbd5e1", fill="#ffffff")
        if not self.wealth_points:
            canvas.create_text(width / 2, height / 2, text="Clique em Calcular para gerar a evolucao patrimonial.", fill="#475569")
            return

        benchmark_series = []
        if getattr(self, "show_ibov_var", None) is not None and self.show_ibov_var.get():
            benchmark_series.append(("Ibovespa", "#f97316", self.wealth_benchmark_points.get("ibovespa", [])))
        if getattr(self, "show_selic_var", None) is not None and self.show_selic_var.get():
            benchmark_series.append(("Selic", "#16a34a", self.wealth_benchmark_points.get("selic", [])))
        if getattr(self, "show_poupanca_var", None) is not None and self.show_poupanca_var.get():
            benchmark_series.append(("Poupanca", "#9333ea", self.wealth_benchmark_points.get("poupanca", [])))
        values = [point[1] for point in self.wealth_points]
        for _name, _color, series in benchmark_series:
            values.extend(value for _period, value in series if value > 0)
        max_value = max(values)
        min_value = min(values)
        if max_value == min_value:
            max_value += Decimal("1")
            min_value = Decimal("0") if min_value >= 0 else min_value - Decimal("1")
        span = max_value - min_value

        for idx in range(5):
            ratio = Decimal(idx) / Decimal("4")
            y = margin_top + float(ratio) * plot_h
            value = max_value - span * ratio
            canvas.create_line(margin_left, y, width - margin_right, y, fill="#e2e8f0")
            canvas.create_text(margin_left - 8, y, text=fmt_money(value), anchor="e", fill="#475569", font=("TkDefaultFont", 8))

        def scaled_points(series: list[tuple[str, Decimal]]) -> list[tuple[float, float]]:
            result: list[tuple[float, float]] = []
            total_points = len(self.wealth_points)
            period_index = {period: idx for idx, (period, *_rest) in enumerate(self.wealth_points)}
            for period, value in series:
                if value <= 0 or period not in period_index:
                    continue
                idx = period_index[period]
                x = margin_left + (plot_w / max(total_points - 1, 1)) * idx
                y = margin_top + float((max_value - value) / span) * plot_h
                result.append((x, y))
            return result

        points: list[tuple[float, float]] = []
        total = len(self.wealth_points)
        for idx, (_period, value, _profit, _tax) in enumerate(self.wealth_points):
            x = margin_left + (plot_w / max(total - 1, 1)) * idx
            y = margin_top + float((max_value - value) / span) * plot_h
            points.append((x, y))
        if len(points) > 1:
            canvas.create_line(*[coord for point in points for coord in point], fill="#2563eb", width=3, smooth=False)
        for name, color, series in benchmark_series:
            chart_points = scaled_points(series)
            if len(chart_points) > 1:
                canvas.create_line(*[coord for point in chart_points for coord in point], fill=color, width=2, dash=(5, 3), smooth=False)
        for idx, (x, y) in enumerate(points):
            period, value, _profit, _tax = self.wealth_points[idx]
            canvas.create_oval(x - 3, y - 3, x + 3, y + 3, fill="#1d4ed8", outline="#1d4ed8")
            if idx == 0 or idx == total - 1 or total <= 12 or idx % max(total // 10, 1) == 0:
                canvas.create_text(x, height - margin_bottom + 18, text=period, anchor="n", fill="#475569", font=("TkDefaultFont", 8))
        last_period, last_value, _last_profit, _last_tax = self.wealth_points[-1]
        canvas.create_text(
            margin_left,
            8,
            text=f"Ultimo ponto: {last_period} | Patrimonio: R$ {fmt_money(last_value)}",
            anchor="nw",
            fill="#0f172a",
        )
        legend_x = width - margin_right - 260
        legend_y = 10
        legend = [("Patrimonio", "#2563eb", None)] + [(name, color, None) for name, color, _series in benchmark_series]
        for idx, (name, color, _dash) in enumerate(legend):
            y = legend_y + idx * 16
            if name == "Patrimonio":
                canvas.create_line(legend_x, y + 6, legend_x + 28, y + 6, fill=color, width=3)
            else:
                canvas.create_line(legend_x, y + 6, legend_x + 28, y + 6, fill=color, width=3, dash=(5, 3))
            canvas.create_text(legend_x + 34, y, text=name, anchor="nw", fill="#0f172a", font=("TkDefaultFont", 8))

    def _refresh_pending_options(self) -> None:
        for item in self.pending_options_tree.get_children():
            self.pending_options_tree.delete(item)
        for year, result in self.results_by_year.items():
            for row in result.pending_options:
                if self._option_has_saved_conference(row.get("codigo", "")):
                    continue
                expiry = row.get("vencimento")
                last_date = row.get("ultima_data")
                qty = money(row.get("quantidade_aberta", 0))
                if money(row.get("quantidade_vendida", 0)) > money(row.get("quantidade_comprada", 0)):
                    qty = -qty
                self.pending_options_tree.insert(
                    "",
                    "end",
                    values=[
                        year,
                        row.get("codigo", ""),
                        expiry.strftime("%d/%m/%Y") if isinstance(expiry, date) else "",
                        fmt_decimal(qty),
                        fmt_money(row.get("premio_liquido", 0)),
                        last_date.strftime("%d/%m/%Y") if isinstance(last_date, date) else "",
                        row.get("situacao", ""),
                    ],
                )
        self._refresh_option_events()

    def _option_has_saved_conference(self, code: str) -> bool:
        code = normalize_ticker(code)
        for row in self.event_table.rows():
            if normalize_ticker(row.get("ativo")) == code:
                event_type = normalize_header(row.get("tipo"))
                if any(marker in event_type for marker in ["exercicio", "virou_po", "po", "ainda_ativa", "vendida_ativa", "ainda_em_aberto", "em_aberto", "recompra", "recomprada", "liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao"]):
                    return True
        return False

    def _refresh_option_events(self) -> None:
        for item in self.option_events_tree.get_children():
            self.option_events_tree.delete(item)
        for row in self.event_table.rows():
            event_type = normalize_header(row.get("tipo"))
            if not any(marker in event_type for marker in ["exercicio", "virou_po", "po", "ainda_ativa", "vendida_ativa", "ainda_em_aberto", "em_aberto", "recompra", "recomprada", "liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao"]):
                continue
            self.option_events_tree.insert(
                "",
                "end",
                values=[
                    row.get("data", ""),
                    row.get("tipo", ""),
                    row.get("ativo", ""),
                    row.get("quantidade", ""),
                    row.get("valor", ""),
                    row.get("observacao", ""),
                ],
            )

    def _refresh_b3_check(self) -> None:
        for item in self.b3_check_tree.get_children():
            self.b3_check_tree.delete(item)
        if not self.results_by_year:
            self.b3_check_status_var.set("Calcule antes de conferir.")
            return
        calculated_signature = getattr(self, "last_calculation_manual_signature", None)
        current_signature = self._manual_data_signature()
        if calculated_signature is not None and calculated_signature != current_signature:
            self.b3_check_status_var.set("Dados manuais mudaram. Clique em Calcular antes de conferir os relatorios B3.")
            return
        files = [Path(path) for path in self.consolidated_files if Path(path).exists()]
        if not files:
            files = sorted((Path.home() / "Downloads").glob("relatorio-consolidado-*.xlsx"))
        checked = 0
        differences = 0
        for path in files:
            period = parse_consolidated_period(path)
            if period is None:
                continue
            year, month = period
            result = self.results_by_year.get(year)
            if result is None:
                continue
            checked += 1
            official = read_consolidated_positions(path)
            calc_positions = result.monthly_positions.get(month, {})
            if month == 12 and not calc_positions:
                calc_positions = result.positions
            calculated_wallet = {code: pos.qty for code, pos in calc_positions.items() if pos.qty > 0 and pos.category not in {"opcoes", "futuro"}}
            calculated_loans = result.monthly_loans.get(month, {})
            calculated_options = self._calculated_option_positions(result, month)
            differences += self._insert_b3_check_rows(year, month, "carteira", official.get("carteira", {}), calculated_wallet, path.name)
            differences += self._insert_b3_check_rows(year, month, "emprestimos", official.get("emprestimos", {}), calculated_loans, path.name)
            differences += self._insert_b3_check_rows(year, month, "opcoes", official.get("opcoes", {}), calculated_options, path.name)
        self.b3_check_status_var.set(f"Relatorios conferidos: {checked}. Diferencas encontradas: {differences}.")

    def _calculated_option_positions(self, result: CalculationResult, month: int) -> dict[str, Decimal]:
        if month == 12:
            rows = result.pending_options
        else:
            rows = [
                row
                for row in result.pending_options
                if isinstance(row.get("ultima_data"), date) and row["ultima_data"].month <= month
                and (not isinstance(row.get("vencimento"), date) or row["vencimento"].month >= month)
            ]
        return {str(row.get("codigo")): money(row.get("quantidade_aberta")) for row in rows if money(row.get("quantidade_aberta")) > 0}

    def _insert_b3_check_rows(
        self,
        year: int,
        month: int,
        kind: str,
        official: dict[str, Decimal],
        calculated: dict[str, Decimal],
        filename: str,
    ) -> int:
        differences = 0
        for code in sorted(set(official) | set(calculated)):
            b3_qty = q2(official.get(code, Decimal("0")))
            calc_qty = q2(calculated.get(code, Decimal("0")))
            diff = q2(calc_qty - b3_qty)
            tag = "ok" if diff == 0 else "diff"
            if diff != 0:
                differences += 1
            self.b3_check_tree.insert(
                "",
                "end",
                values=[year, MONTHS[month - 1], kind, code, fmt_decimal(b3_qty), fmt_decimal(calc_qty), fmt_decimal(diff), filename],
                tags=(tag,),
            )
        return differences

    def _refresh_monthly_portfolio(self, selected_year: int) -> None:
        for item in self.monthly_portfolio_tree.get_children():
            self.monthly_portfolio_tree.delete(item)
        result = self.results_by_year.get(selected_year)
        if result is None:
            return
        for month in range(1, 13):
            loans = result.monthly_loans.get(month, {})
            positions = result.monthly_positions.get(month, {})
            for code, pos in sorted(positions.items()):
                if pos.category in {"opcoes", "futuro"}:
                    continue
                if pos.qty <= 0 and loans.get(code, Decimal("0")) <= 0:
                    continue
                self.monthly_portfolio_tree.insert(
                    "",
                    "end",
                    values=[
                        selected_year,
                        MONTHS[month - 1],
                        code,
                        pos.category,
                        fmt_decimal(pos.qty),
                        fmt_money(pos.avg_price),
                        fmt_money(pos.cost),
                        fmt_decimal(loans.get(code, Decimal("0"))),
                    ],
                )

    def _open_month_value_detail(self, event: tk.Event) -> None:
        if self.month_tree.identify("region", event.x, event.y) != "cell":
            return
        item_id = self.month_tree.identify_row(event.y)
        column_id = self.month_tree.identify_column(event.x)
        if not item_id or not column_id:
            return
        col_index = int(column_id.replace("#", "")) - 1
        columns = list(self.month_tree["columns"])
        if col_index < 0 or col_index >= len(columns):
            return
        column = columns[col_index]
        if column == "mes":
            return
        values = self.month_tree.item(item_id, "values")
        if col_index >= len(values) or money(values[col_index]) == 0:
            return
        month_label = values[0]
        if month_label not in MONTHS:
            return
        selected_year = int(self.selected_year_var.get())
        rows = self._month_detail_rows(selected_year, MONTHS.index(month_label) + 1, column)
        self._show_month_detail_window(selected_year, month_label, column, values[col_index], rows)

    def _month_daytrade_keys(self, year: int, month: int) -> set[tuple[date, str, str, str]]:
        grouped: dict[tuple[date, str, str, str], set[str]] = defaultdict(set)
        for trade in self.trades:
            if trade.dt.year == year and trade.dt.month == month:
                grouped[(trade.dt, trade.code, trade.category, trade.broker)].add(trade.side)
        return {key for key, sides in grouped.items() if {"compra", "venda"}.issubset(sides)}

    def _trade_matches_month_detail(self, trade: Trade, year: int, month: int, column: str, daytrade_keys: set[tuple[date, str, str, str]]) -> bool:
        if trade.dt.year != year or trade.dt.month != month:
            return False
        is_dt = (trade.dt, trade.code, trade.category, trade.broker) in daytrade_keys
        if column in {"daytrade", "base_dt"}:
            return is_dt
        if is_dt:
            return False
        if column in {"normal", "base_normal", "imposto"}:
            return trade.category not in {"fii", "opcoes", "futuro"}
        if column in {"fii", "base_fii"}:
            return trade.category == "fii"
        if column == "opcoes":
            return trade.category == "opcoes"
        if column == "futuro":
            return trade.category == "futuro"
        return False

    def _daytrade_detail_rows(self, year: int, month: int, include_future: bool | None = None) -> list[list[str]]:
        rows: list[list[str]] = []
        grouped: dict[tuple[date, str, str, str], list[Trade]] = defaultdict(list)
        for trade in self.trades:
            if trade.dt.year == year and trade.dt.month == month:
                grouped[(trade.dt, trade.code, trade.category, trade.broker)].append(trade)
        for (_dt, _code, _category, _broker), group in sorted(grouped.items()):
            buys = [item for item in group if item.side == "compra"]
            sells = [item for item in group if item.side == "venda"]
            buy_qty = sum((item.qty for item in buys), Decimal("0"))
            sell_qty = sum((item.qty for item in sells), Decimal("0"))
            dt_qty = min(buy_qty, sell_qty)
            if dt_qty <= 0:
                continue
            buy_avg = sum((item.value for item in buys), Decimal("0")) / buy_qty if buy_qty else Decimal("0")
            sell_avg = sum((item.value for item in sells), Decimal("0")) / sell_qty if sell_qty else Decimal("0")
            sale_value = q2(sell_avg * dt_qty)
            cost = q2(buy_avg * dt_qty)
            profit = q2(sale_value - cost)
            sample = group[0]
            is_future = sample.category == "futuro"
            if include_future is True and not is_future:
                continue
            if include_future is False and is_future:
                continue
            rows.append([
                "day trade",
                sample.dt.strftime("%d/%m/%Y"),
                "compra/venda",
                sample.market,
                sample.code,
                fmt_decimal(dt_qty),
                fmt_money(sell_avg),
                fmt_money(buy_avg),
                fmt_money(cost),
                fmt_money(sale_value),
                fmt_money(profit),
                sample.broker,
            ])
        return rows

    def _sale_detail_rows(self, year: int, month: int, column: str) -> list[list[str]]:
        if column == "daytrade" or column == "base_dt":
            return self._daytrade_detail_rows(year, month, include_future=False)
        if column == "futuro":
            return self._daytrade_detail_rows(year, month, include_future=True)
        if column not in {"normal", "base_normal", "fii", "base_fii", "futuro", "opcoes", "imposto"}:
            return []
        start_year = min(self.results_by_year) if self.results_by_year else year
        if (year - 1) in self.results_by_year:
            initial_positions = self._carry_positions(self.results_by_year[year - 1])
        else:
            initial_positions = self._manual_positions()
        positions = {pos.code: Position(code=pos.code, qty=money(pos.qty), cost=q2(pos.cost), category=pos.category, broker=pos.broker) for pos in initial_positions}
        IRSimpleEngine(year)._apply_events_before_trades(positions, self._events_for_year(self._manual_events(), year, start_year), [])
        daytrade_keys = self._month_daytrade_keys(year, month)
        short_positions: dict[str, dict[str, Any]] = {}
        rows: list[list[str]] = []
        trades = [trade for trade in self.trades if trade.dt.year == year and trade.category != "opcoes"]
        def include_category(category: str) -> bool:
            if column == "normal":
                return category not in {"fii", "opcoes", "futuro"}
            if column in {"base_normal", "imposto"}:
                return category != "fii"
            if column in {"fii", "base_fii"}:
                return category == "fii"
            if column == "futuro":
                return category == "futuro"
            return False
        for trade in sorted(trades, key=lambda item: (item.dt, item.code, item.side)):
            if (trade.dt, trade.code, trade.category, trade.broker) in daytrade_keys:
                continue
            pos = positions.setdefault(trade.code, Position(code=trade.code, category=trade.category, broker=trade.broker))
            if trade.side == "compra":
                remaining_qty = trade.qty
                remaining_value = trade.value
                short = short_positions.get(trade.code)
                if short and short["qty"] > 0:
                    cover_qty = min(remaining_qty, short["qty"])
                    cover_cost = q2(trade.value * cover_qty / trade.qty) if trade.qty else Decimal("0")
                    cover_credit = q2(short["credit"] * cover_qty / short["qty"]) if short["qty"] else Decimal("0")
                    profit = q2(cover_credit - cover_cost)
                    if trade.dt.month == month and include_category(trade.category):
                        avg_sale = cover_credit / cover_qty if cover_qty else Decimal("0")
                        avg_buy = cover_cost / cover_qty if cover_qty else Decimal("0")
                        rows.append(["negociacao", trade.dt.strftime("%d/%m/%Y"), "recompra venda descoberta", trade.market, trade.code, fmt_decimal(cover_qty), fmt_money(avg_sale), fmt_money(avg_buy), fmt_money(cover_cost), fmt_money(cover_credit), fmt_money(profit), trade.broker])
                    short["qty"] = q2(short["qty"] - cover_qty)
                    short["credit"] = q2(short["credit"] - cover_credit)
                    remaining_qty = q2(remaining_qty - cover_qty)
                    remaining_value = q2(remaining_value - cover_cost)
                    if short["qty"] <= 0:
                        del short_positions[trade.code]
                if remaining_qty > 0:
                    pos.qty += remaining_qty
                    pos.cost = q2(pos.cost + remaining_value)
                continue
            sell_remaining = trade.qty
            sell_remaining_value = trade.value
            if pos.qty > 0:
                covered_qty = min(pos.qty, sell_remaining)
                sale_value = q2(trade.value * covered_qty / trade.qty) if trade.qty else Decimal("0")
                avg = pos.avg_price
                cost = q2(avg * covered_qty)
                profit = q2(sale_value - cost)
                if trade.dt.month == month and include_category(trade.category):
                    rows.append(["negociacao", trade.dt.strftime("%d/%m/%Y"), trade.side, trade.market, trade.code, fmt_decimal(covered_qty), fmt_money(trade.price), fmt_money(avg), fmt_money(cost), fmt_money(sale_value), fmt_money(profit), trade.broker])
                pos.qty = q2(pos.qty - covered_qty)
                pos.cost = q2(max(Decimal("0"), pos.cost - cost))
                sell_remaining = q2(sell_remaining - covered_qty)
                sell_remaining_value = q2(sell_remaining_value - sale_value)
                if pos.qty <= 0:
                    pos.qty = Decimal("0")
                    pos.cost = Decimal("0")
            if sell_remaining > 0:
                short = short_positions.setdefault(trade.code, {"qty": Decimal("0"), "credit": Decimal("0")})
                short["qty"] = q2(short["qty"] + sell_remaining)
                short["credit"] = q2(short["credit"] + sell_remaining_value)
                if trade.dt.month == month and include_category(trade.category):
                    rows.append([
                        "negociacao",
                        trade.dt.strftime("%d/%m/%Y"),
                        "venda descoberta aberta",
                        trade.market,
                        trade.code,
                        fmt_decimal(sell_remaining),
                        fmt_money(trade.price),
                        "",
                        "",
                        fmt_money(sell_remaining_value),
                        "",
                        f"Venda sem posicao suficiente no detalhe. Resultado sera apurado quando houver recompra ou ajuste manual; {trade.broker}",
                    ])
        return rows

    def _note_price_lookup(self, year: int, month: int) -> dict[tuple[str, str], dict[str, Decimal]]:
        lookup: dict[tuple[str, str], dict[str, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
        if PdfReader is None:
            return lookup
        for path in self._brokerage_note_files_for_month(year, month):
            try:
                reader = PdfReader(str(path))
                for page in reader.pages:
                    text = page.extract_text() or ""
                    note_date_match = re.search(r"(\d{2}/\d{2}/\d{4})\s*Data de refer", text, re.I)
                    note_date = parse_date(note_date_match.group(1)) if note_date_match else None
                    if note_date is None:
                        dates = [parse_date(value) for value in re.findall(r"\d{2}/\d{2}/\d{4}", text)]
                        note_date = next((value for value in dates if value and value.year == year and value.month == month), None)
                    if note_date is None:
                        continue
                    for code in sorted(set(re.findall(r"\b[A-Z]{4,5}\d{1,4}[A-Z0-9]*\b", text))):
                        normalized = normalize_ticker(code)
                        for match in re.finditer(rf"\b{re.escape(code)}\b", text):
                            snippet = text[match.start(): match.start() + 1400]
                            buy_match = re.search(r"COMPRA[\s\S]{0,260}?R\$\s*(" + NOTE_PRICE_RE + r")", snippet, re.I)
                            sell_match = re.search(r"VENDA[\s\S]{0,260}?R\$\s*(" + NOTE_PRICE_RE + r")", snippet, re.I)
                            avg_buy_match = re.search(r"Pre[cç]o\s+m[eé]dio\s+compra:\s*R\$\s*(" + NOTE_PRICE_RE + r")", snippet, re.I)
                            avg_sell_match = re.search(r"Pre[cç]o\s+m[eé]dio\s+venda:\s*R\$\s*(" + NOTE_PRICE_RE + r")", snippet, re.I)
                            key = (note_date.isoformat(), normalized)
                            if avg_buy_match:
                                lookup[key]["preco_medio_compra"] = note_price(avg_buy_match.group(1))
                            elif buy_match:
                                lookup[key]["preco_medio_compra"] = note_price(buy_match.group(1))
                            if avg_sell_match:
                                lookup[key]["preco_medio_venda"] = note_price(avg_sell_match.group(1))
                            elif sell_match:
                                lookup[key]["preco_medio_venda"] = note_price(sell_match.group(1))
            except Exception as exc:
                debug_log(f"Falha ao ler precos da nota {path}: {exc}")
        return lookup

    def _with_note_prices(self, rows: list[list[str]], year: int, month: int) -> list[list[str]]:
        lookup = self._note_price_lookup(year, month)
        enriched: list[list[str]] = []
        for row in rows:
            item = list(row)
            while len(item) < 12:
                item.append("")
            note_buy = ""
            note_sell = ""
            dt = parse_date(item[1])
            code = normalize_ticker(item[4])
            if dt and code:
                values = lookup.get((dt.isoformat(), code), {})
                note_buy = fmt_decimal(values.get("preco_medio_compra")) if values.get("preco_medio_compra") else ""
                note_sell = fmt_decimal(values.get("preco_medio_venda")) if values.get("preco_medio_venda") else ""
            item.insert(8, note_buy)
            item.insert(9, note_sell)
            enriched.append(item)
        return enriched

    def _option_detail_rows(self, year: int, month: int) -> list[list[str]]:
        rows: list[list[str]] = []
        result = self.results_by_year.get(year)
        events = self._events_for_year(self._manual_events(), year, min(self.results_by_year) if self.results_by_year else year)
        resolved = IRSimpleEngine(year)._resolved_option_events(events)
        resolved_rows = IRSimpleEngine(year)._resolved_option_event_rows(events)
        cutoff = date(year, 12, 31)
        positions: dict[str, dict[str, Any]] = {}
        option_trades = [
            trade
            for trade in self.trades
            if trade.category == "opcoes" and trade.dt <= cutoff and "exercicio" not in normalize_header(trade.market)
        ]
        for trade in sorted(option_trades, key=lambda item: (item.dt, item.code, item.side)):
            pos = positions.setdefault(
                trade.code,
                {
                    "vencimento": trade.expiry,
                    "long_qty": Decimal("0"),
                    "long_cost": Decimal("0"),
                    "short_qty": Decimal("0"),
                    "short_credit": Decimal("0"),
                    "ultima_data": trade.dt,
                    "corretora": trade.broker,
                },
            )
            if trade.expiry and (pos["vencimento"] is None or trade.expiry > pos["vencimento"]):
                pos["vencimento"] = trade.expiry
            pos["ultima_data"] = max(pos["ultima_data"], trade.dt)
            if trade.side == "compra":
                qty_to_close = min(trade.qty, pos["short_qty"])
                if qty_to_close > 0:
                    avg_credit = pos["short_credit"] / pos["short_qty"] if pos["short_qty"] else Decimal("0")
                    buy_cost = q2(trade.price * qty_to_close)
                    sale_value = q2(avg_credit * qty_to_close)
                    profit = q2(sale_value - buy_cost)
                    if trade.dt.year == year and trade.dt.month == month:
                        rows.append([
                            "opcoes",
                            trade.dt.strftime("%d/%m/%Y"),
                            "recompra de opcao vendida",
                            trade.market,
                            trade.code,
                            fmt_decimal(qty_to_close),
                            fmt_money(trade.price),
                            fmt_money(avg_credit),
                            fmt_money(buy_cost),
                            fmt_money(sale_value),
                            fmt_money(profit),
                            f"premio medio recebido na venda=R$ {fmt_money(avg_credit)}; recompra=R$ {fmt_money(trade.price)}; {trade.broker}",
                        ])
                    pos["short_qty"] = q2(pos["short_qty"] - qty_to_close)
                    pos["short_credit"] = q2(max(Decimal("0"), pos["short_credit"] - sale_value))
                remaining = q2(trade.qty - qty_to_close)
                if remaining > 0:
                    pos["long_qty"] = q2(pos["long_qty"] + remaining)
                    pos["long_cost"] = q2(pos["long_cost"] + trade.price * remaining)
                    if trade.dt.year == year and trade.dt.month == month:
                        rows.append([
                            "opcoes",
                            trade.dt.strftime("%d/%m/%Y"),
                            "compra abre posicao",
                            trade.market,
                            trade.code,
                            fmt_decimal(remaining),
                            fmt_money(trade.price),
                            "",
                            fmt_money(q2(trade.price * remaining)),
                            "",
                            "",
                            f"premio pago fica em estoque para calcular venda/exercicio/vencimento; {trade.broker}",
                        ])
            else:
                qty_to_close = min(trade.qty, pos["long_qty"])
                if qty_to_close > 0:
                    avg_cost = pos["long_cost"] / pos["long_qty"] if pos["long_qty"] else Decimal("0")
                    sale_value = q2(trade.price * qty_to_close)
                    cost = q2(avg_cost * qty_to_close)
                    profit = q2(sale_value - cost)
                    if trade.dt.year == year and trade.dt.month == month:
                        rows.append([
                            "opcoes",
                            trade.dt.strftime("%d/%m/%Y"),
                            "venda de opcao comprada",
                            trade.market,
                            trade.code,
                            fmt_decimal(qty_to_close),
                            fmt_money(trade.price),
                            fmt_money(avg_cost),
                            fmt_money(cost),
                            fmt_money(sale_value),
                            fmt_money(profit),
                            f"premio medio pago na compra=R$ {fmt_money(avg_cost)}; venda=R$ {fmt_money(trade.price)}; {trade.broker}",
                        ])
                    pos["long_qty"] = q2(pos["long_qty"] - qty_to_close)
                    pos["long_cost"] = q2(max(Decimal("0"), pos["long_cost"] - cost))
                remaining = q2(trade.qty - qty_to_close)
                if remaining > 0:
                    pos["short_qty"] = q2(pos["short_qty"] + remaining)
                    pos["short_credit"] = q2(pos["short_credit"] + trade.price * remaining)
                    if trade.dt.year == year and trade.dt.month == month:
                        rows.append([
                            "opcoes",
                            trade.dt.strftime("%d/%m/%Y"),
                            "venda abre posicao",
                            trade.market,
                            trade.code,
                            fmt_decimal(remaining),
                            fmt_money(trade.price),
                            "",
                            "",
                            fmt_money(q2(trade.price * remaining)),
                            "",
                            f"premio recebido fica em aberto ate recompra/exercicio/vencimento; {trade.broker}",
                        ])

        for code, pos in sorted(positions.items()):
            expiry = pos.get("vencimento")
            event_type = resolved.get(code)
            if event_type and any(marker in event_type for marker in ["liquidacao_trava", "trava_sem_acao", "sem_exercicio_acao", "exercicio", "virou_po", "po"]):
                continue
            if event_type and ("recompra" in event_type or "recomprada" in event_type):
                event = resolved_rows.get(code)
                event_date = parse_date(event.get("data")) if event else None
                if event and event_date and event_date.year == year and event_date.month == month:
                    value = q2(event.get("valor"))
                    qty = abs(money(event.get("quantidade"))) or pos.get("short_qty") or pos.get("long_qty") or Decimal("0")
                    if pos.get("short_qty", Decimal("0")) > 0:
                        avg_credit = pos["short_credit"] / pos["short_qty"] if pos["short_qty"] else Decimal("0")
                        profit = q2(pos["short_credit"] - value)
                        rows.append(["evento manual", event_date.strftime("%d/%m/%Y"), "recompra informada", "", code, fmt_decimal(qty), fmt_money(value / qty if qty else 0), fmt_money(avg_credit), fmt_money(value), fmt_money(pos["short_credit"]), fmt_money(profit), event.get("observacao", "")])
                    elif pos.get("long_qty", Decimal("0")) > 0:
                        avg_cost = pos["long_cost"] / pos["long_qty"] if pos["long_qty"] else Decimal("0")
                        profit = q2(value - pos["long_cost"])
                        rows.append(["evento manual", event_date.strftime("%d/%m/%Y"), "venda/recompra informada", "", code, fmt_decimal(qty), fmt_money(value / qty if qty else 0), fmt_money(avg_cost), fmt_money(pos["long_cost"]), fmt_money(value), fmt_money(profit), event.get("observacao", "")])
                continue
            if isinstance(expiry, date) and expiry.year == year and expiry.month == month:
                if pos.get("long_qty", Decimal("0")) > 0 and not (event_type and IRSimpleEngine(year)._is_open_option_event(event_type)):
                    avg_cost = pos["long_cost"] / pos["long_qty"] if pos["long_qty"] else Decimal("0")
                    rows.append(["opcoes", expiry.strftime("%d/%m/%Y"), "vencimento virou po", "", code, fmt_decimal(pos["long_qty"]), "", fmt_money(avg_cost), fmt_money(pos["long_cost"]), "0,00", fmt_money(-pos["long_cost"]), "perda do premio medio pago por opcao comprada vencida"])
                if pos.get("short_qty", Decimal("0")) > 0 and not (event_type and IRSimpleEngine(year)._is_open_option_event(event_type)):
                    avg_credit = pos["short_credit"] / pos["short_qty"] if pos["short_qty"] else Decimal("0")
                    rows.append(["opcoes", expiry.strftime("%d/%m/%Y"), "vencimento sem recompra", "", code, fmt_decimal(pos["short_qty"]), "", fmt_money(avg_credit), "0,00", fmt_money(pos["short_credit"]), fmt_money(pos["short_credit"]), "ganho do premio medio recebido por opcao vendida vencida"])
        if result:
            calculated = result.monthly[month - 1].options_result
            detail_total = sum((money(row[10]) for row in rows), Decimal("0"))
            if q2(detail_total) != q2(calculated):
                rows.append([
                    "conferencia",
                    "",
                    "diferenca detalhe x calculo",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    fmt_money(q2(calculated - detail_total)),
                    f"resultado de opcoes no mes=R$ {fmt_money(calculated)}; detalhe listado=R$ {fmt_money(detail_total)}. Pode envolver exercicio de acao-base/evento manual.",
                ])
        return rows

    def _month_detail_rows(self, year: int, month: int, column: str) -> list[list[str]]:
        result = self.results_by_year.get(year)
        rows: list[list[str]] = []
        if result:
            tax = result.monthly[month - 1]
            if column == "base_normal":
                rows.append(["calculo", "", "Base normal", "", "", "", "", "", "", fmt_money(tax.normal_base), "", f"Comum tributavel: resultado normal {fmt_money(tax.normal_result)}, opcoes {fmt_money(tax.options_result)}, futuro {fmt_money(tax.future_result)}; prejuizo anterior {fmt_money(tax.normal_loss_before)}"])
            elif column == "base_dt":
                future_dt = q2(tax.future_dollar_daytrade + tax.future_index_daytrade)
                rows.append(["calculo", "", "Base day trade", "", "", "", "", "", "", fmt_money(tax.daytrade_base), "", f"Resultado DT nao futuro {fmt_money(tax.daytrade_result)}; DT futuro {fmt_money(future_dt)}; prejuizo anterior {fmt_money(tax.daytrade_loss_before)}"])
            elif column == "base_fii":
                rows.append(["calculo", "", "Base FII", "", "", "", "", "", "", fmt_money(tax.fii_base), "", f"Resultado {fmt_money(tax.fii_result)}; prejuizo anterior {fmt_money(tax.fii_loss_before)}"])
            elif column == "imposto":
                rows.append(["calculo", "", "Imposto a pagar", "", "", "", "", "", "", fmt_money(tax.tax_payable), "", f"Imposto devido {fmt_money(tax.tax_due)} menos IRRF compensavel comum/DT/FII."])
        if column == "opcoes":
            rows.extend(self._option_detail_rows(year, month))
        else:
            rows.extend(self._sale_detail_rows(year, month, column))
        daytrade_keys = self._month_daytrade_keys(year, month)
        for trade in sorted(self.trades, key=lambda item: (item.dt, item.code, item.side)):
            if not self._trade_matches_month_detail(trade, year, month, column, daytrade_keys):
                continue
            if column in {"normal", "base_normal", "fii", "base_fii", "futuro", "daytrade", "base_dt", "imposto", "opcoes"}:
                continue
            rows.append([
                "negociacao",
                trade.dt.strftime("%d/%m/%Y"),
                trade.side,
                trade.market,
                trade.code,
                fmt_decimal(trade.qty),
                fmt_money(trade.price),
                "",
                "",
                fmt_money(trade.value),
                "",
                trade.broker,
            ])
        if column in {"opcoes", "imposto"}:
            for event_row in self.event_table.rows():
                event_date = parse_date(event_row.get("data"))
                if not event_date or event_date.year != year or event_date.month != month:
                    continue
                event_type = normalize_header(event_row.get("tipo"))
                if column == "opcoes" and "opcao" not in event_type and "exercicio" not in event_type:
                    continue
                rows.append([
                    "evento manual",
                    event_date.strftime("%d/%m/%Y"),
                    event_row.get("tipo", ""),
                    "",
                    event_row.get("ativo", ""),
                    event_row.get("quantidade", ""),
                    "",
                    "",
                    "",
                    event_row.get("valor", ""),
                    "",
                    event_row.get("observacao", ""),
                ])
        return self._with_note_prices(rows, year, month)

    def _show_month_detail_window(self, year: int, month_label: str, column: str, value: str, rows: list[list[str]]) -> None:
        win = tk.Toplevel(self.root)
        win.title(f"Detalhe do calculo - {month_label}/{year} - {column}")
        win.geometry("1120x520")
        ttk.Label(win, text=f"{month_label}/{year} | {column} = R$ {value}", foreground="#0f172a").pack(fill="x", padx=10, pady=(10, 4))
        frame = ttk.Frame(win, padding=10)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(0, weight=1)
        cols = ["origem", "data", "tipo", "mercado", "ativo", "quantidade", "preco_b3", "preco_medio_calc", "preco_medio_nota", "preco_venda_nota", "custo", "venda", "lucro", "observacao"]
        tree = ttk.Treeview(frame, columns=cols, show="headings", height=16)
        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        for col, width in [
            ("origem", 100),
            ("data", 90),
            ("tipo", 170),
            ("mercado", 130),
            ("ativo", 90),
            ("quantidade", 90),
            ("preco_b3", 90),
            ("preco_medio_calc", 120),
            ("preco_medio_nota", 120),
            ("preco_venda_nota", 120),
            ("custo", 105),
            ("venda", 105),
            ("lucro", 105),
            ("observacao", 380),
        ]:
            tree.heading(col, text=col)
            tree.column(col, width=width, anchor="e" if col in {"quantidade", "preco_b3", "preco_medio_calc", "preco_medio_nota", "preco_venda_nota", "custo", "venda", "lucro"} else "w")
        total_sale = Decimal("0")
        total_profit = Decimal("0")
        for row in rows:
            if len(row) < len(cols):
                row = row + [""] * (len(cols) - len(row))
            total_sale += money(row[11])
            total_profit += money(row[12])
            tree.insert("", "end", values=row)
        tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        subtotals = self._detail_subtotals(rows)
        if subtotals:
            subtotal_text = " | ".join(f"{name}: R$ {fmt_money(total)}" for name, total in subtotals.items())
            ttk.Label(win, text=f"Subtotais do resultado: {subtotal_text}", foreground="#334155", wraplength=1080).pack(fill="x", padx=10, pady=(0, 6))
        buttons = ttk.Frame(win)
        buttons.pack(fill="x", padx=10, pady=(0, 6))
        ttk.Label(buttons, text="Corrigir suposicao selecionada:").pack(side="left")
        ttk.Button(buttons, text="Registrar recompra", command=lambda: self._detail_register_option_fix(tree, "opcao recomprada")).pack(side="left", padx=6)
        ttk.Button(buttons, text="Registrar exercicio", command=lambda: self._detail_register_option_fix(tree, "exercicio de opcoes")).pack(side="left")
        ttk.Button(buttons, text="Registrar virou po", command=lambda: self._detail_register_option_fix(tree, "opcao virou po")).pack(side="left", padx=6)
        ttk.Button(buttons, text="Registrar venda ativa", command=lambda: self._detail_register_option_fix(tree, "opcao vendida ainda ativa")).pack(side="left")
        ttk.Button(buttons, text="Registrar em aberto", command=lambda: self._detail_register_option_fix(tree, "opcao ainda em aberto")).pack(side="left", padx=6)
        ttk.Button(buttons, text="Registrar posicao/migracao", command=lambda: self._detail_register_position_fix(tree)).pack(side="left", padx=6)
        ttk.Label(win, text=f"Totais do detalhamento: venda R$ {fmt_money(total_sale)} | lucro/prejuizo R$ {fmt_money(total_profit)}", foreground="#0f172a").pack(fill="x", padx=10, pady=(0, 8))

    def _detail_register_option_fix(self, tree: ttk.Treeview, event_type: str) -> None:
        selected = tree.selection()
        if not selected:
            messagebox.showinfo(APP_TITLE, "Selecione uma linha do detalhe para corrigir.")
            return
        values = list(tree.item(selected[0], "values"))
        if len(values) < 6:
            return
        code = normalize_ticker(values[4])
        if not code:
            messagebox.showinfo(APP_TITLE, "A linha selecionada nao possui ativo/opcao para registrar.")
            return
        qty = fmt_decimal(abs(money(values[5])))
        default_value = values[10] if event_type == "opcao recomprada" else values[11] if len(values) > 11 else "0,00"
        columns = [
            ("data", "Data do evento", 90),
            ("tipo", "Tipo", 170),
            ("ativo", "Opcao/ativo", 120),
            ("quantidade", "Quantidade", 100),
            ("valor", "Valor", 120),
            ("observacao", "Observacao", 420),
        ]
        defaults = [
            values[1] or date.today().strftime("%d/%m/%Y"),
            event_type,
            code,
            qty,
            default_value or "0,00",
            "Correcao lancada pelo detalhe do calculo; substitui suposicao automatica apos recalcular.",
        ]
        dialog = RowEditor(self.root, columns, defaults)
        self.root.wait_window(dialog)
        if dialog.result is None:
            return
        data, tipo, ativo, quantidade, valor, observacao = dialog.result
        self.event_table.add_row([data, tipo or event_type, ativo or code, quantidade or qty, valor or "0,00", "opcoes", observacao])
        self._refresh_option_events()
        self._save_option_events()
        self.status_var.set(f"Correcao registrada para {ativo or code}. Clique em Recalcular agora.")

    def _detail_register_position_fix(self, tree: ttk.Treeview) -> None:
        selected = tree.selection()
        if not selected:
            messagebox.showinfo(APP_TITLE, "Selecione uma linha do detalhe para registrar uma correcao.")
            return
        values = list(tree.item(selected[0], "values"))
        if len(values) < 6:
            return
        code = normalize_ticker(values[4])
        if not code:
            messagebox.showinfo(APP_TITLE, "A linha selecionada nao possui ativo.")
            return
        qty = fmt_decimal(abs(money(values[5])))
        unit_cost = values[7] or values[6] or "0,00"
        total_cost = fmt_money(q2(money(unit_cost) * money(qty)))
        category = classify_asset(code, values[3] if len(values) > 3 else "")
        columns = [
            ("data", "Data", 90),
            ("tipo", "Tipo", 190),
            ("ativo", "Ativo", 110),
            ("quantidade", "Quantidade", 100),
            ("valor", "Custo/valor total", 130),
            ("observacao", "Observacao", 420),
        ]
        defaults = [
            values[1] or date.today().strftime("%d/%m/%Y"),
            "migracao/transferencia",
            code,
            qty,
            total_cost,
            "Correcao registrada pelo detalhe; use para posicao inicial, migracao, bonificacao, split ou ajuste de custodia.",
        ]
        dialog = RowEditor(self.root, columns, defaults)
        self.root.wait_window(dialog)
        if dialog.result is None:
            return
        data, tipo, ativo, quantidade, valor, observacao = dialog.result
        self.event_table.add_row([data, tipo or "migracao/transferencia", ativo or code, quantidade or qty, valor or total_cost, category, observacao])
        self._save_manual_data_to_db()
        self.status_var.set(f"Correcao patrimonial registrada para {ativo or code}. Clique em Recalcular agora.")
        if not rows:
            ttk.Label(win, text="Nenhum registro individual localizado para esta celula. Verifique eventos manuais e prejuizos acumulados.", foreground="#475569").pack(fill="x", padx=10, pady=(0, 10))

    def _suggested_asset_category(self, code: str, market: str = "") -> str:
        code = normalize_ticker(code)
        for mov in self.movements:
            if mov.code == code and mov.category:
                return mov.category
        for trade in self.trades:
            if trade.code == code and trade.category and trade.category != "normal":
                return trade.category
        category = classify_asset(code, market)
        if category == "normal" and re.fullmatch(r"[A-Z]{4}11", code):
            return "fii"
        return category

    def _parse_alert_solution(self, idx: int, warning: str) -> dict[str, Any] | None:
        is_exceeded_sale = "Venda excede a posicao cadastrada" in warning
        is_active_short = "Venda descoberta em aberto" in warning
        if not is_exceeded_sale and not is_active_short:
            return None
        fields = dict(re.findall(r"([a-z_]+)=([^;.\n]+)", warning))
        code = normalize_ticker(fields.get("ativo") or fields.get("codigo"))
        missing = money(fields.get("falta") or fields.get("qtd_vendida_sem_recompra"))
        if not code or missing <= 0:
            return None
        trade_date = parse_date(fields.get("data"))
        price = money(str(fields.get("preco", "")).replace("R$", ""))
        suggested_value = q2(fields.get("credito_liquido")) if is_active_short else Decimal("0")
        if suggested_value <= 0 and price > 0:
            suggested_value = q2(missing * price)
        before_qty = money(fields.get("posicao_antes"))
        sold_qty = money(fields.get("qtd_vendida"))
        split_factor = q2(sold_qty / before_qty) if before_qty > 0 and sold_qty > before_qty else Decimal("0")
        if is_active_short:
            suggestion = f"Registrar venda descoberta ativa de {fmt_decimal(missing)} {code} ou cadastrar posicao inicial/migracao se a venda nao foi descoberta."
        elif before_qty <= 0:
            suggestion = f"Cadastrar posicao inicial ou migracao de {fmt_decimal(missing)} {code}."
        elif missing <= Decimal("1"):
            suggestion = f"Conferir arredondamento/fracionario ou cadastrar ajuste de {fmt_decimal(missing)} {code}."
        else:
            suggestion = f"Cadastrar posicao inicial/migracao de {fmt_decimal(missing)} {code}; se houve evento societario, avaliar desdobramento/fator {fmt_decimal(split_factor) if split_factor else ''}."
        event_qty = fields.get("qtd_vendida") or fields.get("qtd") or fields.get("qtd_vendida_sem_recompra") or ""
        event = f"{fields.get('tipo', '')} em {fields.get('data', '')}; qtd {event_qty}; preco {fields.get('preco', '')}; corretora {fields.get('corretora', '')}"
        return {
            "idx": str(idx),
            "ativo": code,
            "data": trade_date.strftime("%d/%m/%Y") if trade_date else fields.get("data", ""),
            "falta": fmt_decimal(missing),
            "valor": fmt_money(suggested_value),
            "valor_decimal": suggested_value,
            "categoria": self._suggested_asset_category(code, fields.get("mercado", "")),
            "corretora": fields.get("corretora", ""),
            "sugestao": suggestion,
            "evento": event,
            "fator": fmt_decimal(split_factor) if split_factor else "",
        }

    def _refresh_alert_solutions(self, warnings: list[str]) -> None:
        self.alert_solutions: dict[str, dict[str, Any]] = {}
        for item in self.alert_solution_tree.get_children():
            self.alert_solution_tree.delete(item)
        for idx, warning in enumerate(warnings, start=1):
            solution = self._parse_alert_solution(idx, warning)
            if not solution:
                continue
            item_id = self.alert_solution_tree.insert(
                "",
                "end",
                values=[
                    solution["idx"],
                    solution["ativo"],
                    solution["data"],
                    solution["falta"],
                    solution["valor"],
                    "Pendente",
                    solution["sugestao"],
                    solution["evento"],
                ],
                tags=("pendente",),
            )
            self.alert_solutions[item_id] = solution

    def _selected_alert_solution(self) -> dict[str, Any] | None:
        selected = self.alert_solution_tree.selection()
        if not selected:
            messagebox.showinfo(APP_TITLE, "Selecione uma solucao sugerida.")
            return None
        return self.alert_solutions.get(selected[0])

    def _mark_alert_solution_applied(self, message: str) -> None:
        selected = self.alert_solution_tree.selection()
        if not selected:
            return
        item = selected[0]
        values = list(self.alert_solution_tree.item(item, "values"))
        if len(values) >= 6:
            values[5] = "Aplicado"
        if len(values) >= 7:
            values[6] = message
        self.alert_solution_tree.item(item, values=values, tags=("aplicado",))

    def _register_alert_initial_position(self) -> None:
        solution = self._selected_alert_solution()
        if not solution:
            return
        self.position_table.add_row([
            solution["ativo"],
            solution["falta"],
            solution["valor"],
            solution["categoria"],
            solution["corretora"],
        ])
        self._save_manual_data_to_db()
        self._mark_alert_solution_applied(f"Aplicado: posicao inicial de {solution['falta']} {solution['ativo']}.")
        self.status_var.set(f"Posicao inicial sugerida registrada para {solution['ativo']}. Clique em Recalcular agora.")

    def _register_alert_event(self, event_type: str) -> None:
        solution = self._selected_alert_solution()
        if not solution:
            return
        if "desdobramento" in event_type:
            qty_or_factor = solution.get("fator") or "1"
            value = "0,00"
            observation = f"Ajuste sugerido pelo alerta {solution['idx']}. Confira o fator antes de recalcular."
        else:
            qty_or_factor = solution["falta"]
            value = solution["valor"]
            observation = f"Ajuste sugerido pelo alerta {solution['idx']}: {solution['evento']}"
        self.event_table.add_row([
            solution["data"],
            event_type,
            solution["ativo"],
            qty_or_factor,
            value,
            solution["categoria"],
            observation,
        ])
        self._save_manual_data_to_db()
        self._mark_alert_solution_applied(f"Aplicado: evento '{event_type}' registrado.")
        self.status_var.set(f"Evento '{event_type}' registrado para {solution['ativo']}. Clique em Recalcular agora.")

    def _register_corporate_alert_event(self, event_type: str) -> None:
        solution = self._selected_alert_solution()
        if not solution:
            return
        if "split" in event_type or "desdobramento" in event_type:
            qty_or_factor = solution.get("fator") or ""
            value = "0,00"
            note = (
                f"Evento corporativo sugerido pelo alerta {solution['idx']}. "
                f"Informe o fator do split/desdobramento e justifique a fonte da informacao."
            )
        else:
            qty_or_factor = solution["falta"]
            value = "0,00"
            note = (
                f"Bonificacao sugerida pelo alerta {solution['idx']}. "
                f"Informe a quantidade bonificada, eventual custo atribuido e a fonte da informacao."
            )
        columns = [
            ("data", "Data", 90),
            ("tipo", "Tipo", 160),
            ("ativo", "Ativo", 90),
            ("quantidade", "Quantidade/Fator", 120),
            ("valor", "Valor/Custo", 100),
            ("categoria", "Categoria", 90),
            ("observacao", "Observacao", 360),
        ]
        defaults = [
            solution["data"],
            event_type,
            solution["ativo"],
            qty_or_factor,
            value,
            solution["categoria"],
            note,
        ]
        dialog = RowEditor(self.root, columns, defaults)
        self.root.wait_window(dialog)
        if dialog.result is None:
            return
        self.event_table.add_row(dialog.result)
        self._save_manual_data_to_db()
        self._mark_alert_solution_applied(f"Aplicado: evento corporativo '{dialog.result[1]}' registrado.")
        self.status_var.set(f"Evento corporativo registrado para {solution['ativo']}. Clique em Recalcular agora.")

    def _refresh_results(self) -> None:
        result = self._selected_result()
        if result is None:
            self.refresh_brokerage_compare_view()
            return
        self.refresh_brokerage_compare_view()
        self.result = result
        selected_year = result.year
        self.exercise_var.set(
            f"Declaração exibida: exercício {selected_year + 1} | ano-calendário {selected_year} | "
            f"comparação patrimonial: 31/12/{selected_year - 1} x 31/12/{selected_year}"
        )
        for tree in [self.month_tree, *self.annual_trees.values()]:
            for item in tree.get_children():
                tree.delete(item)
        for item in result.monthly:
            self.month_tree.insert(
                "",
                "end",
                values=[
                    MONTHS[item.month - 1],
                    fmt_money(item.normal_result),
                    fmt_money(item.daytrade_result),
                    fmt_money(item.fii_result),
                    fmt_money(item.options_result),
                    fmt_money(item.future_result),
                    fmt_money(item.normal_base),
                    fmt_money(item.daytrade_base),
                    fmt_money(item.fii_base),
                    fmt_money(item.tax_payable),
                ],
            )
        self.warn_text.delete("1.0", "end")
        warnings = result.warnings or ["Sem alertas criticos. Ainda assim confira custos, IRRF, exercicio de opcoes e eventos societarios."]
        self.warn_text.insert("1.0", "\n\n".join(f"{idx}. {warning}" for idx, warning in enumerate(warnings, start=1)))
        self._refresh_alert_solutions(result.warnings)

        for code, pos in sorted(result.positions.items()):
            if pos.qty <= 0 or pos.category in {"opcoes", "futuro"}:
                continue
            group = annual_asset_group(pos)
            discr = f"{code} - Quantidade: {fmt_decimal(pos.qty)} - Preco medio: R$ {fmt_money(pos.avg_price)}"
            self.annual_trees["bens"].insert("", "end", values=[group, discr, fmt_decimal(pos.previous_qty), fmt_money(pos.previous_cost), fmt_decimal(pos.qty), fmt_money(pos.cost)])
        for row in result.exempt_income:
            self.annual_trees["isentos"].insert("", "end", values=[row["codigo"], row["descricao"], fmt_money(row["valor"])])
        for row in result.taxable_income:
            self.annual_trees["sujeitos"].insert("", "end", values=[row["codigo"], row["descricao"], fmt_money(row["valor"])])
        for row in result.debts:
            self.annual_trees["dividas"].insert(
                "",
                "end",
                values=[row["codigo"], row["descricao"], fmt_money(row["situacao_anterior"]), fmt_money(row["situacao_atual"]), fmt_money(row["valor_pago"])],
            )
        self._refresh_annual_totals(result)
        if self.results_by_year:
            self._refresh_history(selected_year)
            self._refresh_monthly_portfolio(selected_year)
            self._refresh_wealth_view()
            self._refresh_pending_options()

    def _refresh_annual_totals(self, result: CalculationResult) -> None:
        bens_anterior = Decimal("0")
        bens_atual = Decimal("0")
        for pos in result.positions.values():
            if pos.qty <= 0 or pos.category in {"opcoes", "futuro"}:
                continue
            bens_anterior += q2(pos.previous_cost)
            bens_atual += q2(pos.cost)
        isentos = sum((q2(row["valor"]) for row in result.exempt_income), Decimal("0"))
        sujeitos = sum((q2(row["valor"]) for row in result.taxable_income), Decimal("0"))
        dividas_anterior = sum((q2(row["situacao_anterior"]) for row in result.debts), Decimal("0"))
        dividas_atual = sum((q2(row["situacao_atual"]) for row in result.debts), Decimal("0"))
        dividas_pago = sum((q2(row["valor_pago"]) for row in result.debts), Decimal("0"))
        self.annual_total_vars["bens"].set(f"Totais: 31/12 anterior R$ {fmt_money(bens_anterior)} | 31/12 atual R$ {fmt_money(bens_atual)}")
        self.annual_total_vars["isentos"].set(f"Total de rendimentos isentos: R$ {fmt_money(isentos)}")
        self.annual_total_vars["sujeitos"].set(f"Total de rendimentos sujeitos exclusiva: R$ {fmt_money(sujeitos)}")
        self.annual_total_vars["dividas"].set(f"Totais: anterior R$ {fmt_money(dividas_anterior)} | atual R$ {fmt_money(dividas_atual)} | pago R$ {fmt_money(dividas_pago)}")

    def _monthly_rows(self, result: CalculationResult) -> list[list[Any]]:
        return [
            [
                MONTHS[m.month - 1],
                m.normal_result,
                m.daytrade_result,
                m.fii_result,
                m.options_result,
                m.future_result,
                m.normal_base,
                m.daytrade_base,
                m.fii_base,
                m.tax_payable,
            ]
            for m in result.monthly
        ]

    def _annual_rows(self, result: CalculationResult) -> dict[str, list[list[Any]]]:
        bens = []
        asset_cnpjs = self._asset_cnpj_map()
        for code, pos in sorted(result.positions.items()):
            if pos.qty <= 0 or pos.category in {"opcoes", "futuro"}:
                continue
            group = annual_asset_group(pos)
            discr = f"{code} - Quantidade: {fmt_decimal(pos.qty)} - Preco medio: R$ {fmt_money(pos.avg_price)}"
            bens.append([group, asset_cnpjs.get(normalize_ticker(code), ""), discr, fmt_decimal(pos.previous_qty), fmt_money(pos.previous_cost), fmt_decimal(pos.qty), fmt_money(pos.cost)])
        return {
            "bens": bens,
            "isentos": [[row["codigo"], *self._split_description_cnpj(row["descricao"]), fmt_money(row["valor"])] for row in result.exempt_income],
            "sujeitos": [[row["codigo"], *self._split_description_cnpj(row["descricao"]), fmt_money(row["valor"])] for row in result.taxable_income],
            "dividas": [
                [row["codigo"], self._cnpj_from_description(row["descricao"], asset_cnpjs), row["descricao"], fmt_money(row["situacao_anterior"]), fmt_money(row["situacao_atual"]), fmt_money(row["valor_pago"])]
                for row in result.debts
            ],
        }

    def _split_description_cnpj(self, description: str) -> list[str]:
        text = str(description or "")
        match = CNPJ_RE.search(text)
        if not match:
            return ["", text]
        cnpj = match.group(0)
        clean = re.sub(r"\s*-\s*CNPJ\s*" + re.escape(cnpj), "", text, flags=re.I).strip()
        clean = re.sub(r"\s*CNPJ\s*" + re.escape(cnpj), "", clean, flags=re.I).strip(" -")
        return [cnpj, clean]

    def _cnpj_from_description(self, description: str, asset_cnpjs: dict[str, str]) -> str:
        match = CNPJ_RE.search(str(description or ""))
        if match:
            return match.group(0)
        for ticker in re.findall(r"\b[A-Z]{4}\d{1,2}\b", str(description or "").upper()):
            cnpj = asset_cnpjs.get(normalize_ticker(ticker))
            if cnpj:
                return cnpj
        return ""

    def export_excel(self) -> None:
        if self.calculating:
            messagebox.showinfo(APP_TITLE, "Aguarde o calculo terminar antes de exportar.")
            return
        if not self.results_by_year:
            self.calculate()
            if not self.results_by_year:
                return
        if not self.results_by_year or Workbook is None:
            messagebox.showerror(APP_TITLE, "Nao foi possivel exportar Excel. Instale openpyxl.")
            return
        start_year = min(self.results_by_year)
        end_year = max(self.results_by_year)
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=f"IRSimple_{start_year}_{end_year}.xlsx",
            initialdir=str(BASE_DIR),
        )
        if not path:
            return
        wb = Workbook()
        first = True
        for year, result in self.results_by_year.items():
            ws = wb.active if first else wb.create_sheet()
            first = False
            ws.title = f"{year} mensal"
            self._write_sheet(ws, ["Mes", "Resultado normal", "Day trade", "FII", "Opcoes", "Futuro", "Base normal", "Base DT", "Base FII", "Imposto"], self._monthly_rows(result))

            annual = self._annual_rows(result)
            previous_year = year - 1
            sheet_defs = [
                (f"{year} bens", ["codigo", "cnpj", "discriminacao", f"qtd em 31/12/{previous_year}", f"valor em 31/12/{previous_year}", f"qtd em 31/12/{year}", f"valor em 31/12/{year}"], annual["bens"]),
                (f"{year} isentos", ["codigo", "cnpj", "descricao", "valor"], annual["isentos"]),
                (f"{year} exclusiva", ["codigo", "cnpj", "descricao", "valor"], annual["sujeitos"]),
                (f"{year} dividas", ["codigo", "cnpj", "descricao", f"31/12/{previous_year}", f"31/12/{year}", f"valor pago em {year}"], annual["dividas"]),
            ]
            for title, headers, rows in sheet_defs:
                ws = wb.create_sheet(title[:31])
                self._write_sheet(ws, headers, rows)
                for row in ws.iter_rows():
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws = wb.create_sheet(f"{year} alertas")
            for idx, line in enumerate(result.warnings or ["Sem alertas criticos."], start=1):
                ws.cell(idx, 1, line)
            portfolio_rows = []
            for month in range(1, 13):
                loans = result.monthly_loans.get(month, {})
                for code, pos in sorted(result.monthly_positions.get(month, {}).items()):
                    if pos.category in {"opcoes", "futuro"}:
                        continue
                    if pos.qty <= 0 and loans.get(code, Decimal("0")) <= 0:
                        continue
                    portfolio_rows.append([MONTHS[month - 1], code, pos.category, pos.qty, pos.avg_price, pos.cost, loans.get(code, Decimal("0"))])
            ws = wb.create_sheet(f"{year} carteira mensal"[:31])
            self._write_sheet(ws, ["Mes", "Ativo", "Categoria", "Quantidade", "Preco medio", "Custo", "Em emprestimo"], portfolio_rows)
            option_rows = [
                [
                    row.get("codigo", ""),
                    row.get("vencimento").strftime("%d/%m/%Y") if isinstance(row.get("vencimento"), date) else "",
                    row.get("quantidade_aberta", Decimal("0")),
                    row.get("premio_liquido", Decimal("0")),
                    row.get("ultima_data").strftime("%d/%m/%Y") if isinstance(row.get("ultima_data"), date) else "",
                    row.get("situacao", ""),
                ]
                for row in result.pending_options
            ]
            ws = wb.create_sheet(f"{year} opcoes conferir"[:31])
            self._write_sheet(ws, ["Opcao", "Vencimento", "Qtd aberta", "Premio liquido", "Ultima operacao", "Situacao"], option_rows)
        wb.save(path)
        self.status_var.set(f"Excel gerado: {path}")

    def _write_sheet(self, ws: Any, headers: list[str], rows: list[list[Any]]) -> None:
        fill = PatternFill("solid", fgColor="D9EAF7")
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.font = Font(bold=True)
            cell.fill = fill
        for row_idx, row in enumerate(rows, 2):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row_idx, col_idx, float(value) if isinstance(value, Decimal) else value)
                if isinstance(value, Decimal):
                    cell.number_format = '#,##0.00'
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    def _write_tree_sheet(self, ws: Any, tree: ttk.Treeview) -> None:
        headers = list(tree["columns"])
        rows = [list(tree.item(item, "values")) for item in tree.get_children()]
        self._write_sheet(ws, headers, rows)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    def export_pdf(self) -> None:
        if self.calculating:
            messagebox.showinfo(APP_TITLE, "Aguarde o calculo terminar antes de exportar.")
            return
        if not self.results_by_year:
            self.calculate()
            if not self.results_by_year:
                return
        if not self.results_by_year or SimpleDocTemplate is None:
            messagebox.showerror(APP_TITLE, "Nao foi possivel exportar PDF. Instale reportlab.")
            return
        try:
            year = int(self.selected_year_var.get())
        except Exception:
            messagebox.showerror(APP_TITLE, "Selecione um ano-calendario valido para gerar o PDF.")
            return
        result = self.results_by_year.get(year)
        if result is None:
            messagebox.showerror(APP_TITLE, f"O ano-calendario {year} ainda nao foi calculado.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF", "*.pdf")],
            initialfile=f"IRSimple_exercicio_{year + 1}_ano_base_{year}.pdf",
            initialdir=str(BASE_DIR),
        )
        if not path:
            return
        styles = getSampleStyleSheet()
        doc = SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=1 * cm, rightMargin=1 * cm, topMargin=1 * cm, bottomMargin=1 * cm)
        story: list[Any] = []
        self._pdf_header(story, styles, "Impostos em Renda Variavel", year)
        for idx, month in enumerate(result.monthly):
            if idx:
                story.append(PageBreak())
                self._pdf_header(story, styles, "Impostos em Renda Variavel", year)
            self._pdf_variable_month(story, month)

        story.append(PageBreak())
        self._pdf_header(story, styles, "Fundos Imobiliarios", year)
        self._pdf_fii_monthly(story, result)

        annual = self._annual_rows(result)
        previous_year = year - 1
        for key, title, headers in [
            ("bens", "Bens e Direitos", ["Codigo", "CNPJ", "Discriminacao", f"Qtd em 31/12/{previous_year}", f"Situacao em 31/12/{previous_year}", f"Qtd em 31/12/{year}", f"Situacao em 31/12/{year}"]),
            ("isentos", "Rendimentos Isentos e Nao Tributaveis", ["Codigo", "CNPJ", "Descricao", "Valor"]),
            ("sujeitos", "Rendimentos Sujeitos a Tributacao Exclusiva", ["Codigo", "CNPJ", "Descricao", "Valor"]),
            ("dividas", "Onus e Dividas", ["Codigo", "CNPJ", "Discriminacao", f"Situacao em 31/12/{previous_year}", f"Situacao em 31/12/{year}", f"Valor pago em {year}"]),
        ]:
            story.append(PageBreak())
            self._pdf_header(story, styles, title, year)
            self._pdf_table(story, title, headers, annual[key])
        self._pdf_consolidated_reports(story, styles, year)
        doc.build(story)
        self.status_var.set(f"PDF do exercicio {year + 1} gerado: {path}")

    def _pdf_monthly_income_rows(self, result: CalculationResult) -> list[list[Any]]:
        rows: list[list[Any]] = []
        for m in result.monthly:
            month = MONTHS[m.month - 1]
            future_dt = q2(m.future_dollar_daytrade + m.future_index_daytrade)
            future_common = q2(m.future_result - future_dt)
            common_result = q2(m.normal_result + m.options_result + future_common)
            daytrade_result = q2(m.daytrade_result + future_dt)
            common_tax = q2(m.normal_base * Decimal("0.15"))
            daytrade_tax = q2(m.daytrade_base * Decimal("0.20"))
            rows.extend(
                [
                    [month, "Mercado a Vista", "Mercado a vista - acoes", fmt_money(m.normal_result), fmt_money(m.daytrade_result)],
                    [month, "Mercado de Opcoes", "Mercado opcoes - acoes", fmt_money(m.options_result), "-"],
                    [month, "Mercado Futuro", "Mercado futuro - dolar dos EUA", fmt_money(m.future_dollar_common), fmt_money(m.future_dollar_daytrade)],
                    [month, "Mercado Futuro", "Mercado futuro - indices", fmt_money(m.future_index_common), fmt_money(m.future_index_daytrade)],
                    [month, "Resultados", "RESULTADO LIQUIDO DO MES", fmt_money(common_result), fmt_money(daytrade_result)],
                    [month, "Resultados", "Resultado negativo ate o mes anterior", fmt_money(m.normal_loss_before), fmt_money(m.daytrade_loss_before)],
                    [month, "Resultados", "BASE DE CALCULO DO IMPOSTO", fmt_money(m.normal_base), fmt_money(m.daytrade_base)],
                    [month, "Resultados", "Prejuizo a compensar", fmt_money(m.normal_loss_after), fmt_money(m.daytrade_loss_after)],
                    [month, "Resultados", "Aliquota do imposto", "15,00 %", "20,00 %"],
                    [month, "Resultados", "IMPOSTO DEVIDO", fmt_money(common_tax), fmt_money(daytrade_tax)],
                    [month, "Consolidacao do Mes", "Total de imposto devido", fmt_money(m.tax_due), ""],
                    [month, "Consolidacao do Mes", "Imposto a pagar", fmt_money(m.tax_payable), ""],
                ]
            )
        return rows

    def _pdf_variable_month(self, story: list[Any], m: MonthlyTax) -> None:
        month = MONTHS[m.month - 1]
        future_dt = q2(m.future_dollar_daytrade + m.future_index_daytrade)
        future_common = q2(m.future_result - future_dt)
        common_result = q2(m.normal_result + m.options_result + future_common)
        daytrade_result = q2(m.daytrade_result + future_dt)
        common_tax = q2(m.normal_base * Decimal("0.15"))
        daytrade_tax = q2(m.daytrade_base * Decimal("0.20"))
        rows = [
            ["Mercado a Vista", "Operacoes Comuns", "Day Trade"],
            ["Mercado a vista - acoes", fmt_money(m.normal_result), fmt_money(m.daytrade_result)],
            ["Mercado a vista - ouro", "-", "-"],
            ["Mercado a vista - ouro at. fin. fora bolsa", "-", "-"],
            ["Mercado de Opcoes", "Operacoes Comuns", "Day Trade"],
            ["Mercado opcoes - acoes", fmt_money(m.options_result), "-"],
            ["Mercado opcoes - ouro", "-", "-"],
            ["Mercado opcoes - fora de bolsa", "-", "-"],
            ["Mercado opcoes - outros", "-", "-"],
            ["Mercado Futuro", "Operacoes Comuns", "Day Trade"],
            ["Mercado futuro - dolar dos EUA", fmt_money(m.future_dollar_common), fmt_money(m.future_dollar_daytrade)],
            ["Mercado futuro - indices", fmt_money(m.future_index_common), fmt_money(m.future_index_daytrade)],
            ["Mercado futuro - juros", "-", "-"],
            ["Mercado futuro - outros", fmt_money(q2(future_common - m.future_dollar_common - m.future_index_common)), fmt_money(q2(future_dt - m.future_dollar_daytrade - m.future_index_daytrade))],
            ["Mercado a Termo", "Operacoes Comuns", "Day Trade"],
            ["Mercado a termo - acoes/ouro", "-", "-"],
            ["Mercado a termo - outros", "-", "-"],
            ["Resultados", "Operacoes Comuns", "Day Trade"],
            ["RESULTADO LIQUIDO DO MES", fmt_money(common_result), fmt_money(daytrade_result)],
            ["Resultado negativo ate o mes anterior", fmt_money(m.normal_loss_before), fmt_money(m.daytrade_loss_before)],
            ["BASE DE CALCULO DO IMPOSTO", fmt_money(m.normal_base), fmt_money(m.daytrade_base)],
            ["Prejuizo a compensar", fmt_money(m.normal_loss_after), fmt_money(m.daytrade_loss_after)],
            ["Aliquota do imposto", "15%", "20%"],
            ["IMPOSTO DEVIDO", fmt_money(common_tax), fmt_money(daytrade_tax)],
            ["Consolidacao do Mes", "", ""],
            ["Total do imposto devido", fmt_money(common_tax), fmt_money(daytrade_tax)],
            ["IR fonte de Day Trade no mes", "", fmt_money(m.irrf_daytrade_month)],
            ["IR fonte de Day Trade nos meses anteriores", "", fmt_money(m.irrf_daytrade_before)],
            ["IR fonte de Day Trade a compensar", "", fmt_money(m.irrf_daytrade_after)],
            ["IR fonte(Lei no 11.033/2004) no mes", fmt_money(m.irrf_common_month), ""],
            ["IR fonte(Lei no 11.033/2004) nos meses anteriores", fmt_money(m.irrf_common_before), ""],
            ["IR fonte(Lei no 11.033/2004) a compensar", fmt_money(m.irrf_common_after), ""],
            ["Imposto a pagar", fmt_money(max(Decimal("0"), common_tax - m.irrf_common_before - m.irrf_common_month)), fmt_money(max(Decimal("0"), daytrade_tax - m.irrf_daytrade_before - m.irrf_daytrade_month))],
            ["Imposto pago (valor + imposto acumulado + multa + juros)", "0,00", "0,00"],
        ]
        self._pdf_table(story, f"Ganhos Liquidos ou Perdas em {month}", ["", "", ""], rows, section_rows={0, 4, 9, 14, 17, 24})

    def _pdf_fii_monthly(self, story: list[Any], result: CalculationResult) -> None:
        rows = [
            [
                MONTHS[m.month - 1],
                fmt_money(m.fii_result),
                fmt_money(m.fii_loss_before),
                fmt_money(m.fii_base),
                fmt_money(m.fii_loss_after),
                "20,00 %",
                fmt_money(m.fii_base * Decimal("0.20")),
                fmt_money(m.irrf_fii_before),
                fmt_money(m.irrf_fii_month),
                fmt_money(m.irrf_fii_after),
                fmt_money(max(Decimal("0"), m.fii_base * Decimal("0.20") - m.irrf_fii_before - m.irrf_fii_month)),
                "-",
            ]
            for m in result.monthly
        ]
        self._pdf_table(
            story,
            "Ganhos Liquidos ou Perdas",
            [
                "Mes",
                "Resultado liquido no mes",
                "Resultado negativo ate o mes anterior",
                "Base de calculo do imposto",
                "Prejuizo a compensar",
                "Aliquota do imposto",
                "Imposto devido",
                "Saldo do imposto retido nos meses anteriores",
                "Imposto retido no mes",
                "Imposto a compensar",
                "Imposto a pagar",
                "Imposto pago",
            ],
            rows,
        )

    def _pdf_consolidated_reports(self, story: list[Any], styles: Any, year: int) -> None:
        reports = self._consolidated_files_for_year(year)
        if not reports:
            return
        for path in reports:
            try:
                sheets = read_consolidated_sheets(path)
            except Exception as exc:
                story.append(PageBreak())
                self._pdf_header(story, styles, "Relatorio Consolidado B3", year)
                self._pdf_table(story, path.name, ["Erro"], [[f"Falha ao ler arquivo: {exc}"]])
                continue
            for sheet in sheets:
                story.append(PageBreak())
                self._pdf_header(story, styles, "Relatorio Consolidado B3", year)
                rows = sheet.get("rows") or [[""]]
                columns = [str(col) for col in (sheet.get("columns") or ["Sem dados"])]
                normalized_rows = [list(row[: len(columns)]) + [""] * max(0, len(columns) - len(row)) for row in rows]
                title = f"{path.name} - {sheet.get('name') or 'Sheet'}"
                self._pdf_table(story, title, columns, normalized_rows)

    def _consolidated_files_for_year(self, year: int) -> list[Path]:
        paths: list[Path] = []
        for file in self.consolidated_files:
            path = Path(file)
            if not path.exists():
                continue
            period = parse_consolidated_period(path)
            if period and period[0] == year:
                paths.append(path)
        return sorted(paths, key=lambda path: (parse_consolidated_period(path) or (year, 99))[1])

    def _pdf_header(self, story: list[Any], styles: Any, subtitle: str, year: int) -> None:
        story.append(Paragraph("DECLARACAO ANUAL DE IMPOSTO DE RENDA", styles["Title"]))
        story.append(Paragraph(f"Nome: {self.name_var.get() or '-'}", styles["Normal"]))
        story.append(Paragraph(f"{subtitle} - Ano base: {year} - Exercicio: {year + 1}", styles["Heading2"]))
        story.append(Spacer(1, 0.25 * cm))

    def _pdf_table(self, story: list[Any], title: str, headers: list[str], rows: list[list[Any]], section_rows: set[int] | None = None) -> None:
        styles = getSampleStyleSheet()
        story.append(Paragraph(title, styles["Heading2"]))
        normal = styles["Normal"]
        normal.fontSize = 7
        normal.leading = 8
        source_rows = rows if rows else [["-" for _ in headers]]
        data = [[Paragraph(str(cell), normal) for cell in headers]] + [[Paragraph(str(cell), normal) for cell in row] for row in source_rows]
        table = Table(data, repeatRows=1)
        style_commands = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaeaea")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c9c9c9")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]
        for idx in section_rows or set():
            row = idx + 1
            style_commands.extend(
                [
                    ("BACKGROUND", (0, row), (-1, row), colors.HexColor("#f1f1f1")),
                    ("FONTNAME", (0, row), (-1, row), "Helvetica-Bold"),
                ]
            )
        table.setStyle(TableStyle(style_commands))
        story.append(table)
        story.append(Spacer(1, 0.2 * cm))

    def run(self) -> int:
        self.root.mainloop()
        return 0


def main() -> int:
    app = IRSimpleApp()
    return app.run()


if __name__ == "__main__":
    raise SystemExit(main())
